from app.core.resources import resource_base, resource_path
import os
import json
import base64
import math
import re
import io
import unicodedata
from typing import Dict, Any, Optional, Tuple, List, Set
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import sys

import openpyxl
from openpyxl.utils import get_column_letter

from PySide6.QtCore import QUrl, QObject, Slot
from PySide6.QtWidgets import (
    QMainWindow, QWidget, QVBoxLayout, QFileDialog, QMessageBox
)
from PySide6.QtWebEngineWidgets import QWebEngineView
from PySide6.QtWebEngineCore import QWebEnginePage, QWebEngineProfile
from PySide6.QtWebChannel import QWebChannel

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from PIL import Image

from app.ui.topbar import TopBar

# =========================
# HTML (UI + Leaflet)
# =========================

HTML = r"""<!doctype html>
<html lang="pt-br">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1.0" />
  <title>Mapa de Calor</title>

  <link rel="stylesheet" href="https://unpkg.com/leaflet@1.9.4/dist/leaflet.css" crossorigin=""/>

  <style>
    html,body{height:100%;margin:0;font-family:system-ui,Arial,sans-serif;}
    .app{display:grid;grid-template-columns:460px 1fr;height:100%}
    .panel{border-right:1px solid #eee;padding:14px 16px 16px;overflow:auto;background:#fff}
    h1{font-size:18px;margin:0 0 12px}
    h2{font-size:14px;margin:16px 0 8px}
    label{font-size:13px;display:block;margin:8px 0 6px}
    small{color:#666;font-size:12px}
    input[type="file"],button,input[type="number"],select{
      font:inherit;padding:8px 10px;border-radius:10px;border:1px solid #ddd;background:#fff
    }
    .btn-primary{background:#1E90FF;border-color:#1E90FF;color:#fff}
    .btn-outline{background:#fff;color:#111;border-color:#ccc}
    .btn-sm{padding:6px 8px;border-radius:10px;font-size:12px}
    .inline{display:flex;gap:8px;align-items:center;flex-wrap:wrap}
    .muted{color:#666;font-size:12px}

    #map{height:100%;width:100%;background:#e9e9e9}
    .leaflet-tile{filter:grayscale(1) brightness(1) contrast(1)}

    .manual-row{display:flex; flex-direction: column; gap:8px; padding:16px 0; border-bottom:1px dashed #eee}
    .manual-header{display:flex; align-items: center; justify-content: space-between; gap: 10px;}
    .manual-name{font-size:13px; font-weight: 600; color: #333; flex: 1; white-space: nowrap; overflow: hidden; text-overflow: ellipsis;}
    .num{width:70px; text-align:center; border-radius:10px; border:1px solid #ddd; padding:6px 8px; font-weight:700; font-size:14px}

    .box{
      border:1px solid #eee;border-radius:14px;padding:10px 10px 8px;background:#fafafa;
      margin-bottom:10px;
    }
    .stats{
      display:flex; gap:10px; flex-wrap:wrap; font-size:12px; color:#333;
      margin-top:8px;
    }
    .pill{
      padding:4px 8px;border-radius:999px;background:#fff;border:1px solid #e6e6e6;
      font-weight:700;
      cursor: help;
      white-space: nowrap;
    }

    /* Dropdown custom */
    .dd{position:relative;}
    .dd-btn{
      width:100%;
      text-align:left;
      padding:10px 10px;
      border-radius:12px;
      border:1px solid #ddd;
      background:#fff;
      cursor:pointer;
      font:inherit;
      display:flex;
      align-items:center;
      justify-content:space-between;
      gap:10px;
    }
    .dd-btn:disabled{opacity:.6;cursor:not-allowed}
    .dd-arrow{color:#666;font-size:12px}
    .dd-menu{
      position:absolute;
      left:0; right:0;
      top:calc(100% + 6px);
      background:#fff;
      border:1px solid #ddd;
      border-radius:12px;
      max-height:260px;
      overflow:auto;
      box-shadow:0 12px 28px rgba(0,0,0,.10);
      z-index:2000;
    }
    .dd-item{
      padding:10px 10px;
      cursor:pointer;
      font-size:13px;
      color:#111;
      user-select:none;
    }
    .dd-item:hover{background:#f2f2f2}
    .dd-item.is-selected{background:#ededed;font-weight:800}

    .badge{
      display:inline-flex;align-items:center;justify-content:center;
      font-size:10px;font-weight:900;
      padding:2px 6px;border-radius:999px;
      border:1px solid #ddd;background:#fff;color:#333;
      margin-left:6px;
    }
    .badge.manual{border-color:#cfd8ff;background:#eef2ff;color:#1b2a6b;}
    .badge.auto{border-color:#d9d9d9;background:#f6f6f6;color:#333;}

    input[type="color"]{
      width: 34px; height: 34px; padding: 0; border: 1px solid #ddd; border-radius: 10px;
      cursor: pointer; background: none;
    }

    #pptQueueList{white-space:pre-wrap; line-height:1.35}
  </style>
</head>
<body>
<div class="app">
  <aside class="panel">
    <h1>Mapa de Calor</h1>

    <div class="box">
      <h2 style="margin:0 0 8px">0) Excel (consulta)</h2>
      <label>Importar Relatório Descritivo Quantitivo</label>
      <div class="inline">
        <button id="btnPickExcel" class="btn-outline">Importar Excel</button>
        <small id="excelStatus" class="muted"></small>
      </div>

      <h2>1) KMLs</h2>
      <label>Importar KML (vários):</label>
      <input id="kmlInput" type="file" accept=".kml" multiple />

      <h2>2) Configurações Visuais</h2>
      <div class="inline">
        <label for="baseSelect" style="margin:0">Fundo:</label>
        <select id="baseSelect">
          <option value="osm">OSM (Padrão)</option>
          <option value="esri">Satélite (Esri)</option>
          <option value="none">Sem mapa</option>
        </select>
        <button id="btnToggleLabels" class="btn-outline">Ocultar números</button>
      </div>

      <div style="margin-top:10px">
        <label style="margin:0 0 4px">Opacidade das cores (%)</label>
        <div class="inline">
          <input id="opacityRange" type="range" min="0" max="100" value="60" style="width:260px" />
          <input id="opacityNumber" type="number" min="0" max="100" value="60" style="width:70px" />
        </div>
      </div>

      <div style="margin-top:10px">
        <label style="margin:0 0 4px">Espessura das bordas (px)</label>
        <div class="inline">
          <input id="borderRange" type="range" min="0" max="10" step="0.5" value="2" style="width:260px" />
          <input id="borderNumber" type="number" min="0" max="10" step="0.5" value="2" style="width:70px" />
        </div>
      </div>

      <div style="margin-top:10px">
        <label style="margin:0 0 4px">Tamanho dos números (px)</label>
        <div class="inline">
          <input id="fontRange" type="range" min="10" max="80" step="1" value="28" style="width:260px" />
          <input id="fontNumber" type="number" min="10" max="80" step="1" value="28" style="width:70px" />
        </div>
      </div>

      <h2>3) Configurações dos casos e Segmentos</h2>
      <label>Caso</label>
      <div class="dd" id="caseDD" data-placeholder="—"></div>

      <label>Segmento</label>
      <div class="dd" id="segmentDD" data-placeholder="—"></div>

      <label>Base (Total) (manual)</label>
      <div class="inline">
        <input id="baseOverride" type="number" min="1" step="1" placeholder="(vazio = usar base do Excel)" style="width:180px" />
        <button id="btnUseExcelBase" class="btn-outline" type="button">Usar base do Excel</button>
      </div>
      <small id="excelBaseHint" class="muted"></small>

      <div class="stats" id="statsRow" style="display:none">
        <span class="pill" id="statMin">Mín: —</span>
        <span class="pill" id="statMean">Média —</span>
        <span class="pill" id="statMax">Máx: —</span>
      </div>

      <small class="muted">Passe o mouse nos pills para ver o detalhamento do cálculo (células/valores usados).</small>

      <h2>4) Pintura do Mapa</h2>
      <h3 style="margin:12px 0 6px;font-size:14px">Pintura automática</h3>

      <div class="inline" style="margin:6px 0 8px">
        <button id="btnApplyMap" class="btn-primary" type="button">Aplicar ao mapa</button>
        <label style="display:flex;gap:8px;align-items:center;margin:0;font-size:12px;color:#333">
          <input id="chkNoOverwriteManual" type="checkbox" checked />
          Não sobrescrever pintura manual
        </label>
      </div>

      <div class="box" style="background:#fff;margin-top:8px">
        <div class="inline" style="justify-content:space-between; width:100%">
          <div style="font-weight:780">Paleta ativa (automática)</div>
          <div class="muted" id="paletteGroupHint">Grupo: —</div>
        </div>

        <div class="inline" style="margin-top:8px">
          <div class="muted" style="font-weight:900;width:60px">Baixo</div>
          <input id="palLow" type="color" value="#87cefa" title="Baixo">
          <div class="muted" style="font-weight:900;width:60px;margin-left:10px">Médio</div>
          <input id="palMid" type="color" value="#1e90ff" title="Médio">
          <div class="muted" style="font-weight:900;width:60px;margin-left:10px">Alto</div>
          <input id="palHigh" type="color" value="#121958" title="Alto">
          <button id="btnResetPalette" class="btn-outline btn-sm" type="button" style="margin-left:8px">Restaurar padrão</button>
        </div>

        <div class="muted" id="autoPaintStatus" style="margin-top:10px">—</div>
      </div>

      <h2 style="margin:12px 0 6px;font-size:14px">PPTX (vários mapas)</h2>
      <div class="inline" style="margin:6px 0 8px">
        <button id="btnPickLogo" class="btn-outline" type="button">Carregar logo</button>
        <small id="logoStatus" class="muted">Sem logo</small>
      </div>


      <div class="inline" style="margin:6px 0 8px">
        <button id="btnAddToQueue" class="btn-outline" type="button">Adicionar à fila</button>
        <button id="btnClearQueue" class="btn-outline" type="button">Limpar fila</button>
        <button id="btnExportPptx" class="btn-primary" type="button">Exportar PPTX</button>
        <small id="pptQueueStatus" class="muted"></small>
      </div>

      <div id="pptQueueList" class="muted"></div>
    </div>


    <h2>5) Numeração</h2>
    <div id="manualPanel" class="muted">Carregue KMLs para listar as malhas…</div>
    <div class="inline" style="margin-top:12px">
      <button id="btnClearPaint" class="btn-outline">Limpar pintura</button>
      <small class="muted">Limpa todas as regiões (manual + automática)</small>
    </div>


    <h2>6) Exportação</h2>
    <div class="inline">
      <button id="btnExport" class="btn-primary">Exportar PNG</button>
      <small id="expStatus"></small>
    </div>
  </aside>

  <main id="map"></main>
</div>

<script src="qrc:///qtwebchannel/qwebchannel.js"></script>
<script>
  window.QtBridge = null;
  new QWebChannel(qt.webChannelTransport, function(channel){
    window.QtBridge = channel.objects.QtBridge;
  });

  function Dropdown(rootEl){
    const placeholder = rootEl.dataset.placeholder || '—';
    let items = [];
    let value = '';
    let enabled = false;

    const btn = document.createElement('button');
    btn.type = 'button';
    btn.className = 'dd-btn';
    btn.disabled = true;
    btn.innerHTML = `<span class="dd-text">${placeholder}</span><span class="dd-arrow">▼</span>`;

    const menu = document.createElement('div');
    menu.className = 'dd-menu';
    menu.hidden = true;

    rootEl.appendChild(btn);
    rootEl.appendChild(menu);

    function close(){ menu.hidden = true; }
    function open(){ if(enabled && items.length) menu.hidden = false; }
    function setText(t){ btn.querySelector('.dd-text').textContent = t || placeholder; }

    function render(){
      menu.innerHTML = '';
      items.forEach(it=>{
        const div = document.createElement('div');
        div.className = 'dd-item' + (it===value ? ' is-selected' : '');
        div.textContent = it;
        div.addEventListener('click', ()=>{
          value = it;
          setText(it);
          render();
          close();
          api.onChange && api.onChange(it);
        });
        menu.appendChild(div);
      });
    }

    btn.addEventListener('click', ()=>{ if(menu.hidden) open(); else close(); });
    document.addEventListener('click', (e)=>{ if(!rootEl.contains(e.target)) close(); });

    const api = {
      setEnabled(v){
        enabled = !!v;
        btn.disabled = !enabled;
        if(!enabled){ value=''; setText(''); close(); }
      },
      setItems(arr){
        items = Array.isArray(arr) ? arr.slice() : [];
        value = '';
        setText('');
        render();
        close();
      },
      getValue(){ return value; },
      clear(){ items=[]; value=''; setText(''); render(); close(); },
      onChange: null
    };
    return api;
  }

  function el(id){ return document.getElementById(id); }

  window.__ddCase = Dropdown(document.getElementById('caseDD'));
  window.__ddSeg  = Dropdown(document.getElementById('segmentDD'));
  window.__segmentsByCase = {};
  window.__baseByCase = {};

  // último stats "cru" (para PPTX)
  window.__lastStatsRaw = null; // {caseKey, segment, min, mean, max}

  // Paletas por grupo (editáveis)
  window.__palettes = {
    good: { low:'#87cefa', mid:'#1e90ff', high:'#121958' },
    bad:  { low:'#ffdf9e', mid:'#ffa500', high:'#843c0c' },
    defaults: {
      good: { low:'#87cefa', mid:'#1e90ff', high:'#121958' },
      bad:  { low:'#ffdf9e', mid:'#ffa500', high:'#843c0c' },
    }
  };

  function normSeg(s){
    return String(s||'')
      .normalize('NFD').replace(/[\u0300-\u036f]/g,'')
      .toLowerCase().trim();
  }
  function isGoodSegment(seg){
    const t = normSeg(seg);
    return t.includes('certamente voto') || t.includes('aprovacao') || t.includes('aprovação');
  }
  function activeGroup(){
    const seg = window.__ddSeg.getValue();
    if(!seg) return null;
    return isGoodSegment(seg) ? 'good' : 'bad';
  }
  function syncPaletteUI(){
    const g = activeGroup();
    if(!g){
      el('paletteGroupHint').textContent = 'Grupo: —';
      return;
    }
    el('paletteGroupHint').textContent = 'Grupo: ' + (g==='good' ? 'Bom' : 'Ruim');
    el('palLow').value  = window.__palettes[g].low;
    el('palMid').value  = window.__palettes[g].mid;
    el('palHigh').value = window.__palettes[g].high;
  }
    el('pptQueueList').addEventListener('click', (ev)=>{
        const btn = ev.target.closest('button[data-action="removeSlide"]');
        if(!btn) return;
        const idx = parseInt(btn.dataset.index, 10);
        if(Number.isFinite(idx) && window.QtBridge?.removePptItem){
            window.QtBridge.removePptItem(idx);
        }
    });

  function updatePaletteFromUI(){
    const g = activeGroup();
    if(!g) return;
    window.__palettes[g].low  = el('palLow').value;
    window.__palettes[g].mid  = el('palMid').value;
    window.__palettes[g].high = el('palHigh').value;
  }

  el('palLow').addEventListener('input', updatePaletteFromUI);
  el('palMid').addEventListener('input', updatePaletteFromUI);
  el('palHigh').addEventListener('input', updatePaletteFromUI);

  el('btnResetPalette').addEventListener('click', ()=>{
    const g = activeGroup();
    if(!g) return;
    window.__palettes[g] = JSON.parse(JSON.stringify(window.__palettes.defaults[g]));
    syncPaletteUI();
  });

  window.__setCases = function(payloadJson){
    try{
      const data = JSON.parse(payloadJson);
      const status  = document.getElementById('excelStatus');
      const statsRow = document.getElementById('statsRow');

      window.__segmentsByCase = data.segmentsByCase || {};
      window.__baseByCase = data.baseByCase || {};

      window.__ddCase.setItems(data.cases || []);
      window.__ddCase.setEnabled((data.cases||[]).length>0);

      window.__ddSeg.clear();
      window.__ddSeg.setEnabled(false);

      let msg = (data.cases||[]).length ? `${data.cases.length} caso(s) encontrado(s)` : 'Nenhum caso encontrado';
      if(data.debugPath){ msg += ` | Debug: ${data.debugPath}`; }
      status.textContent = msg;

      statsRow.style.display = 'none';
      el('autoPaintStatus').textContent = '—';
    }catch(e){ console.error(e); }
  };

  // Agora recebe também rawMin/rawMean/rawMax (números) no final
  window.__setStats = function(minV, meanV, maxV, tipMin, tipMean, tipMax, rawMin, rawMean, rawMax){
    const statsRow = document.getElementById('statsRow');
    statsRow.style.display = 'flex';

    const elMin  = document.getElementById('statMin');
    const elMean = document.getElementById('statMean');
    const elMax  = document.getElementById('statMax');

    elMin.textContent  = `Mín: ${minV}`;
    elMean.textContent = `Média ${meanV}`;
    elMax.textContent  = `Máx: ${maxV}`;

    elMin.title  = tipMin || '';
    elMean.title = tipMean || '';
    elMax.title  = tipMax || '';

    const caseKey = window.__ddCase.getValue();
    const seg = window.__ddSeg.getValue();
    if(typeof rawMin === 'number' && typeof rawMean === 'number' && typeof rawMax === 'number'){
      window.__lastStatsRaw = { caseKey, segment: seg, min: rawMin, mean: rawMean, max: rawMax };
    } else {
      window.__lastStatsRaw = null;
    }
  };

    window.__setPptQueueState = function(payloadJson){
        try{
            const data = JSON.parse(payloadJson);
            el('pptQueueStatus').textContent = `Fila: ${data.count} mapa(s)`;
            el('logoStatus').textContent = data.logoLoaded ? 'Logo carregada' : 'Sem logo';

            const box = el('pptQueueList');
            box.innerHTML = '';

            (data.items || []).forEach((it, idx)=>{
                const index = (typeof it === 'string') ? idx : it.index;
                const label = (typeof it === 'string') ? it : it.label;

                const row = document.createElement('div');
                row.className = 'inline';
                row.style.justifyContent = 'space-between';
                row.style.gap = '10px';
                row.style.padding = '6px 0';
                row.style.borderBottom = '1px dashed #eee';

                const text = document.createElement('div');
                text.textContent = `${index+1}. ${label}`;
                text.style.flex = '1';
                text.style.whiteSpace = 'nowrap';
                text.style.overflow = 'hidden';
                text.style.textOverflow = 'ellipsis';

                const btn = document.createElement('button');
                btn.className = 'btn-outline btn-sm';
                btn.type = 'button';
                btn.textContent = 'Remover';
                btn.dataset.action = 'removeSlide';
                btn.dataset.index = String(index);

                row.appendChild(text);
                row.appendChild(btn);
                box.appendChild(row);
            });

        }catch(e){ console.error(e); }
    };


  // Excel UI
  el('btnPickExcel').addEventListener('click', ()=>{
    if(window.QtBridge?.pickExcel) window.QtBridge.pickExcel();
  });

  window.__ddCase.onChange = (caseKey)=>{
    const segs = (window.__segmentsByCase && window.__segmentsByCase[caseKey]) ? window.__segmentsByCase[caseKey] : [];
    window.__ddSeg.setItems(segs);
    window.__ddSeg.setEnabled(!!caseKey && segs.length>0);
    el('statsRow').style.display = 'none';

    const excelBase = (window.__baseByCase && window.__baseByCase[caseKey]) ? window.__baseByCase[caseKey] : '';
    const baseInput = el('baseOverride');
    const hint = el('excelBaseHint');

    hint.textContent = excelBase ? `Base do Excel (detectada): ${excelBase}` : 'Base do Excel: não detectada para este caso.';
    if(excelBase) baseInput.value = excelBase;

    syncPaletteUI();
  };

  window.__ddSeg.onChange = (segment)=>{
    syncPaletteUI();
    const caseKey = window.__ddCase.getValue();
    const baseOverride = (el('baseOverride').value || '').trim();
    if(caseKey && segment && window.QtBridge?.computeStats){
      window.QtBridge.computeStats(caseKey, segment, baseOverride);
    }
  };

  el('baseOverride').addEventListener('input', ()=>{
    const caseKey = window.__ddCase.getValue();
    const segment = window.__ddSeg.getValue();
    const baseOverride = (el('baseOverride').value || '').trim();
    if(caseKey && segment && window.QtBridge?.computeStats){
      window.QtBridge.computeStats(caseKey, segment, baseOverride);
    }
  });

  el('btnUseExcelBase').addEventListener('click', ()=>{
    const caseKey = window.__ddCase.getValue();
    const excelBase = (window.__baseByCase && window.__baseByCase[caseKey]) ? window.__baseByCase[caseKey] : '';
    if (excelBase) el('baseOverride').value = excelBase;

    const segment = window.__ddSeg.getValue();
    const baseOverride = (el('baseOverride').value || '').trim();
    if(caseKey && segment && window.QtBridge?.computeStats){
      window.QtBridge.computeStats(caseKey, segment, baseOverride);
    }
  });

  el('btnApplyMap').addEventListener('click', ()=>{
    const caseKey = window.__ddCase.getValue();
    const segment = window.__ddSeg.getValue();
    const baseOverride = (el('baseOverride').value || '').trim();
    const noOverwriteManual = !!el('chkNoOverwriteManual').checked;

    if(!caseKey || !segment){
      el('autoPaintStatus').textContent = 'Selecione um Caso e um Segmento antes de aplicar.';
      return;
    }
    if(window.QtBridge?.applyAutoPaint){
      window.QtBridge.applyAutoPaint(caseKey, segment, baseOverride, noOverwriteManual);
    }
  });

  // PPTX UI
  el('btnPickLogo').addEventListener('click', ()=>{
    if(window.QtBridge?.pickLogo) window.QtBridge.pickLogo();
  });

  el('btnClearQueue').addEventListener('click', ()=>{
    if(window.QtBridge?.clearPptQueue) window.QtBridge.clearPptQueue();
  });

  el('btnExportPptx').addEventListener('click', ()=>{
    if(window.QtBridge?.exportPptx) window.QtBridge.exportPptx();
  });

  el('btnAddToQueue').addEventListener('click', async ()=>{
    if(window.__enqueuePptSlide) await window.__enqueuePptSlide();
  });
</script>

<script src="https://unpkg.com/leaflet@1.9.4/dist/leaflet.js" crossorigin=""></script>
<script src="https://unpkg.com/leaflet-image/leaflet-image.js"></script>

<script type="module">
import { kml as kmlToGeoJSON } from "https://unpkg.com/@tmcw/togeojson?module";

function el(id){ return document.getElementById(id); }

/* ===== MAPA ===== */
const map = L.map('map', { zoomControl:true, preferCanvas:true }).setView([-15.78,-47.88], 11);
const canvasRenderer = L.canvas({ padding:0.5 });

const baseOSM  = L.tileLayer('https://{s}.tile.openstreetmap.org/{z}/{x}/{y}.png',
  { maxZoom:19, attribution:'© OpenStreetMap contributors', crossOrigin:'anonymous' });
const baseEsri = L.tileLayer('https://server.arcgisonline.com/ArcGIS/rest/services/World_Imagery/MapServer/tile/{z}/{y}/{x}',
  { maxZoom:19, attribution:'Tiles © Esri', crossOrigin:'anonymous' });

let currentBase = baseOSM.addTo(map);
let labelsVisible = true;

const kmlGroup = L.layerGroup().addTo(map);
const boundsGroup = L.featureGroup().addTo(map);

const layersByKey = new Map();
const displayName = new Map();
const colorByKey  = new Map();
const paintSourceByKey = new Map(); // "manual" | "auto"
const numberByKey = new Map();
const labelByKey  = new Map();
const labelPositionByKey = new Map();
const usedKeys    = new Set();

/* ===== NOVO: dados estruturados p/ legenda do PPT ===== */
const legendNameByKey  = new Map(); // key -> nome (ex.: "AP 5")
const legendValueByKey = new Map(); // key -> valor (%)
const legendClsByKey   = new Map(); // key -> low|mid|high (opcional)

let fillOpacityPct = 60;
let numFontSize = 28;
let borderWeight = 2;
const numFontFamily = `"Aptos (Corpo)", "Aptos", Calibri, DIN, Arial, sans-serif`;

/* ===== util normalização/match ===== */
function stripNumericPrefix(name){
  return String(name||'').replace(/^\s*\d+\s*[-–—.:]\s*/,'').trim();
}
function matchKey(name){
  let s = String(name || '');
  s = s.replace(/\r\n/g, '\n');
  s = s.replace(/-\s*\n?\s*/g, ''); // cola "Flumi- nense" -> "Fluminense"
  s = s.replace(/\u00AD/g, '');     // soft-hyphen invisível
  s = s
    .normalize('NFD').replace(/[\u0300-\u036f]/g,'')
    .toLowerCase()
    .replace(/[^a-z0-9]+/g,' ')
    .trim();
  return s;
}
function isGoodSegment(seg){
  const t = String(seg||'').normalize('NFD').replace(/[\u0300-\u036f]/g,'').toLowerCase().trim();
  return t.includes('certamente voto') || t.includes('aprovacao') || t.includes('aprovação');
}

/* ===== CONTROLES ===== */
const opacityRange  = el('opacityRange'), opacityNumber = el('opacityNumber');
const fontRange     = el('fontRange'), fontNumber = el('fontNumber');
const borderRange   = el('borderRange'), borderNumber = el('borderNumber');

opacityRange.addEventListener('input', e=> setOpacityPct(Number(e.target.value)));
opacityNumber.addEventListener('input', e=> setOpacityPct(Number(e.target.value)));
fontRange.addEventListener('input', e=> setFontSize(Number(e.target.value)));
fontNumber.addEventListener('input', e=> setFontSize(Number(e.target.value)));
borderRange.addEventListener('input', e=> setBorderWeight(Number(e.target.value)));
borderNumber.addEventListener('input', e=> setBorderWeight(Number(e.target.value)));

function setOpacityPct(v){ fillOpacityPct = Math.max(0, Math.min(100, v)); opacityRange.value = opacityNumber.value = fillOpacityPct; repaintAll(); }
function setFontSize(px){ numFontSize = Math.max(10, Math.min(80, px)); fontRange.value = fontNumber.value = numFontSize; numberByKey.forEach((_,k)=> updateLabel(k)); }
function setBorderWeight(w){ borderWeight = Math.max(0, Math.min(10, w)); borderRange.value = borderNumber.value = borderWeight; repaintAll(); }

const kmlInput=el('kmlInput'), manualPanel=el('manualPanel');
const btnClear=el('btnClearPaint'), btnExport=el('btnExport');
const baseSelect=el('baseSelect'), expStatus=el('expStatus'), btnToggleLabels=el('btnToggleLabels');

kmlInput.addEventListener('change', async (e)=>{
  resetAll();
  for(const file of e.target.files){
    if(!file.name.toLowerCase().endsWith('.kml')) continue;
    const text=await file.text();
    const dom=new DOMParser().parseFromString(text,'text/xml');
    const gj=kmlToGeoJSON(dom);

    const baseName=file.name.replace(/\.[A-Za-z0-9]+$/,'');
    const key=uniqueKey(canon(baseName));
    displayName.set(key, baseName);

    const layer = L.geoJSON(gj, {
      renderer: canvasRenderer,
      style: defaultStyle,
      onEachFeature: (f,ly)=>{
        f.properties=f.properties||{};
        f.properties.__fileKey=key;
        ly.bindPopup(`<b>${escapeHtml(baseName)}</b>`);
        registerLayer(key, ly);
      }
    });
    layer.addTo(kmlGroup); boundsGroup.addLayer(layer);
  }
  if(boundsGroup.getLayers().length){ try{ map.fitBounds(boundsGroup.getBounds(),{padding:[20,20]}); }catch{} }
  rebuildManualUI();
});

baseSelect.addEventListener('change', ()=>{
  if(currentBase) map.removeLayer(currentBase);
  const val = baseSelect.value;
  if(val==='esri') currentBase = baseEsri.addTo(map);
  else if(val==='osm') currentBase = baseOSM.addTo(map);
  else currentBase = null;
});

btnToggleLabels.addEventListener('click', ()=>{
  labelsVisible = !labelsVisible;
  btnToggleLabels.textContent = labelsVisible ? 'Ocultar números' : 'Mostrar números';
  labelByKey.forEach(m=> labelsVisible ? m.addTo(map) : map.removeLayer(m));
});

/* ===== Pintura manual/limpeza individual ===== */
manualPanel.addEventListener('click', (ev)=>{
  const btn = ev.target.closest('[data-action="clear"]');
  if(btn){
    const key = btn.dataset.key;
    clearRegionPaint(key);
    rebuildManualUI();
    return;
  }
});

manualPanel.addEventListener('input', (ev)=>{
  const input = ev.target.closest('.num');
  const colorPicker = ev.target.closest('.custom-color-input');
  if(input){
    const key=input.dataset.key, val=(input.value||'').trim();
    if(!val){ numberByKey.delete(key); removeLabel(key); }
    else { numberByKey.set(key,val); updateLabel(key); }
  }
  if(colorPicker){
    applyColor(colorPicker.dataset.key, colorPicker.value, 'manual');
    rebuildManualUI();
  }
});

function applyColor(key, hex, source='manual'){
  colorByKey.set(key, hex);
  paintSourceByKey.set(key, source);
  (layersByKey.get(key)||[]).forEach(ly=> setStyle(ly, {fillColor:hex}));
}

function clearRegionPaint(key){
  colorByKey.delete(key);
  paintSourceByKey.delete(key);

  // ✅ limpa dados da legenda do PPT para esta região
  legendNameByKey.delete(key);
  legendValueByKey.delete(key);
  legendClsByKey.delete(key);

  (layersByKey.get(key)||[]).forEach(ly=>{
    try{ ly.unbindPopup(); }catch{}
    setStyle(ly, {fillColor:'#ffffff'});
  });
}

btnClear.addEventListener('click', ()=>{
  colorByKey.clear();
  paintSourceByKey.clear();

  // ✅ limpa dados da legenda do PPT
  legendNameByKey.clear();
  legendValueByKey.clear();
  legendClsByKey.clear();

  layersByKey.forEach((list)=>{
    list.forEach(ly=>{
      try{ ly.unbindPopup(); }catch{}
    });
  });

  repaintAll();
  rebuildManualUI();
  el('autoPaintStatus').textContent = 'Pintura limpa.';
});

/* ===== UI KML list ===== */
function rebuildManualUI(){
  if(layersByKey.size===0){
    manualPanel.innerHTML = '<span class="muted">Carregue KMLs para listar as malhas…</span>';
    return;
  }
  const rows=[];
  for(const key of Array.from(layersByKey.keys()).sort()){
    const name = escapeHtml(displayName.get(key)||key);
    const currentColor = colorByKey.get(key);
    const val = escapeHtml(numberByKey.get(key)||'');
    const customHex = currentColor || '#ffffff';
    const src = paintSourceByKey.get(key);

    const badge = src ? `<span class="badge ${src==='manual'?'manual':'auto'}">${src==='manual'?'MANUAL':'AUTO'}</span>` : '';

    rows.push(`
      <div class="manual-row">
        <div class="manual-header">
           <div class="manual-name" title="${name}">${name} ${badge}</div>
           <div class="inline" style="gap:6px">
             <button class="btn-outline btn-sm" type="button" data-action="clear" data-key="${escapeHtml(key)}">Limpar</button>
             <input class="num" type="text" placeholder="#" value="${val}" data-key="${escapeHtml(key)}" />
           </div>
        </div>
        <div class="inline">
          <input type="color" class="custom-color-input" title="Cor (manual)" data-key="${escapeHtml(key)}" value="${customHex}">
          <small class="muted">Cor (manual)</small>
        </div>
      </div>
    `);
  }
  manualPanel.innerHTML = rows.join('');
}

/* ===== Labels ===== */
function updateLabel(key){
  const text=numberByKey.get(key); if(!text) return removeLabel(key);
  let pos = labelPositionByKey.get(key) || getCenterForKey(key);
  if(!pos) return;
  const { url,w,h } = makeSvgLabelIcon(text);
  const icon=L.icon({ iconUrl:url, iconSize:[w,h], iconAnchor:[w/2,h/2] });
  let marker=labelByKey.get(key);
  if(marker){ marker.setLatLng(pos); marker.setIcon(icon); }
  else {
    marker = L.marker(pos, { icon, draggable:true, zIndexOffset:1000 });
    marker.on('dragend', (e)=> labelPositionByKey.set(key, e.target.getLatLng()));
    if(labelsVisible) marker.addTo(map);
    labelByKey.set(key, marker);
  }
}

function removeLabel(key){ const m=labelByKey.get(key); if(m) map.removeLayer(m); labelByKey.delete(key); }

function getCenterForKey(key){
  const list=layersByKey.get(key)||[]; if(!list.length) return null;
  let bounds=null;
  for(const ly of list){
    let b=null; if(ly.getBounds) b=ly.getBounds();
    if(!b) continue; bounds = bounds ? bounds.extend(b) : b;
  }
  return bounds ? bounds.getCenter() : null;
}

function makeSvgLabelIcon(text){
  const t=String(text); const fs=numFontSize, pad=6;
  const w=Math.max(32, Math.round(0.65*fs)*t.length + pad*2); const h=fs + pad*2;
  const svg=`<svg xmlns="http://www.w3.org/2000/svg" width="${w}" height="${h}" viewBox="0 0 ${w} ${h}">
    <style>text{font:700 ${fs}px ${numFontFamily}; fill:#fff; stroke:#000; stroke-width:4px; paint-order:stroke;}</style>
    <text x="${w/2}" y="${h-pad}" text-anchor="middle">${escapeHtml(t)}</text>
  </svg>`;
  return { url:'data:image/svg+xml;utf8,'+encodeURIComponent(svg), w, h };
}

/* ===== Estilo ===== */
function defaultStyle(){ return { color:'#000', weight:borderWeight, opacity:1, lineJoin:'round', fillOpacity:.15, fillColor:'#ffffff' }; }
function setStyle(ly, s){
  const key = ly.feature?.properties?.__fileKey;
  const fill = s.fillColor || colorByKey.get(key) || '#ffffff';
  const final = { fillColor: fill, color: '#000', weight: borderWeight, opacity: 1, fillOpacity: fillOpacityPct/100 };
  if(ly.setStyle) ly.setStyle(final);
}
function repaintAll(){
  layersByKey.forEach((list)=>{
    list.forEach(ly => setStyle(ly, {}));
  });
}

/* ===== Auto paint ===== */
function buildKmlNameIndex(){
  const idx = new Map(); // normalizedName -> key
  for(const key of layersByKey.keys()){
    const base = displayName.get(key) || key;
    const a1 = matchKey(base);
    const a2 = matchKey(stripNumericPrefix(base));
    if(a1 && !idx.has(a1)) idx.set(a1, key);
    if(a2 && !idx.has(a2)) idx.set(a2, key);
  }
  return idx;
}

window.__applyAutoPaint = function(payloadJson){
  let payload;
  try{ payload = JSON.parse(payloadJson); }catch(e){ console.error(e); return; }

  if(layersByKey.size===0){
    el('autoPaintStatus').textContent = 'Carregue os KMLs antes de aplicar a pintura.';
    return;
  }

  const idx = buildKmlNameIndex();

  const seg = payload.segment;
  const minV = payload.min;
  const maxV = payload.max;
  const noOverwriteManual = !!payload.noOverwriteManual;

  const group = isGoodSegment(seg) ? 'good' : 'bad';
  const pal = window.__palettes[group];

  const missing = [];
  const skippedManual = [];
  let painted = 0;

  for(const item of payload.items){
    const excelName = item.name;
    const v = item.value;
    const cls = item.cls; // low|mid|high

    const n1 = matchKey(excelName);
    const n2 = matchKey(stripNumericPrefix(excelName));
    const key = idx.get(n1) || idx.get(n2);

    if(!key){
      missing.push(excelName);
      continue;
    }

    // ✅ Guarda dados para a tabela do PPT (não depende do tooltip)
    legendNameByKey.set(key, excelName);
    legendValueByKey.set(key, v);
    legendClsByKey.set(key, cls);

    const currentSource = paintSourceByKey.get(key);
    const shouldPaint = !(noOverwriteManual && currentSource === 'manual');

    const color = (cls==='low') ? pal.low : (cls==='high') ? pal.high : pal.mid;

    if(!shouldPaint){
      skippedManual.push(excelName);
    } else {
      applyColor(key, color, 'auto');
      painted++;
    }

    // Popup
    const labelCls = (cls==='low') ? 'Baixo' : (cls==='mid') ? 'Médio' : 'Alto';
    const html = `
      <div style="font-family:system-ui,Arial,sans-serif;min-width:240px">
        <div style="font-weight:900;margin-bottom:4px">${escapeHtml(excelName)}</div>
        <div><b>Segmento:</b> ${escapeHtml(seg)}</div>
        <div><b>Valor:</b> ${String(v).replace('.',',')}</div>
        <div><b>Categoria:</b> ${labelCls}</div>
        <div style="margin-top:6px;color:#666;font-size:12px">
          <b>Regras:</b> &lt;Min = Baixo | Min..Max = Médio | &gt;Max = Alto<br>
          <b>Min:</b> ${String(minV).replace('.',',')} | <b>Max:</b> ${String(maxV).replace('.',',')}
        </div>
      </div>
    `;
    (layersByKey.get(key)||[]).forEach(ly=>{
      ly.bindPopup(html);
    });
  }

  const st = [];
  st.push(`Pintadas: ${painted}`);
  if(skippedManual.length) st.push(`Ignoradas (manual): ${skippedManual.length}`);
  if(missing.length){
    const preview = missing.slice(0, 8).join(', ');
    st.push(`Não encontradas no KML: ${missing.length}${preview?` (ex.: ${preview})`:''}`);
  }
  el('autoPaintStatus').textContent = st.join(' | ');

  rebuildManualUI();
};

/* ===== Captura PNG (para Export e PPTX) ===== */
async function makeMapPngDataUrl(){
  const toRestore = [];
  if(labelsVisible){
    labelByKey.forEach(m=>{ if(map.hasLayer(m)){ toRestore.push(m); map.removeLayer(m);} });
  }
  await waitTilesLoaded(4000);

  return await new Promise((resolve, reject)=>{
    window.leafletImage(map, (err, canvas)=>{
      if(err){ toRestore.forEach(m=>m.addTo(map)); return reject(err); }
      const out = document.createElement('canvas'); out.width = canvas.width; out.height = canvas.height;
      const ctx = out.getContext('2d'); ctx.drawImage(canvas, 0, 0);

      // desenha números manuais por cima
      ctx.textAlign = 'center'; ctx.textBaseline = 'middle'; ctx.lineJoin = 'round';
      numberByKey.forEach((txt, key)=>{
        const pos = labelPositionByKey.get(key) || getCenterForKey(key); if(!pos) return;
        const pt = map.latLngToContainerPoint(pos);
        ctx.font = `700 ${numFontSize}px ${numFontFamily}`;
        ctx.strokeStyle = '#000'; ctx.lineWidth = Math.max(2, numFontSize*0.15);
        ctx.fillStyle = '#fff';
        ctx.strokeText(String(txt), pt.x, pt.y);
        ctx.fillText(String(txt), pt.x, pt.y);
      });

      const dataUrl = out.toDataURL('image/png');
      toRestore.forEach(m=>m.addTo(map));
      resolve(dataUrl);
    });
  });
}

/* ===== PPTX enqueue (um slide por captura) ===== */
window.__enqueuePptSlide = async function(){
  const caseKey = window.__ddCase?.getValue?.() || '';
  const segment = window.__ddSeg?.getValue?.() || '';
  if(!caseKey || !segment){
    el('pptQueueStatus').textContent = 'Selecione Caso e Segmento antes de adicionar.';
    return;
  }

  const stats = window.__lastStatsRaw;
  if(!stats || stats.caseKey !== caseKey || stats.segment !== segment){
    el('pptQueueStatus').textContent = 'Calcule as estatísticas (selecione o segmento) antes de adicionar.';
    return;
  }

  const group = isGoodSegment(segment) ? 'good' : 'bad';
  const pal = window.__palettes[group];

  // ✅ monta legenda estruturada (nome, cor, valor)
  const legend = [];
  legendValueByKey.forEach((val, key) => {
    if (typeof val !== 'number') return;
    const name = (numberByKey.get(key) || '').toString().trim() || (legendNameByKey.get(key) || displayName.get(key) || key);
    const color = colorByKey.get(key) || '#ffffff';
    legend.push({ name, value: val, color });
  });
  legend.sort((a,b) => String(a.name).localeCompare(String(b.name), 'pt-BR', { numeric:true, sensitivity:'base' }));

  const meta = {
    caseKey,
    segment,
    group,
    palette: { low: pal.low, mid: pal.mid, high: pal.high },
    stats: { min: stats.min, mean: stats.mean, max: stats.max },
    legend
  };

  if(!window.QtBridge?.addToPptQueue){
    el('pptQueueStatus').textContent = 'Bridge não disponível.';
    return;
  }

  el('pptQueueStatus').textContent = 'Capturando mapa…';
  try{
    const dataUrl = await makeMapPngDataUrl();
    window.QtBridge.addToPptQueue(JSON.stringify(meta), dataUrl);
    el('pptQueueStatus').textContent = 'Adicionado à fila.';
  }catch(e){
    console.error(e);
    el('pptQueueStatus').textContent = 'Erro ao capturar o mapa.';
  }
};

/* ===== Export ===== */
btnExport.addEventListener('click', async ()=>{
  try{
    btnExport.disabled = true; expStatus.textContent = 'A gerar PNG…';
    const dataUrl = await makeMapPngDataUrl();
    if(window.QtBridge?.savePng) window.QtBridge.savePng(dataUrl);
  }finally{ btnExport.disabled = false; expStatus.textContent = ''; }
});

function registerLayer(k, ly){ if(!layersByKey.has(k)) layersByKey.set(k, []); layersByKey.get(k).push(ly); }
function resetAll(){
  kmlGroup.clearLayers(); boundsGroup.clearLayers();
  layersByKey.clear(); displayName.clear(); colorByKey.clear(); paintSourceByKey.clear();
  numberByKey.clear(); labelPositionByKey.clear();
  labelByKey.forEach(m=>map.removeLayer(m)); labelByKey.clear();
  usedKeys.clear();

  // ✅ limpa dados da legenda do PPT
  legendNameByKey.clear();
  legendValueByKey.clear();
  legendClsByKey.clear();

  manualPanel.innerHTML='<span class="muted">Carregue KMLs para listar as malhas…</span>';
}
function canon(s){ return s.normalize('NFD').replace(/\p{Diacritic}/gu,'').toLowerCase().replace(/[^a-z0-9]+/g,' ').trim(); }
function uniqueKey(base){ let k=base, i=2; while(usedKeys.has(k)){ k=`${base} #${i++}` } usedKeys.add(k); return k; }
function escapeHtml(s){ return String(s).replace(/[&<>"']/g, m=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":"&#39;"}[m])); }

function waitTilesLoaded(timeout=4000){
  const layers = []; map.eachLayer(l=>{ if(l instanceof L.TileLayer && map.hasLayer(l)) layers.push(l); });
  if(!layers.length) return Promise.resolve();
  return new Promise(res=>{
    let done=false;
    const check=()=>{ if(done) return; if(layers.every(tl=>!tl._loading)) { done=true; res(); } };
    layers.forEach(tl=>{ if(tl._loading){ tl.once('load', ()=> check()); } });
    setTimeout(()=>{ if(!done) res(); }, timeout);
    check();
  });
}
</script>
</body>
</html>
"""



# =========================
# Excel parsing (com debug)
# =========================

SEGMENT_LABELS = [
    "Certamente voto",
    "Certamente não voto",
    "Aprovação",
    "Reprovação",
]

REQUIRED_SEGMENT_PAIRS = [
    {"Certamente voto", "Certamente não voto"},
    {"Aprovação", "Reprovação"},
]

MARKER_LABELS = {
    "municipio", "municipios",
    "bairro", "bairros",
    "distrito", "distritos",
    "subprefeitura", "subprefeituras",
    "zona", "zonas",
    "mesorregiao", "mesorregioes",
    "microrregiao", "microrregioes",
    "regiao administrativa", "regioes administrativas",
    "regiao", "regioes",
    "area de planejamento", "areas de planejamento",
    "ap", "ap areas de planejamento",
    "rpa", "za",
    "zona administrativa", "zonas administrativas",
    "zona administrativa da capital", "zonas administrativas da capital",
    "uf", "estado", "estados", "capital", "capitais",
}


IGNORE_HEADER_WORDS = {"total", "base", "geral"}


def normalize_text(value: Any) -> str:
    text = "" if value is None else str(value)
    text = text.strip()
    text = unicodedata.normalize("NFD", text)
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    text = text.lower()
    text = " ".join(text.split())
    return text


def is_marker_text(value: Any) -> bool:
    """
    Reconhece apenas rótulos territoriais usados como marcador de tabela.
    Evita confundir palavras comuns dentro de títulos e perguntas.
    """
    normalized = normalize_text(value)
    if not normalized:
        return False

    cleaned = re.sub(r"[^a-z0-9 ]+", " ", normalized)
    cleaned = re.sub(r"\s+", " ", cleaned).strip()

    if cleaned in MARKER_LABELS:
        return True

    for label in MARKER_LABELS:
        if cleaned.startswith(label + " "):
            remainder = cleaned[len(label):].strip()
            if remainder in {"", "total", "percentual", "porcentagem", "segmento"}:
                return True

    return False


def parse_number(cell) -> Optional[float]:
    raw = cell.value
    if raw is None:
        return None

    num_format = (cell.number_format or "").lower()

    if isinstance(raw, (int, float)):
        value = float(raw)
        if "%" in num_format and 0 <= value <= 1.0:
            return value * 100.0
        return value

    text = str(raw).strip()
    if not text:
        return None

    has_percent = "%" in text or "%" in num_format
    text = text.replace("%", "").strip()

    if "," not in text and re.match(r"^\d{1,3}(\.\d{3})+$", text):
        text = text.replace(".", "")

    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    else:
        text = text.replace(",", ".")

    try:
        value = float(text)
    except ValueError:
        return None

    if has_percent and 0 <= value <= 1.0:
        value *= 100.0

    return value


def get_row_label_text(sheet, row_index: int, max_label_cols: int = 3) -> Optional[str]:
    for col_index in range(1, max_label_cols + 1):
        v = sheet.cell(row_index, col_index).value
        if isinstance(v, str) and v.strip():
            return v.strip()
    return None


def row_has_segment_label(sheet, row_index: int) -> Optional[str]:
    label = get_row_label_text(sheet, row_index, max_label_cols=3)
    if not label:
        return None
    label_norm = normalize_text(label)
    for seg in SEGMENT_LABELS:
        if normalize_text(seg) in label_norm:
            return seg
    return None


def row_is_base(sheet, row_index: int) -> bool:
    label = get_row_label_text(sheet, row_index, max_label_cols=3)
    if not label:
        return False
    return normalize_text(label).startswith("base")


def row_contains_marker(sheet, row_index: int, max_check_cols: int = 10) -> bool:
    for col_index in range(1, max_check_cols + 1):
        v = sheet.cell(row_index, col_index).value
        if isinstance(v, str) and v.strip() and is_marker_text(v):
            return True
    return False


def find_segment_rows_below_marker(sheet, marker_row: int, max_row: int, scan_limit: int = 500) -> List[int]:
    segment_rows: List[int] = []
    end_row = min(max_row, marker_row + scan_limit)
    for row_index in range(marker_row, end_row + 1):
        if row_has_segment_label(sheet, row_index):
            segment_rows.append(row_index)
    return segment_rows


def pick_header_row_from_segments(sheet, first_segment_row: int, max_col: int, lookback: int = 10) -> Optional[int]:
    if first_segment_row <= 1:
        return None

    candidates: List[Tuple[int, int]] = []
    start_row = max(1, first_segment_row - lookback)
    end_row = first_segment_row - 1

    for row_index in range(start_row, end_row + 1):
        if row_has_segment_label(sheet, row_index) or row_is_base(sheet, row_index):
            continue

        text_cells = 0
        numeric_cells = 0
        nonempty_cells = 0

        for col_index in range(2, max_col + 1):
            v = sheet.cell(row_index, col_index).value
            if v in (None, ""):
                continue
            nonempty_cells += 1
            if isinstance(v, str) and v.strip():
                text_cells += 1
            elif isinstance(v, (int, float)):
                numeric_cells += 1

        if text_cells >= 2:
            score = text_cells * 10 - numeric_cells * 20 + nonempty_cells
            candidates.append((score, row_index))

    if not candidates:
        return end_row if end_row >= 1 else None

    candidates.sort(reverse=True)
    return candidates[0][1]


def compute_last_table_col(sheet, header_row: int, first_segment_row: int, max_col: int) -> int:
    last_col = 1
    for col_index in range(1, max_col + 1):
        if sheet.cell(header_row, col_index).value not in (None, ""):
            last_col = col_index
        if sheet.cell(first_segment_row, col_index).value not in (None, ""):
            last_col = max(last_col, col_index)
    return last_col


def find_total_column(sheet, header_row: int, last_table_col: int) -> Tuple[int, bool]:
    """
    Procura a coluna Total em várias linhas próximas ao cabeçalho.
    Não assume automaticamente que a coluna B é o Total.
    """
    scan_start = max(1, header_row - 8)
    scan_end = min(sheet.max_row or header_row, header_row + 2)

    for scan_row in range(scan_end, scan_start - 1, -1):
        for col_index in range(1, last_table_col + 1):
            value = sheet.cell(scan_row, col_index).value
            if value is None:
                continue
            if normalize_text(value) in {"total", "geral", "total geral"}:
                return col_index, True

    for merged in sheet.merged_cells.ranges:
        if merged.max_row < scan_start or merged.min_row > scan_end:
            continue
        value = sheet.cell(merged.min_row, merged.min_col).value
        if value is not None and normalize_text(value) in {"total", "geral", "total geral"}:
            return merged.min_col, True

    return 0, False



def extract_base_total_from_row(
    sheet,
    base_row: int,
    total_col: int,
    last_table_col: int,
) -> Tuple[Optional[float], Optional[int]]:
    """
    Lê a base total. Primeiro usa a coluna Total detectada. Se ela não existir,
    usa o maior valor numérico da linha Base, que normalmente é a base geral.
    """
    if total_col and total_col > 0:
        value = parse_number(sheet.cell(base_row, total_col))
        if isinstance(value, (int, float)) and value > 0:
            return float(value), total_col

    candidates: List[Tuple[float, int]] = []
    for col_index in range(2, last_table_col + 1):
        value = parse_number(sheet.cell(base_row, col_index))
        if isinstance(value, (int, float)) and value > 0:
            candidates.append((float(value), col_index))

    if not candidates:
        return None, None

    candidates.sort(key=lambda item: item[0], reverse=True)
    return candidates[0]

def extract_region_columns(sheet, header_row: int, last_table_col: int) -> Tuple[List[str], List[int], List[str]]:
    region_names: List[str] = []
    region_cols: List[int] = []
    raw_header_labels: List[str] = []

    for col_index in range(2, last_table_col + 1):
        header_value = sheet.cell(header_row, col_index).value
        if header_value is None:
            continue
        header_text = str(header_value).strip()
        if not header_text:
            continue

        raw_header_labels.append(header_text)
        header_norm = normalize_text(header_text)

        if header_norm in IGNORE_HEADER_WORDS:
            continue
        if header_norm == "total":
            continue
        if is_marker_text(header_text) and not header_norm.startswith("ap "):
            continue

        region_names.append(header_text)
        region_cols.append(col_index)

    return region_names, region_cols, raw_header_labels


def find_title_above(sheet, header_row: int, first_col: int, last_col: int) -> Optional[str]:
    merged_ranges = list(sheet.merged_cells.ranges)

    for scan_row in range(max(1, header_row - 30), header_row)[::-1]:
        for rng in merged_ranges:
            if rng.min_row == scan_row and rng.max_row == scan_row and rng.min_col <= first_col and rng.max_col >= last_col:
                v = sheet.cell(rng.min_row, rng.min_col).value
                if v is not None and str(v).strip():
                    return str(v).strip()

        a_val = sheet.cell(scan_row, 1).value
        if a_val is not None and str(a_val).strip():
            empty_count = 0
            width = last_col - first_col + 1
            for col_index in range(first_col, last_col + 1):
                if sheet.cell(scan_row, col_index).value in (None, ""):
                    empty_count += 1
            if empty_count >= int(0.85 * width):
                return str(a_val).strip()

    return None


def detect_table_kind(marker_text: str, header_labels: List[str]) -> str:
    """
    Identifica automaticamente o tipo de recorte territorial.
    """
    nm = normalize_text(marker_text)
    header_norms = [normalize_text(h) for h in header_labels]

    if (
        "ap (areas de planejamento)" in nm
        or "area de planejamento" in nm
        or "areas de planejamento" in nm
        or nm == "ap"
        or any(h == "ap" or h.startswith("ap ") for h in header_norms)
    ):
        return "Áreas de Planejamento"

    if nm == "rpa" or nm.startswith("rpa ") or any(h == "rpa" or h.startswith("rpa ") for h in header_norms):
        return "RPA"

    if (
        nm == "za"
        or "zona administrativa" in nm
        or "zonas administrativas" in nm
        or any(h == "za" or h.startswith("za ") for h in header_norms)
    ):
        return "ZA"

    if "regiao administrativa" in nm or "regioes administrativas" in nm:
        return "Regiões Administrativas"

    if "mesorregiao" in nm or "mesorregioes" in nm:
        return "Mesorregiões"

    if "microrregiao" in nm or "microrregioes" in nm:
        return "Microrregiões"

    if "municipio" in nm or "municipios" in nm:
        return "Municípios"

    if "bairro" in nm or "bairros" in nm:
        return "Bairros"

    if "distrito" in nm or "distritos" in nm:
        return "Distritos"

    if "subprefeitura" in nm or "subprefeituras" in nm:
        return "Subprefeituras"

    if nm == "zona" or "zonas" in nm:
        return "Zonas"

    if nm == "uf" or nm.startswith("uf ") or any(h == "uf" for h in header_norms):
        return "UF"

    if "estado" in nm or "estados" in nm:
        return "Estados"

    if "capital" in nm or "capitais" in nm:
        return "Capitais"

    if "regiao" in nm or "regioes" in nm:
        return "Regiões"

    return "Tabela"


def parse_excel_cases_with_debug(excel_path: str) -> Tuple[Dict[str, Dict[str, Any]], Dict[str, Any]]:
    workbook = openpyxl.load_workbook(excel_path, data_only=True)

    cases_by_key: Dict[str, Dict[str, Any]] = {}
    debug_candidates: List[Dict[str, Any]] = []
    seen_tables: Set[Tuple[str, int, int]] = set()

    for sheet in workbook.worksheets:
        max_row = min(sheet.max_row or 0, 20000)
        max_col = min(sheet.max_column or 0, 300)
        if max_row == 0 or max_col == 0:
            continue

        for marker_row in range(1, max_row + 1):
            # Marcadores territoriais costumam ficar nas primeiras colunas.
            for marker_col in range(1, min(max_col, 12) + 1):
                marker_value = sheet.cell(marker_row, marker_col).value
                if not isinstance(marker_value, str):
                    continue
                if not is_marker_text(marker_value):
                    continue

                marker_text = str(marker_value).strip()

                segment_rows = find_segment_rows_below_marker(sheet, marker_row, max_row, scan_limit=500)
                if not segment_rows:
                    continue

                first_segment_row = min(segment_rows)

                header_row = pick_header_row_from_segments(sheet, first_segment_row, max_col, lookback=10)
                if header_row is None:
                    continue

                last_table_col = compute_last_table_col(sheet, header_row, first_segment_row, max_col)
                table_dedupe_key = (sheet.title, header_row, last_table_col)
                if table_dedupe_key in seen_tables:
                    continue
                seen_tables.add(table_dedupe_key)

                total_col, _found_total = find_total_column(sheet, header_row, last_table_col)
                region_names, region_cols, raw_header_labels = extract_region_columns(sheet, header_row, last_table_col)
                if len(region_names) < 2:
                    continue

                table_kind = detect_table_kind(marker_text, raw_header_labels)

                base_total: Optional[float] = None
                base_row: Optional[int] = None

                segment_values_by_label: Dict[str, Dict[str, float]] = {}
                segment_sources_by_label: Dict[str, Dict[str, Dict[str, Any]]] = {}
                segment_row_by_label: Dict[str, int] = {}

                blank_run = 0
                scan_start = header_row + 1
                scan_end = min(max_row, header_row + 1200)

                for data_row in range(scan_start, scan_end + 1):
                    row_is_empty = True
                    for col_index in range(1, min(last_table_col, 20) + 1):
                        if sheet.cell(data_row, col_index).value not in (None, ""):
                            row_is_empty = False
                            break
                    if row_is_empty:
                        blank_run += 1
                    else:
                        blank_run = 0
                    if blank_run >= 2 and data_row > first_segment_row:
                        break

                    if data_row > first_segment_row and row_contains_marker(sheet, data_row, max_check_cols=12):
                        break

                    label_text = get_row_label_text(sheet, data_row, max_label_cols=3)
                    if not label_text:
                        continue
                    label_norm = normalize_text(label_text)

                    if label_norm.startswith("base"):
                        if base_total is None:
                            detected_base, detected_col = extract_base_total_from_row(
                                sheet,
                                data_row,
                                total_col,
                                last_table_col,
                            )
                            base_total = detected_base
                            base_row = data_row
                            if detected_col:
                                total_col = detected_col
                        continue

                    matched_segment = row_has_segment_label(sheet, data_row)
                    if not matched_segment:
                        continue

                    if matched_segment in segment_row_by_label:
                        continue

                    segment_row_by_label[matched_segment] = data_row
                    segment_values_by_label.setdefault(matched_segment, {})
                    segment_sources_by_label.setdefault(matched_segment, {})

                    for region_name, region_col in zip(region_names, region_cols):
                        cell = sheet.cell(data_row, region_col)
                        parsed = parse_number(cell)
                        if parsed is None:
                            continue
                        addr = f"{get_column_letter(region_col)}{data_row}"
                        segment_values_by_label[matched_segment][region_name] = float(parsed)
                        segment_sources_by_label[matched_segment][region_name] = {
                            "addr": addr,
                            "row": data_row,
                            "col": region_col,
                            "raw": cell.value,
                            "num_format": cell.number_format,
                            "parsed": float(parsed),
                        }

                if base_total is None or not isinstance(base_total, (int, float)) or base_total <= 0:
                    continue

                segment_keys = set(segment_values_by_label.keys())
                ok = any(pair.issubset(segment_keys) for pair in REQUIRED_SEGMENT_PAIRS)
                if not ok:
                    continue

                min_required_fill = max(3, int(0.70 * len(region_names)))
                if any(len(segment_values_by_label.get(seg, {})) < min_required_fill for seg in segment_keys):
                    pass

                title = find_title_above(sheet, header_row, 1, last_table_col) or "(Sem título)"
                case_key = (
                    f"{title} — {table_kind} — Aba: {sheet.title} — "
                    f"Header: {header_row} — Base: {base_row if base_row else '—'}"
                )
                if case_key in cases_by_key:
                    case_key = f"{case_key} — marcador({marker_row},{marker_col})"

                cases_by_key[case_key] = {
                    "__meta": {
                        "title": title,
                        "kind": table_kind,
                        "sheet": sheet.title,
                        "marker_row": marker_row,
                        "marker_col": marker_col,
                        "marker_text": marker_text,
                        "first_segment_row": first_segment_row,
                        "header_row": header_row,
                        "last_table_col": last_table_col,
                        "total_col": total_col,
                        "region_names": region_names,
                        "region_cols": region_cols,
                        "base_row": base_row,
                        "base_total": float(base_total),
                        "segment_rows": segment_row_by_label,
                        "segment_sources": segment_sources_by_label,
                    }
                }

                for seg_label, region_map in segment_values_by_label.items():
                    cases_by_key[case_key][seg_label] = region_map

    debug_report = {"excel_path": excel_path, "cases_count": len(cases_by_key)}
    return cases_by_key, debug_report


# =========================
# PPT helpers
# =========================

def dec_round_half_up(value: float, decimals: int = 1) -> Decimal:
    """
    Arredonda em Decimal com ROUND_HALF_UP (0,55 -> 0,6).
    """
    q = Decimal("1").scaleb(-decimals)  # decimals=1 -> 0.1
    return Decimal(str(value)).quantize(q, rounding=ROUND_HALF_UP)

def round_half_up(value: float, decimals: int = 1) -> float:
    return float(dec_round_half_up(value, decimals))

def fmt_1_half_up(value: float) -> str:
    """
    Retorna string com 1 casa decimal e vírgula, usando HALF_UP.
    """
    d = dec_round_half_up(value, 1)
    return str(d).replace(".", ",")


def _hex_to_rgb(hex_color: str) -> RGBColor:
    s = (hex_color or "").strip()
    if s.startswith("#"):
        s = s[1:]
    if len(s) != 6:
        return RGBColor(0, 0, 0)
    r = int(s[0:2], 16)
    g = int(s[2:4], 16)
    b = int(s[4:6], 16)
    return RGBColor(r, g, b)


def _fmt_1(x: float) -> str:
    return fmt_1_half_up(float(x))


def _is_good_segment(seg: str) -> bool:
    t = (seg or "").strip()
    t = unicodedata.normalize("NFD", t)
    t = "".join(ch for ch in t if unicodedata.category(ch) != "Mn")
    t = t.lower()
    return ("certamente voto" in t) or ("aprovacao" in t) or ("aprovação" in t)


def build_pptx(
    out_path: str,
    slides: List[Dict[str, Any]],
    logo_bytes: Optional[bytes]
) -> None:
    """
    Gera os slides em 16:9 mantendo o mapa e o cabeçalho como estavam.

    Alterações visuais desta versão:
    - tabela da direita com apenas 2 colunas: Regiões | %;
    - a própria célula do percentual recebe a cor da região;
    - a antiga legenda horizontal abaixo do mapa foi removida;
    - Mín / Entre / Máx agora aparecem em uma pequena tabela abaixo
      da tabela de regiões, com as cores aplicadas nos rótulos.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank = prs.slide_layouts[6]

    def text_color_for_fill(hex_color: str) -> RGBColor:
        """
        Retorna preto ou branco conforme a luminosidade da cor de fundo.
        """
        s = (hex_color or "").strip().lstrip("#")
        if len(s) != 6:
            return RGBColor(0, 0, 0)

        try:
            r = int(s[0:2], 16)
            g = int(s[2:4], 16)
            b = int(s[4:6], 16)
        except ValueError:
            return RGBColor(0, 0, 0)

        luminance = (0.299 * r + 0.587 * g + 0.114 * b) / 255.0
        return RGBColor(0, 0, 0) if luminance >= 0.56 else RGBColor(255, 255, 255)

    def set_cell_border(cell, color: str = "BFBFBF", width: int = 12700) -> None:
        """
        Aplica borda visível nos quatro lados da célula.
        width=12700 equivale aproximadamente a 1 pt.
        """
        from pptx.oxml.xmlchemy import OxmlElement
        from pptx.oxml.ns import qn

        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()

        for edge in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
            ln = tcPr.find(qn(edge))
            if ln is None:
                ln = OxmlElement(edge)
                tcPr.append(ln)

            ln.set("w", str(width))
            ln.set("cap", "flat")
            ln.set("cmpd", "sng")
            ln.set("algn", "ctr")

            solidFill = ln.find(qn("a:solidFill"))
            if solidFill is None:
                solidFill = OxmlElement("a:solidFill")
                ln.append(solidFill)

            srgbClr = solidFill.find(qn("a:srgbClr"))
            if srgbClr is None:
                srgbClr = OxmlElement("a:srgbClr")
                solidFill.append(srgbClr)
            srgbClr.set("val", color)

            prstDash = ln.find(qn("a:prstDash"))
            if prstDash is None:
                prstDash = OxmlElement("a:prstDash")
                ln.append(prstDash)
            prstDash.set("val", "solid")

    for item in slides:
        seg = item["segment"]
        case_title = (item.get("case_title") or "").strip()
        title_text = f"{case_title} - {seg}".strip(" -") if case_title else seg

        pal = item["palette"]
        stats = item["stats"]
        png_bytes = item["png_bytes"]

        region_rows = item.get("legend_rows") or []
        has_table = bool(region_rows)

        slide = prs.slides.add_slide(blank)

        # =====================================================
        # CABEÇALHO: mantido
        # =====================================================
        slide_left = Inches(0.38)
        slide_right = Inches(0.38)

        logo_reserved_w = Inches(2.05)
        logo_gap = Inches(0.18)

        title_left = slide_left
        title_top = Inches(0.08)
        title_w = prs.slide_width - slide_left - slide_right - logo_reserved_w - logo_gap
        title_h = Inches(0.62)

        title_box = slide.shapes.add_textbox(
            title_left,
            title_top,
            title_w,
            title_h
        )
        tf = title_box.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.margin_left = 0
        tf.margin_right = 0
        tf.margin_top = 0
        tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.text = title_text
        p.font.bold = True
        p.font.name = "DIN"

        if len(title_text) <= 75:
            p.font.size = Pt(20)
        elif len(title_text) <= 115:
            p.font.size = Pt(18)
        else:
            p.font.size = Pt(16)

        if logo_bytes:
            stream = io.BytesIO(logo_bytes)
            logo_pic = slide.shapes.add_picture(
                stream,
                left=prs.slide_width - slide_right - logo_reserved_w,
                top=Inches(0.10),
                height=Inches(0.42)
            )
            logo_pic.left = prs.slide_width - slide_right - logo_pic.width

        # =====================================================
        # ÁREA PRINCIPAL: dimensões gerais mantidas
        # =====================================================
        content_left = Inches(0.72)
        content_right = Inches(0.62)
        content_top = Inches(0.76)

        map_h = Inches(5.28)

        table_gap = Inches(0.32)

        # Mantemos a mesma largura total da tabela antiga:
        # 1.55 + 0.34 + 0.66 = 2.55 pol.
        # Agora ela é dividida em apenas 2 colunas.
        name_w = Inches(1.55)
        pct_w = Inches(1.00)
        table_w = name_w + pct_w if has_table else Inches(0)

        total_content_w = prs.slide_width - content_left - content_right
        map_box_w = (
            total_content_w - table_gap - table_w
            if has_table
            else total_content_w
        )

        table_left = content_left + map_box_w + table_gap

        # =====================================================
        # MAPA: mantido exatamente com lógica "contain"
        # =====================================================
        img = Image.open(io.BytesIO(png_bytes))
        iw, ih = img.size
        img_ratio = iw / ih if ih else 1.0
        box_ratio = float(map_box_w) / float(map_h)

        if img_ratio >= box_ratio:
            pic_w = int(map_box_w)
            pic_h = int(float(map_box_w) / img_ratio)
        else:
            pic_h = int(map_h)
            pic_w = int(float(map_h) * img_ratio)

        pic_left = int(content_left) + int((float(map_box_w) - pic_w) / 2)
        pic_top = int(content_top) + int((float(map_h) - pic_h) / 2)

        map_stream = io.BytesIO(png_bytes)
        slide.shapes.add_picture(
            map_stream,
            pic_left,
            pic_top,
            width=pic_w,
            height=pic_h
        )

        # =====================================================
        # TABELA À DIREITA: Regiões | %
        # A cor agora fica na MESMA célula do percentual.
        # =====================================================
        if has_table:
            def natural_key(value: str):
                parts = re.split(r"(\d+)", str(value).lower())
                return [
                    int(part) if part.isdigit() else part
                    for part in parts
                ]

            rows_sorted = sorted(
                region_rows,
                key=lambda row: natural_key(row.get("name", ""))
            )

            max_rows = 18
            trimmed = rows_sorted[:max_rows]

            rows = len(trimmed) + 1
            cols = 2

            # Reserva espaço inferior para a tabela Mín / Entre / Máx.
            limits_gap = Inches(0.16)
            limits_h = Inches(1.42)
            max_region_table_h = map_h - limits_gap - limits_h

            if len(trimmed) <= 8:
                desired_row_h = Inches(0.52)
            elif len(trimmed) <= 12:
                desired_row_h = Inches(0.42)
            elif len(trimmed) <= 16:
                desired_row_h = Inches(0.34)
            else:
                desired_row_h = Inches(0.29)

            desired_table_h = desired_row_h * rows
            region_table_h = min(desired_table_h, max_region_table_h)

            tbl_shape = slide.shapes.add_table(
                rows,
                cols,
                table_left,
                content_top,
                table_w,
                region_table_h
            )
            tbl = tbl_shape.table

            tbl.columns[0].width = name_w
            tbl.columns[1].width = pct_w

            row_h = int(region_table_h / rows) if rows else int(region_table_h)
            for row_index in range(rows):
                tbl.rows[row_index].height = row_h

            headers = ["Regiões", "%"]
            for col_index, header in enumerate(headers):
                cell = tbl.cell(0, col_index)
                cell.text = header
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(245, 245, 245)

                paragraph = cell.text_frame.paragraphs[0]
                paragraph.font.bold = True
                paragraph.font.size = Pt(12)
                paragraph.font.name = "DIN"
                paragraph.font.color.rgb = RGBColor(0, 0, 0)
                paragraph.alignment = PP_ALIGN.CENTER
                set_cell_border(cell)

            font_size = 12

            for row_index, row in enumerate(trimmed, start=1):
                name = str(row.get("name", "")).strip()
                value = float(row.get("value", 0.0))
                color = str(row.get("color", "#ffffff")).strip() or "#ffffff"

                # Coluna da região
                cell_name = tbl.cell(row_index, 0)
                cell_name.text = name
                cell_name.fill.solid()
                cell_name.fill.fore_color.rgb = RGBColor(255, 255, 255)

                p_name = cell_name.text_frame.paragraphs[0]
                p_name.font.size = Pt(font_size)
                p_name.font.bold = True
                p_name.font.name = "DIN"
                p_name.font.color.rgb = RGBColor(0, 0, 0)
                p_name.alignment = PP_ALIGN.CENTER
                set_cell_border(cell_name)

                # Coluna do percentual + cor
                cell_pct = tbl.cell(row_index, 1)
                cell_pct.text = f"{_fmt_1(value)}%"
                cell_pct.fill.solid()
                cell_pct.fill.fore_color.rgb = _hex_to_rgb(color)

                p_pct = cell_pct.text_frame.paragraphs[0]
                p_pct.font.size = Pt(font_size)
                p_pct.font.bold = True
                p_pct.font.name = "DIN"
                p_pct.font.color.rgb = text_color_for_fill(color)
                p_pct.alignment = PP_ALIGN.CENTER
                set_cell_border(cell_pct)

            # =================================================
            # TABELA MÍN / ENTRE / MÁX
            # Fica abaixo da tabela de regiões, como na referência.
            # =================================================
            limits_top = content_top + region_table_h + limits_gap

            limits_rows = 4
            limits_cols = 2

            label_w = Inches(0.92)
            value_w = table_w - label_w

            limits_shape = slide.shapes.add_table(
                limits_rows,
                limits_cols,
                table_left,
                limits_top,
                table_w,
                limits_h
            )
            limits_tbl = limits_shape.table

            limits_tbl.columns[0].width = label_w
            limits_tbl.columns[1].width = value_w

            header_h = Inches(0.38)
            data_h = int((limits_h - header_h) / 3)

            limits_tbl.rows[0].height = int(header_h)
            for i in range(1, 4):
                limits_tbl.rows[i].height = data_h

            # Cabeçalho simples, sem inventar período do relatório.
            c00 = limits_tbl.cell(0, 0)
            c00.text = ""
            c00.fill.solid()
            c00.fill.fore_color.rgb = RGBColor(255, 255, 255)
            set_cell_border(c00)

            c01 = limits_tbl.cell(0, 1)
            c01.text = "Limites"
            c01.fill.solid()
            c01.fill.fore_color.rgb = RGBColor(255, 255, 255)

            p_header = c01.text_frame.paragraphs[0]
            p_header.font.name = "DIN"
            p_header.font.size = Pt(12)
            p_header.font.bold = True
            p_header.font.color.rgb = RGBColor(0, 0, 0)
            p_header.alignment = PP_ALIGN.CENTER
            set_cell_border(c01)

            min_v = float(stats["min"])
            max_v = float(stats["max"])

            limits_data = [
                ("Mín.:", f"{_fmt_1(min_v)}%", pal["low"]),
                (
                    "Entre:",
                    f"{_fmt_1(min_v)}% a {_fmt_1(max_v)}%",
                    pal["mid"]
                ),
                ("Máx.:", f"{_fmt_1(max_v)}%", pal["high"]),
            ]

            for idx, (label, value_text, color_hex) in enumerate(limits_data, start=1):
                label_cell = limits_tbl.cell(idx, 0)
                label_cell.text = label
                label_cell.fill.solid()
                label_cell.fill.fore_color.rgb = _hex_to_rgb(color_hex)

                p_label = label_cell.text_frame.paragraphs[0]
                p_label.font.name = "DIN"
                p_label.font.size = Pt(12)
                p_label.font.bold = True
                p_label.font.color.rgb = text_color_for_fill(color_hex)
                p_label.alignment = PP_ALIGN.CENTER
                set_cell_border(label_cell)

                value_cell = limits_tbl.cell(idx, 1)
                value_cell.text = value_text
                value_cell.fill.solid()
                value_cell.fill.fore_color.rgb = RGBColor(255, 255, 255)

                p_value = value_cell.text_frame.paragraphs[0]
                p_value.font.name = "DIN"
                p_value.font.size = Pt(12)
                p_value.font.bold = True
                p_value.font.color.rgb = RGBColor(0, 0, 0)
                p_value.alignment = PP_ALIGN.CENTER
                set_cell_border(value_cell)

        # A antiga legenda horizontal abaixo do mapa foi removida de propósito.

    prs.save(out_path)


# =========================
# Qt Bridge
# =========================

class QtBridge(QObject):
    def __init__(self, main_window: QMainWindow, web_view: QWebEngineView):
        super().__init__()
        self.main_window = main_window
        self.web_view = web_view
        self.cases_by_key: Dict[str, Dict[str, Any]] = {}

        self.logo_bytes: Optional[bytes] = None
        self.ppt_queue: List[Dict[str, Any]] = []

    def _push_ppt_state(self):
        items = []
        for i, it in enumerate(self.ppt_queue):
            # label que aparece na lista
            label = f"{it.get('case_title','')} - {it.get('segment','')}".strip(" -")
            items.append({"index": i, "label": label})

        payload = {
            "count": len(self.ppt_queue),
            "items": items,
            "logoLoaded": bool(self.logo_bytes),
        }
        self.web_view.page().runJavaScript(
            f"window.__setPptQueueState({json.dumps(json.dumps(payload, ensure_ascii=False))});"
        )

    @Slot(int)
    def removePptItem(self, index: int):
        try:
            if 0 <= index < len(self.ppt_queue):
                del self.ppt_queue[index]
            self._push_ppt_state()
        except Exception as e:
            QMessageBox.critical(self.main_window, "Erro", f"Falha ao remover item da fila.\n\n{e}")

    @Slot()
    def pickExcel(self):
        excel_path, _ = QFileDialog.getOpenFileName(
            self.main_window,
            "Selecionar arquivo Excel",
            "",
            "Excel (*.xlsx);;Todos (*.*)"
        )
        if not excel_path:
            return

        try:
            self.cases_by_key, debug_report = parse_excel_cases_with_debug(excel_path)

            debug_path = os.path.splitext(excel_path)[0] + ".processamentos_relatoria_debug.json"
            try:
                with open(debug_path, "w", encoding="utf-8") as f:
                    json.dump(debug_report, f, ensure_ascii=False, indent=2)
            except Exception:
                debug_path = ""

            segments_by_case: Dict[str, List[str]] = {}
            base_by_case: Dict[str, str] = {}

            for case_key, case_data in self.cases_by_key.items():
                segments_by_case[case_key] = [k for k in case_data.keys() if not k.startswith("__")]
                meta = case_data.get("__meta", {}) if isinstance(case_data, dict) else {}
                base_total = meta.get("base_total", None)
                base_by_case[case_key] = (f"{float(base_total):.0f}" if isinstance(base_total, (int, float)) else "")

            payload = {
                "cases": list(self.cases_by_key.keys()),
                "segmentsByCase": segments_by_case,
                "baseByCase": base_by_case,
                "debugPath": debug_path,
            }
            self.web_view.page().runJavaScript(f"window.__setCases({json.dumps(json.dumps(payload, ensure_ascii=False))});")

        except Exception as e:
            QMessageBox.critical(self.main_window, "Erro ao ler Excel", f"Não consegui analisar o Excel.\n\nDetalhes: {e}")

    @Slot(str, str, str)
    def computeStats(self, case_key: str, segment_label: str, base_override_text: str):
        try:
            case_data = self.cases_by_key.get(case_key, {})
            meta = case_data.get("__meta", {}) if isinstance(case_data, dict) else {}

            excel_base_total = meta.get("base_total", None)

            manual_base: Optional[float] = None
            txt = (base_override_text or "").strip()
            if txt:
                try:
                    manual_base = float(txt.replace(".", "").replace(",", "."))
                except ValueError:
                    manual_base = None

            base_total_used = manual_base if (manual_base is not None and manual_base > 0) else (
                float(excel_base_total) if isinstance(excel_base_total, (int, float)) and excel_base_total > 0 else None
            )
            base_source = "manual" if (manual_base is not None and manual_base > 0) else "excel"

            region_map = case_data.get(segment_label, {}) if isinstance(case_data, dict) else {}
            items = [(name, val) for name, val in region_map.items() if isinstance(val, (int, float))]
            if not items:
                self.web_view.page().runJavaScript("window.__setStats('—','—','—','','','',null,null,null);")
                return

            values = [v for (_, v) in items]
            count = len(values)
            total_sum = sum(values)
            mean = total_sum / count

            non_occ = 100.0 - mean
            margin_error: Optional[float] = None
            if base_total_used is not None and base_total_used > 0:
                margin_error = 1.96 * math.sqrt((mean * non_occ) / float(base_total_used))

            min_value = (mean - margin_error) if margin_error is not None else None
            max_value = (mean + margin_error) if margin_error is not None else None

            def fmt_stat(x: Optional[float]) -> str:
                if x is None:
                    return "—"
                return fmt_1_half_up(float(x))

            def fmt_dbg(x: Optional[float]) -> str:
                if x is None:
                    return "—"
                return f"{x:.2f}".replace(".", ",")

            segment_sources = (meta.get("segment_sources", {}) or {}).get(segment_label, {}) or {}
            used_lines: List[str] = []
            for region_name, value in items:
                src = segment_sources.get(region_name, {}) or {}
                used_lines.append(
                    f"{region_name}={fmt_dbg(value)} | cell={src.get('addr')} | raw={src.get('raw')} | fmt={src.get('num_format')} | parsed={fmt_dbg(src.get('parsed')) if isinstance(src.get('parsed'), (int,float)) else src.get('parsed')}"
                )
            values_str = "\n".join(used_lines[:300])

            location = (
                f"Aba: {meta.get('sheet')}\n"
                f"Marcador: ({meta.get('marker_row')},{meta.get('marker_col')}) '{meta.get('marker_text')}'\n"
                f"Header: {meta.get('header_row')} | 1º segmento: {meta.get('first_segment_row')} | Base: {meta.get('base_row')} | Total col: {meta.get('total_col')}\n"
            )

            tip_mean = (
                f"{location}\n"
                f"Segmento: {segment_label}\n"
                f"Obs.: Total NÃO entra na média.\n\n"
                f"Soma = {fmt_dbg(total_sum)}\n"
                f"n = {count}\n"
                f"Média = Soma/n = {fmt_dbg(total_sum)} / {count} = {fmt_dbg(mean)}\n\n"
                f"Células/valores usados:\n{values_str}"
            )

            if margin_error is None:
                tip_min = tip_mean + "\n\nBase inválida/não definida -> Erro/Mín/Máx não calculados."
                tip_max = tip_min
            else:
                tip_common = (
                    f"{tip_mean}\n\n"
                    f"Base usada ({base_source}) = {fmt_dbg(float(base_total_used))}\n"
                    f"Não ocorrência = 100 - Média = 100 - {fmt_dbg(mean)} = {fmt_dbg(non_occ)}\n"
                    f"Erro = 1,96 * √((Média * Não ocorrência) / Base)\n"
                    f"Erro = 1,96 * √(({fmt_dbg(mean)} * {fmt_dbg(non_occ)}) / {fmt_dbg(float(base_total_used))}) = {fmt_dbg(margin_error)}\n"
                )
                tip_min = tip_common + f"\nMínima = Média - Erro = {fmt_dbg(mean)} - {fmt_dbg(margin_error)} = {fmt_dbg(min_value)}"
                tip_max = tip_common + f"\nMáxima = Média + Erro = {fmt_dbg(mean)} + {fmt_dbg(margin_error)} = {fmt_dbg(max_value)}"

            raw_mean = round_half_up(mean, 1)
            raw_min  = round_half_up(min_value, 1) if isinstance(min_value, (int, float)) else None
            raw_max  = round_half_up(max_value, 1) if isinstance(max_value, (int, float)) else None

            js = (
                "window.__setStats("
                f"{json.dumps(fmt_stat(min_value))}, {json.dumps(fmt_stat(mean))}, {json.dumps(fmt_stat(max_value))}, "
                f"{json.dumps(tip_min)}, {json.dumps(tip_mean)}, {json.dumps(tip_max)}, "
                f"{json.dumps(raw_min)}, {json.dumps(raw_mean)}, {json.dumps(raw_max)}"
                ");"
            )
            self.web_view.page().runJavaScript(js)

        except Exception as e:
            QMessageBox.warning(self.main_window, "Erro", f"Falha ao calcular estatísticas.\n\n{e}")

    @Slot(str, str, str, bool)
    def applyAutoPaint(self, case_key: str, segment_label: str, base_override_text: str, no_overwrite_manual: bool):
        try:
            case_data = self.cases_by_key.get(case_key, {})
            meta = case_data.get("__meta", {}) if isinstance(case_data, dict) else {}

            excel_base_total = meta.get("base_total", None)

            manual_base: Optional[float] = None
            txt = (base_override_text or "").strip()
            if txt:
                try:
                    manual_base = float(txt.replace(".", "").replace(",", "."))
                except ValueError:
                    manual_base = None

            base_total_used = manual_base if (manual_base is not None and manual_base > 0) else (
                float(excel_base_total) if isinstance(excel_base_total, (int, float)) and excel_base_total > 0 else None
            )
            if base_total_used is None:
                QMessageBox.warning(self.main_window, "Base inválida", "Defina uma Base (Total) válida para aplicar ao mapa.")
                return

            region_map = case_data.get(segment_label, {}) if isinstance(case_data, dict) else {}
            items = [(name, val) for name, val in region_map.items() if isinstance(val, (int, float))]
            if not items:
                QMessageBox.information(self.main_window, "Sem dados", "Não há valores numéricos para este segmento.")
                return

            values = [v for (_, v) in items]
            mean = sum(values) / len(values)
            non_occ = 100.0 - mean
            margin_error = 1.96 * math.sqrt((mean * non_occ) / float(base_total_used))
            min_value = mean - margin_error
            max_value = mean + margin_error

            min_r = round_half_up(min_value, 1)
            max_r = round_half_up(max_value, 1)

            out_items = []
            for name, v in items:
                v_r = round_half_up(float(v), 1)

                if v_r < min_r:
                    cls = "low"
                elif v_r > max_r:
                    cls = "high"
                else:
                    cls = "mid"

                out_items.append({"name": name, "value": v_r, "cls": cls})

            payload = {
                "segment": segment_label,
                "min": min_r,
                "max": max_r,
                "noOverwriteManual": bool(no_overwrite_manual),
                "items": out_items
            }

            self.web_view.page().runJavaScript(f"window.__applyAutoPaint({json.dumps(json.dumps(payload, ensure_ascii=False))});")

        except Exception as e:
            QMessageBox.critical(self.main_window, "Erro", f"Falha ao aplicar pintura automática.\n\n{e}")

    @Slot()
    def pickLogo(self):
        path, _ = QFileDialog.getOpenFileName(
            self.main_window,
            "Selecionar logo (PNG/JPG)",
            "",
            "Imagens (*.png *.jpg *.jpeg);;Todos (*.*)"
        )
        if not path:
            return
        try:
            with open(path, "rb") as f:
                self.logo_bytes = f.read()
            self._push_ppt_state()
        except Exception as e:
            QMessageBox.critical(self.main_window, "Erro", f"Não consegui carregar a logo.\n\n{e}")

    @Slot(str, str)
    def addToPptQueue(self, meta_json: str, data_url: str):
        try:
            meta = json.loads(meta_json)

            if not data_url.startswith("data:image/png"):
                QMessageBox.warning(self.main_window, "PPTX", "Imagem inválida (esperado PNG).")
                return

            _, b64 = data_url.split(",", 1)
            png_bytes = base64.b64decode(b64)

            case_key = meta.get("caseKey", "")
            segment = meta.get("segment", "")
            palette = meta.get("palette", {})
            stats = meta.get("stats", {})

            # ✅ NOVO: legenda estruturada (Nome | Cor | %)
            legend_in = meta.get("legend", [])
            legend_rows: List[Dict[str, Any]] = []
            if isinstance(legend_in, list):
                for row in legend_in:
                    if not isinstance(row, dict):
                        continue
                    name = str(row.get("name", "")).strip()
                    color = str(row.get("color", "")).strip() or "#ffffff"
                    value = row.get("value", None)
                    if name and isinstance(value, (int, float)):
                        legend_rows.append({"name": name, "color": color, "value": float(value)})

            if not case_key or not segment:
                QMessageBox.warning(self.main_window, "PPTX", "Meta incompleta (case/segment).")
                return

            if not all(k in palette for k in ("low", "mid", "high")):
                QMessageBox.warning(self.main_window, "PPTX", "Paleta incompleta.")
                return

            if not all(k in stats for k in ("min", "mean", "max")):
                QMessageBox.warning(self.main_window, "PPTX", "Estatísticas incompletas (min/mean/max).")
                return

            # título do caso (pra lista do usuário)
            case_title = ""
            case_data = self.cases_by_key.get(case_key, {})
            meta_case = case_data.get("__meta", {}) if isinstance(case_data, dict) else {}
            case_title = meta_case.get("title", "") or case_key.split(" — ")[0]

            self.ppt_queue.append({
                "case_key": case_key,
                "case_title": case_title,
                "segment": segment,
                "palette": palette,
                "stats": {
                    "min": float(stats["min"]),
                    "mean": float(stats["mean"]),
                    "max": float(stats["max"]),
                },
                "png_bytes": png_bytes,
                "legend_rows": legend_rows,  # ✅ NOVO
            })

            self._push_ppt_state()

        except Exception as e:
            QMessageBox.critical(self.main_window, "Erro", f"Falha ao adicionar à fila.\n\n{e}")

    @Slot()
    def clearPptQueue(self):
        self.ppt_queue.clear()
        self._push_ppt_state()

    @Slot()
    def exportPptx(self):
        if not self.ppt_queue:
            QMessageBox.information(self.main_window, "PPTX", "A fila está vazia. Adicione mapas antes de exportar.")
            return

        out_path, _ = QFileDialog.getSaveFileName(
            self.main_window,
            "Salvar PPTX",
            "mapas_calor.pptx",
            "PowerPoint (*.pptx)"
        )
        if not out_path:
            return
        if not out_path.lower().endswith(".pptx"):
            out_path += ".pptx"

        try:
            build_pptx(out_path, self.ppt_queue, self.logo_bytes)
            QMessageBox.information(self.main_window, "PPTX", f"Apresentação salva com {len(self.ppt_queue)} slide(s).")
        except Exception as e:
            QMessageBox.critical(self.main_window, "Erro", f"Falha ao gerar PPTX.\n\n{e}")

    @Slot(str)
    def savePng(self, data_url: str):
        try:
            if not data_url.startswith("data:image/png"):
                QMessageBox.warning(self.main_window, "Exportação", "Formato inesperado ao exportar PNG.")
                return

            _, b64 = data_url.split(",", 1)
            raw = base64.b64decode(b64)

            out_path, _ = QFileDialog.getSaveFileName(
                self.main_window,
                "Salvar PNG",
                "mapa_calor.png",
                "PNG (*.png)"
            )
            if not out_path:
                return
            if not out_path.lower().endswith(".png"):
                out_path += ".png"

            with open(out_path, "wb") as f:
                f.write(raw)

        except Exception as e:
            QMessageBox.critical(self.main_window, "Erro ao salvar PNG", str(e))


class HeatmapPage(QWidget):
    def __init__(self, main_window: QMainWindow, profile: QWebEngineProfile):
        super().__init__()
        self.main_window = main_window

        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(0)  # topbar “cola” no conteúdo

        self.topbar = TopBar(
            title="Mapa de Calor",
            icon_path=resource_path("assets/icons/heatmap.ico"),
            back_icon_path=resource_path("assets/icons/arrow-left.ico"),
            show_help=True
          )

        self.btn_back = self.topbar.btn_back
        layout.addWidget(self.topbar)

        self.view = QWebEngineView(self)
        layout.addWidget(self.view, 1)

        # ✅ Usa profile compartilhado do app
        page = QWebEnginePage(profile, self.view)
        self.view.setPage(page)

        # ✅ baseUrl exclusivo desta aba (isola storage por origem)
        self.view.setHtml(HTML, QUrl("https://heatmap.app.local/"))

        self.bridge = QtBridge(main_window, self.view)
        channel = QWebChannel(self.view.page())
        channel.registerObject("QtBridge", self.bridge)
        self.view.page().setWebChannel(channel)

        self.view.loadFinished.connect(lambda ok: self.bridge._push_ppt_state())

    def cleanup(self):
        pass