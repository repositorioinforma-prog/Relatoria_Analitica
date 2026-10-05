from app.core.resources import resource_base, resource_path
import os
import json
import re
import unicodedata
import math
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import numpy as np
import pandas as pd

from PySide6.QtCore import QUrl, QObject, Slot, Signal
from PySide6.QtWidgets import QWidget, QVBoxLayout, QFileDialog, QMessageBox, QMainWindow
from PySide6.QtWebEngineWidgets import QWebEngineView
from PySide6.QtWebEngineCore import QWebEngineProfile, QWebEnginePage
from PySide6.QtWebChannel import QWebChannel

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from app.ui.topbar import TopBar


# ============================================================
# CONFIG: itens sugeridos no final (igual Crosstabs)
# ============================================================
FIXED_BOTTOM_LABELS = [
    "Não sei",
    "Branco/Nulo",
    "Branco / Nulo",
    "Não votarei",
    "Não irá votar",
    "Não sabe",
    "Sem resposta",
    "(Sem resposta)",
]




def _round_half_up(x: float, ndigits: int = 1) -> float:
    q = Decimal("1").scaleb(-ndigits)
    return float(Decimal(str(x)).quantize(q, rounding=ROUND_HALF_UP))


def _canon(s: str) -> str:
    s = (s or "").strip()
    s = unicodedata.normalize("NFD", s)
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    s = s.lower()
    s = re.sub(r"[^a-z0-9]+", " ", s).strip()
    return s


def _is_peso_name(name: str) -> bool:
    return "peso" in _canon(name)


def _safe_str(v: Any) -> str:
    if v is None:
        return ""
    try:
        if isinstance(v, float) and math.isnan(v):
            return ""
    except Exception:
        pass
    return str(v).strip()


def _apply_fixed_bottom_order(labels: List[str], fixed_bottom: List[str]) -> List[str]:
    fixed_canon = [_canon(x) for x in fixed_bottom]
    present = labels[:]

    fixed_present: List[str] = []
    for f in fixed_canon:
        for lab in present:
            if _canon(lab) == f and lab not in fixed_present:
                fixed_present.append(lab)
                break

    rest = [lab for lab in present if _canon(lab) not in set(fixed_canon)]
    return rest + fixed_present


def _sheet_name_safe(name: str) -> str:
    bad = r"[\[\]\:\*\?\/\\]"
    name = re.sub(bad, " ", name)
    name = re.sub(r"\s+", " ", name).strip()
    return name or "A.R"


def _unique_sheet_name(base: str, used: set) -> str:
    base = _sheet_name_safe(base)
    if len(base) > 31:
        base = base[:31].rstrip()

    if base not in used:
        used.add(base)
        return base

    i = 2
    while True:
        suffix = f" ({i})"
        cut = 31 - len(suffix)
        cand = (base[:cut].rstrip() + suffix)
        if cand not in used:
            used.add(cand)
            return cand
        i += 1


# VBA colors
FILL_RED = "FBE2D5"    # RGB(251, 226, 213)
FONT_RED = "7E350E"    # RGB(126, 53, 14)
FILL_BLUE = "DAE9F8"   # RGB(218, 233, 248)
FONT_BLUE = "153D64"   # RGB(21, 61, 100)


def _cell_style_for_z(z: float) -> Optional[Tuple[str, str]]:
    if z is None or not np.isfinite(z):
        return None
    if z <= -1.85:
        return (FILL_RED, FONT_RED)
    if z >= 1.85:
        return (FILL_BLUE, FONT_BLUE)
    return None


def _round_obs_to_integer(ct: pd.DataFrame) -> pd.DataFrame:
    return ct.apply(
        lambda s: s.map(
            lambda v: float(Decimal(str(v)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
        )
    )


def _pct_matrix(ct: pd.DataFrame, mode: str) -> pd.DataFrame:
    """
    Igual Crosstabs:
      - col: percentual por coluna
      - row: percentual por linha
      - total: percentual no total
    """
    if mode == "col":
        denom = ct.sum(axis=0).replace(0, float("nan"))
        p = ct.divide(denom, axis=1) * 100.0
    elif mode == "row":
        denom = ct.sum(axis=1).replace(0, float("nan"))
        p = ct.divide(denom, axis=0) * 100.0
    elif mode == "total":
        denom = float(ct.values.sum())
        p = (ct / denom) * 100.0 if denom else ct * 0.0
    else:
        p = ct * 0.0

    p = p.fillna(0.0)
    p = p.apply(lambda s: s.map(lambda v: _round_half_up(float(v), 1)))
    return p


def _adjusted_residuals(obs: np.ndarray) -> np.ndarray:
    """
    Adjusted Residual (SPSS-like):
      z_ij = (O_ij - E_ij) / sqrt( E_ij * (1 - r_i) * (1 - c_j) )
    """
    O = obs.astype(float)
    total = float(np.nansum(O))
    if not np.isfinite(total) or total <= 0:
        return np.full_like(O, np.nan)

    row_sum = np.nansum(O, axis=1)
    col_sum = np.nansum(O, axis=0)
    E = np.outer(row_sum, col_sum) / total

    with np.errstate(divide="ignore", invalid="ignore"):
        ri = row_sum / total
        cj = col_sum / total
        denom = np.sqrt(E * (1.0 - ri[:, None]) * (1.0 - cj[None, :]))
        z = (O - E) / denom

    z[~np.isfinite(z)] = np.nan
    return z


# ============================================================
# DEBUG-ONLY START
# Helpers usados apenas para os arquivos de debug.
# Pode remover este bloco quando o problema estiver resolvido.
# ============================================================
def _expected_from_ct(ct: pd.DataFrame) -> np.ndarray:
    O = ct.to_numpy(dtype=float)
    total = float(np.nansum(O))
    if not np.isfinite(total) or total <= 0:
        return np.full_like(O, np.nan)

    row_sum = np.nansum(O, axis=1)
    col_sum = np.nansum(O, axis=0)
    E = np.outer(row_sum, col_sum) / total
    E[~np.isfinite(E)] = np.nan
    return E


def _safe_label_from_map(map_: Dict[Any, str], code: Any) -> str:
    return _safe_str(map_.get(code, _safe_str(code)))
# ============================================================
# DEBUG-ONLY END
# ============================================================

HTML = r'''<!doctype html>
<html lang="pt-br">
<head>
  <meta charset="utf-8"/>
  <meta name="viewport" content="width=device-width,initial-scale=1"/>
  <title>Análise Residual</title>
  <script src="qrc:///qtwebchannel/qwebchannel.js"></script>
  <style>
    html,body{height:100%;margin:0;font-family:system-ui,Arial,sans-serif;background:#fff;}
    .app{display:grid;grid-template-columns:560px 1fr;height:100%}
    .panel{border-right:1px solid #eee;padding:14px 16px 16px;overflow:auto;background:#fff}
    h1{font-size:18px;margin:0 0 12px}
    h2{font-size:14px;margin:16px 0 8px}
    label{font-size:13px;display:block;margin:8px 0 6px}
    small{color:#666;font-size:12px}
    button,input,select{font:inherit;padding:8px 10px;border-radius:10px;border:1px solid #ddd;background:#fff}
    button{cursor:pointer}
    .btn-primary{background:#1E90FF;border-color:#1E90FF;color:#fff}
    .btn-outline{background:#fff;color:#111;border-color:#ccc}
    .btn-sm{padding:6px 8px;border-radius:10px;font-size:12px}
    .inline{display:flex;gap:8px;align-items:center;flex-wrap:wrap}
    .muted{color:#666;font-size:12px}
    .box{border:1px solid #eee;border-radius:14px;padding:10px 10px 8px;background:#fafafa;margin-bottom:10px;}

    .listbox{border:1px solid #e6e6e6;border-radius:12px;background:#fff;padding:6px;max-height:230px;overflow:auto;}
    .row{display:flex;align-items:center;gap:8px;padding:4px 4px}
    .row input{margin:0}
    .row .lab{font-size:12px;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;flex:1;min-width:0;}

    .search{width:100%;box-sizing:border-box}

    .weightList{border:1px solid #e6e6e6;border-radius:12px;background:#fff;padding:6px;max-height:180px;overflow:auto;}
    .weightItem{display:flex;align-items:center;gap:8px;width:100%;padding:8px 10px;border:none;border-radius:10px;background:#fff;color:#111;text-align:left;cursor:pointer;box-sizing:border-box}
    .weightItem:hover{background:#f5f5f5}
    .weightItem.active{background:#eaf3ff;color:#0b57d0;font-weight:800}
    .weightDot{width:10px;height:10px;border-radius:999px;border:1px solid #bdbdbd;flex:0 0 10px;background:#fff}
    .weightItem.active .weightDot{background:#0b57d0;border-color:#0b57d0}
    .weightText{font-size:12px;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;flex:1;min-width:0}

    .renameList{border:1px solid #e6e6e6;border-radius:12px;background:#fff;padding:8px;max-height:210px;overflow:auto;}
    .renameItem{display:grid;grid-template-columns:90px 1fr;gap:8px;align-items:center;padding:6px 2px;border-bottom:1px solid #f2f2f2}
    .renameItem:last-child{border-bottom:none}
    .renameKey{font-size:12px;color:#333;font-weight:700;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}
    .renameInp{width:100%;box-sizing:border-box;font-size:12px}

    .queue{border:1px solid #e6e6e6;border-radius:12px;background:#fff;padding:8px;max-height:190px;overflow:auto;}
    .qitem{display:flex;align-items:center;justify-content:space-between;gap:8px;padding:6px 4px;border-bottom:1px solid #f2f2f2}
    .qitem:last-child{border-bottom:none}
    .qtitle{font-size:12px;font-weight:700;color:#111;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;flex:1}
    .qbtn{padding:4px 8px;border-radius:10px;font-size:12px;border:1px solid #ccc;background:#fff;cursor:pointer}

    #main{height:100%;background:#fff;overflow:hidden;display:flex;flex-direction:column;}
    #headerbar{padding:10px 12px;border-bottom:1px solid #eee;display:flex;align-items:center;justify-content:space-between;gap:10px;}
    #title{font-weight:800;font-size:14px;color:#111;margin:0;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;}
    #subtitle{font-size:12px;color:#666;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;}
    #wrap{flex:1;overflow:auto;padding:12px;}

    .card{margin:0 0 18px 0;border:1px solid #eee;border-radius:14px;overflow:hidden}
    .cardHead{padding:10px 12px;background:#fafafa;border-bottom:1px solid #eee;display:flex;gap:10px;align-items:center;justify-content:space-between;flex-wrap:wrap}
    .cardTitleWrap{display:flex;gap:8px;align-items:center;flex-wrap:wrap;min-width:340px;flex:1}
    .cardTitleLabel{font-weight:900;font-size:13px;white-space:nowrap}
    .titleInput{flex:1;min-width:240px;padding:8px 10px;border-radius:10px;border:1px solid #ddd;font:inherit;}
    .cardBtns{display:flex;gap:8px;flex-wrap:wrap}

    .tableWrap{overflow:auto;border:1px solid #eee;border-radius:12px;margin-top:10px}
    table{border-collapse:collapse;font-size:12px;min-width:720px;background:#fff;}
    th,td{border:1px solid #e6e6e6;padding:8px 8px;background:#fff;text-align:center;font-variant-numeric:tabular-nums;white-space:nowrap;}
    th{font-weight:900;position:sticky;top:0;z-index:5;background:#f7f7f7;}
    td.firstcol, th.firstcol{position:sticky;left:0;z-index:6;background:#f7f7f7;font-weight:900;text-align:left;}

    .cell{display:flex;flex-direction:column;gap:2px;line-height:1.1}
    .pct{font-size:11px;opacity:.85}
    .z{font-size:12px;font-weight:800}
    .badge{font-size:12px;padding:4px 8px;border-radius:999px;border:1px solid #e6e6e6;background:#fff;white-space:nowrap;}

    .modal-backdrop{
      position:fixed; inset:0; background:rgba(0,0,0,.25);
      display:none; align-items:center; justify-content:center; z-index:9999;
    }
    .modal-backdrop.open{display:flex}
    .modal{
      width:min(720px, calc(100vw - 32px));
      max-height:min(80vh, 760px);
      background:#fff; border-radius:16px; border:1px solid #e6e6e6;
      box-shadow:0 16px 40px rgba(0,0,0,.18);
      display:flex; flex-direction:column; gap:10px; padding:14px;
    }
    .modal h3{margin:0;font-size:14px}
    .modal .rows{border:1px solid #eee;border-radius:12px;overflow:auto;flex:1;min-height:120px;}
    .modal .r{
      display:grid; grid-template-columns:1fr 1fr auto; gap:10px;
      align-items:center; padding:10px 10px; border-bottom:1px solid #f0f0f0; background:#fff;
    }
    .modal .r:last-child{border-bottom:none}
    .modal .orig{font-weight:800;font-size:13px;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}
    .modal .edit{width:100%;box-sizing:border-box}
    .modal .actions{display:flex;gap:6px;flex-wrap:nowrap}
    .modal .actions button{padding:6px 8px;border-radius:10px}
    .modal .footer{display:flex;justify-content:flex-end;gap:8px}
  </style>
</head>
<body>
<div class="app">
  <aside class="panel">
    <h1>Análise Residual (A.R)</h1>

    <div class="box">
      <h2 style="margin:0 0 8px">0) Dados</h2>
      <label>Importar banco</label>
      <div class="inline">
        <button id="btnPick" class="btn-outline" type="button">Importar CSV/XLSX/SAV</button>
        <small id="dataStatus" class="muted"></small>
      </div>
    </div>

    <div class="box">
      <h2 style="margin:0 0 8px">1) Modo</h2>
      <select id="mode" style="width:100%">
        <option value="batch">1) Gerar automaticamente (1 tabela por variável LINHA)</option>
        <option value="queue">2) Gerar 1 por vez + Adicionar à fila</option>
      </select>
      <div class="muted" style="margin-top:8px" id="modeHint"></div>
    </div>

    <div class="box">
      <h2 style="margin:0 0 8px">2) Peso / Percentual</h2>

      <label>Variável de peso (obrigatório)</label>
      <div id="weightList" class="weightList"></div>
      <small class="muted" id="weightHint">Mostra primeiro colunas que contêm “peso”.</small>

      <label style="margin-top:10px">Percentual exibido</label>
      <select id="pctMode" style="width:220px">
        <option value="col" selected>Coluna (%)</option>
        <option value="row">Linha (%)</option>
        <option value="total">Total (%)</option>
      </select>

      <label style="margin-top:10px">Casas decimais (Z)</label>
      <input id="decimals" type="number" min="0" max="10" step="1" value="1" style="width:120px"/>
    </div>

    <div class="box">
      <h2 style="margin:0 0 8px">3) Selecionar variáveis</h2>

      <div style="margin-bottom:12px">
        <label>Variáveis LINHA</label>
        <input id="searchRow" class="search" type="text" placeholder="Buscar..."/>
        <div id="rowsList" class="listbox" style="margin-top:8px"></div>

        <div class="inline" style="margin-top:10px">
          <button id="btnAllRows" class="btn-outline btn-sm" type="button">Todos Linhas</button>
          <button id="btnNoneRows" class="btn-outline btn-sm" type="button">Nenhum Linhas</button>
        </div>
      </div>

      <div>
        <label>Variáveis COLUNA</label>
        <input id="searchCol" class="search" type="text" placeholder="Buscar..."/>
        <div id="colsList" class="listbox" style="margin-top:8px"></div>

        <div class="inline" style="margin-top:10px">
          <button id="btnAllCols" class="btn-outline btn-sm" type="button">Todos Colunas</button>
          <button id="btnNoneCols" class="btn-outline btn-sm" type="button">Nenhum Colunas</button>
        </div>
      </div>

      <div class="inline" style="margin-top:12px;justify-content:space-between">
        <span class="badge" id="selBadge">0 linhas · 0 colunas</span>
        <small id="status" class="muted"></small>
      </div>

      <div class="inline" style="margin-top:10px">
        <button id="btnPreview" class="btn-primary" type="button">Gerar prévia</button>
        <button id="btnAddQueue" class="btn-outline" type="button">Adicionar à fila</button>
      </div>
    </div>

    <div class="box">
      <h2 style="margin:0 0 8px">4) Renomear variáveis</h2>
      <small class="muted">Isso muda apenas o nome exibido/Exportação (não altera o banco).</small>
      <div id="renameList" class="renameList"></div>
    </div>

    <div class="box">
      <h2 style="margin:0 0 8px">5) Fila / Exportação</h2>
      <div class="inline" style="margin-top:6px">
        <button id="btnClearQueue" class="btn-outline" type="button">Limpar fila</button>
        <button id="btnExportXlsx" class="btn-primary" type="button">Exportar XLSX</button>
        <small id="expStatus" class="muted"></small>
      </div>
      <div id="queue" class="queue" style="margin-top:10px"></div>
      <small class="muted">No modo 1 (automático), gere a prévia e exporte.</small>
    </div>
  </aside>

  <main id="main">
    <div id="headerbar">
      <div style="min-width:0">
        <div id="title">Prévia</div>
        <div id="subtitle">Selecione variáveis e gere a prévia.</div>
      </div>
      <span class="badge" id="meta">—</span>
    </div>
    <div id="wrap">
      <div class="muted">Nada para mostrar ainda.</div>
    </div>
  </main>
</div>

<div class="modal-backdrop" id="modalBackdrop">
  <div class="modal">
    <h3 id="modalTitle">Ordenar linhas</h3>
    <div class="rows" id="modalRows"></div>
    <div class="footer">
      <button class="btn-outline" id="btnModalCancel" type="button">Cancelar</button>
      <button class="btn-primary" id="btnModalSave" type="button">Salvar</button>
    </div>
  </div>
</div>

<script>
  window.QtBridge = null;
  new QWebChannel(qt.webChannelTransport, function(channel){
    window.QtBridge = channel.objects.QtBridge;
  });

  function el(id){ return document.getElementById(id); }

  let __vars = [];
  let __labels = {};
  let __dispMap = {};
  let __weights = [];
  let __selectedWeight = "";

  let __selRows = new Set();
  let __selCols = new Set();

  let __queue = [];
  let __preview = null;

  let __modal = {
    open:false,
    mode:'rows',
    tableIndex:-1,
    blockIndex:-1,
    items:[]
  };

  function dispVar(v){
    return (__dispMap && __dispMap[v]) ? __dispMap[v] : ((__labels && __labels[v]) ? __labels[v] : v);
  }

  function updateBadge(){
    el('selBadge').textContent = `${__selRows.size} linhas · ${__selCols.size} colunas`;
  }

  function mode(){
    return el('mode').value || 'batch';
  }

  function updateModeUI(){
    const m = mode();
    if(m === 'batch'){
      el('modeHint').textContent = 'Selecione VÁRIAS LINHAS e VÁRIAS COLUNAS. O app gera 1 tabela por LINHA, com todas as COLUNAS juntas.';
      el('btnAddQueue').style.display = 'none';
    } else {
      el('modeHint').textContent = 'Selecione 1 LINHA e várias COLUNAS. Gere a prévia e clique em “Adicionar à fila”.';
      el('btnAddQueue').style.display = 'inline-block';
    }
  }

  function renderWeightPicker(){
    const box = el('weightList');
    box.innerHTML = '';

    const list = (__weights && __weights.length) ? __weights.slice() : __vars.slice();

    if(!list.length){
      box.innerHTML = '<small class="muted">Nenhuma variável disponível.</small>';
      return;
    }

    list.forEach(v=>{
      const btn = document.createElement('button');
      btn.type = 'button';
      btn.className = 'weightItem' + ((__selectedWeight === v) ? ' active' : '');

      const dot = document.createElement('span');
      dot.className = 'weightDot';

      const txt = document.createElement('span');
      txt.className = 'weightText';
      txt.textContent = dispVar(v);
      txt.title = v;

      btn.appendChild(dot);
      btn.appendChild(txt);

      btn.addEventListener('click', ()=>{
        __selectedWeight = v;
        renderWeightPicker();
      });

      box.appendChild(btn);
    });
  }

  function renderVarList(containerId, allVars, selectedSet, searchId, enforceSingle){
    const box = el(containerId);
    box.innerHTML = '';
    const q = (el(searchId).value || '').toLowerCase().trim();

    const show = allVars.filter(v=>{
      const dn = String(dispVar(v)).toLowerCase();
      const vn = String(v).toLowerCase();
      return !q || dn.includes(q) || vn.includes(q);
    });

    for(const v of show){
      const row = document.createElement('div');
      row.className = 'row';

      const cb = document.createElement('input');
      cb.type = 'checkbox';
      cb.checked = selectedSet.has(v);

      cb.addEventListener('change', ()=>{
        if(cb.checked){
          if(enforceSingle){
            selectedSet.clear();
            selectedSet.add(v);
            renderAll();
          } else {
            selectedSet.add(v);
            updateBadge();
            refreshRenameList();
          }
        } else {
          selectedSet.delete(v);
          updateBadge();
          refreshRenameList();
        }
      });

      const lab = document.createElement('div');
      lab.className = 'lab';
      lab.textContent = dispVar(v);
      lab.title = v;

      row.appendChild(cb);
      row.appendChild(lab);
      box.appendChild(row);
    }
  }

  function renderSelectionLists(){
    const enforceSingleRow = (mode() === 'queue');
    renderVarList('rowsList', __vars, __selRows, 'searchRow', enforceSingleRow);
    renderVarList('colsList', __vars, __selCols, 'searchCol', false);
    updateBadge();
  }

  function refreshRenameList(){
    const box = el('renameList');

    const focused = document.activeElement;
    const focusedVar = (focused && focused.classList && focused.classList.contains('renameInp'))
      ? focused.getAttribute('data-var')
      : null;
    const selStart = (focused && typeof focused.selectionStart === 'number') ? focused.selectionStart : null;
    const selEnd = (focused && typeof focused.selectionEnd === 'number') ? focused.selectionEnd : null;

    box.innerHTML = '';

    const selected = Array.from(new Set([...Array.from(__selRows.values()), ...Array.from(__selCols.values())]));
    if(!selected.length){
      box.innerHTML = '<small class="muted">Selecione variáveis para renomear.</small>';
      return;
    }

    for(const v of selected){
      const row = document.createElement('div');
      row.className = 'renameItem';

      const key = document.createElement('div');
      key.className = 'renameKey';
      key.textContent = v;
      key.title = v;

      const inp = document.createElement('input');
      inp.className = 'renameInp';
      inp.type = 'text';
      inp.value = dispVar(v);
      inp.placeholder = v;
      inp.setAttribute('data-var', v);

      inp.addEventListener('input', ()=>{
        const nv = (inp.value || '').trim();
        __dispMap[v] = nv ? nv : ((__labels[v] || '').trim() ? __labels[v] : v);
      });

      row.appendChild(key);
      row.appendChild(inp);
      box.appendChild(row);
    }

    if(focusedVar){
      const newInp = box.querySelector(`.renameInp[data-var="${focusedVar}"]`);
      if(newInp){
        newInp.focus();
        try{
          if(selStart !== null && selEnd !== null){
            newInp.setSelectionRange(selStart, selEnd);
          }
        }catch(_){}
      }
    }
  }

  function renderAll(){
    renderSelectionLists();
    refreshRenameList();
  }

  function renderQueue(){
    const box = el('queue');
    box.innerHTML = '';

    if(mode() !== 'queue'){
      box.innerHTML = '<small class="muted">Fila só funciona no modo 2.</small>';
      return;
    }

    if(!__queue.length){
      box.innerHTML = '<small class="muted">Fila vazia.</small>';
      return;
    }

    __queue.forEach((it, idx)=>{
      const row = document.createElement('div');
      row.className = 'qitem';

      const t = document.createElement('div');
      t.className = 'qtitle';
      t.textContent = it.title || `Tabela ${idx+1}`;

      const btn = document.createElement('button');
      btn.className = 'qbtn';
      btn.textContent = 'Remover';
      btn.addEventListener('click', ()=> window.QtBridge?.removeQueueItem?.(idx));

      row.appendChild(t);
      row.appendChild(btn);
      box.appendChild(row);
    });
  }

  function fmtZ(v, d){
    if(v === null || v === undefined || !isFinite(v)) return '—';
    const dd = Math.max(0, Math.min(10, Number(d) || 1));
    return Number(v).toFixed(dd);
  }

  function fmtPct(v){
    if(v === null || v === undefined || !isFinite(v)) return '—';
    return Number(v).toFixed(1) + '%';
  }

  function styleForZ(z){
    if(z === null || z === undefined || !isFinite(z)) return null;
    if(z <= -1.85) return {bg:'rgb(251,226,213)', fg:'rgb(126,53,14)'};
    if(z >= 1.85) return {bg:'rgb(218,233,248)', fg:'rgb(21,61,100)'};
    return null;
  }

  function reorderBlockColsByLabels(block, sourceItems, targetItems){
    if(!block || !Array.isArray(block.col_cats)) return block;

    const oldColCats = (block.col_cats || []).slice();
    const oldPercent = (block.percent || []).map(r => (r || []).slice());
    const oldZ = (block.z || []).map(r => (r || []).slice());
    const oldObs = (block.debug_obs || []).map(r => (r || []).slice());
    const oldExp = (block.debug_exp || []).map(r => (r || []).slice());
    const oldColCodes = Array.isArray(block.debug_col_codes) ? block.debug_col_codes.slice() : [];
    const oldColDisplay = Array.isArray(block.debug_col_display) ? block.debug_col_display.slice() : [];

    const idxs = targetItems.map(item => Number(item.id)).filter(idx => !Number.isNaN(idx) && idx >= 0 && idx < oldColCats.length);
    const newLabels = targetItems.map(item => (item.text || "").trim());

    block.col_cats = idxs.map((idx, pos) => newLabels[pos] || oldColCats[idx]);
    block.percent = oldPercent.map(row => idxs.map(i => row[i]));
    block.z = oldZ.map(row => idxs.map(i => row[i]));
    block.debug_obs = oldObs.map(row => idxs.map(i => row[i]));
    block.debug_exp = oldExp.map(row => idxs.map(i => row[i]));
    block.debug_col_codes = idxs.map(i => oldColCodes[i] || "");
    block.debug_col_display = idxs.map((idx, pos) => newLabels[pos] || oldColDisplay[idx] || oldColCats[idx]);

    return block;
    }

  function reorderBlockRowsByLabels(block, sourceItems, targetItems){
    if(!block || !Array.isArray(block.row_cats)) return block;

    const colCount = Array.isArray(block.col_cats) ? block.col_cats.length : 0;

    const oldRowCats = (block.row_cats || []).slice();
    const oldPercent = (block.percent || []).slice();
    const oldZ = (block.z || []).slice();
    const oldObs = (block.debug_obs || []).slice();
    const oldExp = (block.debug_exp || []).slice();
    const oldRowCodes = Array.isArray(block.debug_row_codes) ? block.debug_row_codes.slice() : [];
    const oldRowDisplay = Array.isArray(block.debug_row_display) ? block.debug_row_display.slice() : [];

    const newRowCats = [];
    const newPercent = [];
    const newZ = [];
    const newObs = [];
    const newExp = [];
    const newRowCodes = [];
    const newRowDisplay = [];

    targetItems.forEach(item=>{
        const idx = Number(item.id);

        if(Number.isNaN(idx) || idx < 0 || idx >= oldRowCats.length){
        newRowCats.push(item.text || "");
        newPercent.push(Array(colCount).fill(0));
        newZ.push(Array(colCount).fill(null));
        newObs.push(Array(colCount).fill(0));
        newExp.push(Array(colCount).fill(null));
        newRowCodes.push("");
        newRowDisplay.push(item.text || "");
        } else {
        const newLabel = (item.text || "").trim() || oldRowCats[idx];
        newRowCats.push(newLabel);
        newPercent.push(oldPercent[idx] || Array(colCount).fill(0));
        newZ.push(oldZ[idx] || Array(colCount).fill(null));
        newObs.push(oldObs[idx] || Array(colCount).fill(0));
        newExp.push(oldExp[idx] || Array(colCount).fill(null));
        newRowCodes.push(oldRowCodes[idx] || "");
        newRowDisplay.push(newLabel);
        }
    });

    block.row_cats = newRowCats;
    block.percent = newPercent;
    block.z = newZ;
    block.debug_obs = newObs;
    block.debug_exp = newExp;
    block.debug_row_codes = newRowCodes;
    block.debug_row_display = newRowDisplay;

    return block;
    }
  function buildCombinedTable(tableObj, decimals){
    const blocks = tableObj.blocks || [];
    if(!blocks.length){
      const empty = document.createElement('div');
      empty.className = 'muted';
      empty.textContent = 'Sem blocos para exibir.';
      return empty;
    }

    const canonicalRows = (blocks[0].row_cats || []).slice();

    const table = document.createElement('table');
    const thead = document.createElement('thead');

    const tr1 = document.createElement('tr');
    const th0 = document.createElement('th');
    th0.className = 'firstcol';
    th0.rowSpan = 2;
    th0.textContent = dispVar(tableObj.row_var || '');
    tr1.appendChild(th0);

    blocks.forEach((b, bi)=>{
      const th = document.createElement('th');
      th.colSpan = Math.max(1, (b.col_cats || []).length);

      const headBox = document.createElement('div');
      headBox.style.display = 'flex';
      headBox.style.alignItems = 'center';
      headBox.style.justifyContent = 'space-between';
      headBox.style.gap = '8px';

      const txt = document.createElement('span');
      txt.textContent = dispVar(b.col_var || '');

      const bc = document.createElement('button');
      bc.type = 'button';
      bc.className = 'btn-outline btn-sm';
      bc.textContent = 'Ordenar colunas';
      bc.addEventListener('click', (ev)=>{
        ev.stopPropagation();
        openColModal(tableObj.__tableIndex, bi);
      });

      headBox.appendChild(txt);
      headBox.appendChild(bc);
      th.appendChild(headBox);
      tr1.appendChild(th);
    });

    const tr2 = document.createElement('tr');
    blocks.forEach(b=>{
      (b.col_cats || []).forEach(cc=>{
        const th = document.createElement('th');
        th.textContent = cc;
        tr2.appendChild(th);
      });
    });

    thead.appendChild(tr1);
    thead.appendChild(tr2);
    table.appendChild(thead);

    const tbody = document.createElement('tbody');

    for(let i=0;i<canonicalRows.length;i++){
      const trPct = document.createElement('tr');
      const trZ = document.createElement('tr');

      const td0 = document.createElement('td');
      td0.className = 'firstcol';
      td0.rowSpan = 2;
      td0.textContent = canonicalRows[i];
      trPct.appendChild(td0);

      blocks.forEach(b=>{
        const pct = b.percent || [];
        const z = b.z || [];

        for(let j=0;j<(b.col_cats || []).length;j++){
          const tdPct = document.createElement('td');
          const tdZ = document.createElement('td');

          tdPct.textContent = fmtPct((pct[i] || [])[j]);
          tdZ.textContent = fmtZ((z[i] || [])[j], decimals);

          const st = styleForZ((z[i] || [])[j]);
          if(st){
            tdPct.style.background = st.bg;
            tdPct.style.color = st.fg;
            tdPct.style.fontWeight = '800';

            tdZ.style.background = st.bg;
            tdZ.style.color = st.fg;
            tdZ.style.fontWeight = '800';
          }

          trPct.appendChild(tdPct);
          trZ.appendChild(tdZ);
        }
      });

      tbody.appendChild(trPct);
      tbody.appendChild(trZ);
    }

    table.appendChild(tbody);

    const wrap = document.createElement('div');
    wrap.className = 'tableWrap';
    wrap.appendChild(table);
    return wrap;
  }

    function openRowModal(tableIndex){
        if(!__preview || !__preview.ok) return;
        const t = (__preview.tables || [])[tableIndex];
        if(!t) return;

        const labels = ((t.blocks || [])[0]?.row_cats || []).slice();
        __modal.open = true;
        __modal.mode = 'rows';
        __modal.tableIndex = tableIndex;
        __modal.blockIndex = -1;
        __modal.items = labels.map((l, idx) => ({id: idx, old: l, text: l}));

        el('modalTitle').textContent = `Ordenar linhas: ${dispVar(t.row_var || '')}`;
        renderModalRows();
        el('modalBackdrop').classList.add('open');
   }

    function openColModal(tableIndex, blockIndex){
        if(!__preview || !__preview.ok) return;
        const t = (__preview.tables || [])[tableIndex];
        if(!t) return;

        const b = (t.blocks || [])[blockIndex];
        if(!b) return;

        const labels = (b.col_cats || []).slice();
        __modal.open = true;
        __modal.mode = 'cols';
        __modal.tableIndex = tableIndex;
        __modal.blockIndex = blockIndex;
        __modal.items = labels.map((l, idx) => ({id: idx, old: l, text: l}));

        el('modalTitle').textContent = `Ordenar colunas: ${dispVar(t.row_var || '')} × ${dispVar(b.col_var || '')}`;
        renderModalRows();
        el('modalBackdrop').classList.add('open');
    }
    function renderModalRows(){
    const box = el('modalRows');
    box.innerHTML = '';

    __modal.items.forEach((it, idx)=>{
      const row = document.createElement('div');
      row.className = 'r';

      const orig = document.createElement('div');
      orig.className = 'orig';
      orig.textContent = it.old;
      orig.title = it.old;

      const edit = document.createElement('input');
      edit.className = 'edit';
      edit.type = 'text';
      edit.value = it.text;
      edit.addEventListener('input', ()=>{
        __modal.items[idx].text = edit.value;
      });

      const actions = document.createElement('div');
      actions.className = 'actions';

      const up = document.createElement('button');
      up.type = 'button';
      up.className = 'btn-outline btn-sm';
      up.textContent = '↑';
      up.disabled = idx === 0;
      up.addEventListener('click', ()=>{
        if(idx <= 0) return;
        const tmp = __modal.items[idx-1];
        __modal.items[idx-1] = __modal.items[idx];
        __modal.items[idx] = tmp;
        renderModalRows();
      });

      const down = document.createElement('button');
      down.type = 'button';
      down.className = 'btn-outline btn-sm';
      down.textContent = '↓';
      down.disabled = idx >= (__modal.items.length - 1);
      down.addEventListener('click', ()=>{
        if(idx >= (__modal.items.length - 1)) return;
        const tmp = __modal.items[idx+1];
        __modal.items[idx+1] = __modal.items[idx];
        __modal.items[idx] = tmp;
        renderModalRows();
      });

      actions.appendChild(up);
      actions.appendChild(down);

      row.appendChild(orig);
      row.appendChild(edit);
      row.appendChild(actions);
      box.appendChild(row);
    });
  }

  function closeModal(){
    __modal.open = false;
    __modal.mode = 'rows';
    __modal.tableIndex = -1;
    __modal.blockIndex = -1;
    __modal.items = [];
    el('modalBackdrop').classList.remove('open');
  }

  function applyRowModal(){
    if(!__modal.open || !__preview || !__preview.ok) return closeModal();

    const t = (__preview.tables || [])[__modal.tableIndex];
    if(!t) return closeModal();

    const sourceItems = __modal.items
        .slice()
        .sort((a,b)=> Number(a.id) - Number(b.id));

    const targetItems = __modal.items.slice();

    if(__modal.mode === 'rows'){
        (t.blocks || []).forEach(b=>{
        reorderBlockRowsByLabels(b, sourceItems, targetItems);
        });
    } else if(__modal.mode === 'cols'){
        const b = (t.blocks || [])[__modal.blockIndex];
        if(b){
        reorderBlockColsByLabels(b, sourceItems, targetItems);
        }
    }

    renderPreview(__preview);
    closeModal();
    }

  function renderPreview(payload){
    const wrap = el('wrap');
    wrap.innerHTML = '';

    if(!payload || !payload.ok){
      wrap.innerHTML = '<div class="muted">Nada para mostrar.</div>';
      return;
    }

    const tables = payload.tables || [];
    if(!tables.length){
      wrap.innerHTML = '<div class="muted">Nada para mostrar.</div>';
      return;
    }

    const decimals = payload.decimals || 1;

    tables.forEach((t, ti)=>{
      t.__tableIndex = ti;

      const card = document.createElement('div');
      card.className = 'card';

      const head = document.createElement('div');
      head.className = 'cardHead';

      const tw = document.createElement('div');
      tw.className = 'cardTitleWrap';

      const lab = document.createElement('div');
      lab.className = 'cardTitleLabel';
      lab.textContent = `Tabela: ${dispVar(t.row_var)}`;

      const inp = document.createElement('input');
      inp.className = 'titleInput';
      inp.type = 'text';
      inp.value = t.title || (`A.R ${dispVar(t.row_var)}`);
      inp.addEventListener('input', ()=>{
        t.title = (inp.value || '').trim();
      });

      tw.appendChild(lab);
      tw.appendChild(inp);

      const btns = document.createElement('div');
      btns.className = 'cardBtns';

      const badge = document.createElement('span');
      badge.className = 'badge';
      badge.textContent = `${(t.blocks || []).length} bloco(s)`;
      btns.appendChild(badge);

      const br = document.createElement('button');
      br.type = 'button';
      br.className = 'btn-outline btn-sm';
      br.textContent = 'Ordenar linhas';
      br.addEventListener('click', ()=> openRowModal(ti));
      btns.appendChild(br);

      head.appendChild(tw);
      head.appendChild(btns);
      card.appendChild(head);

      card.appendChild(buildCombinedTable(t, decimals));
      wrap.appendChild(card);
    });
  }

  el('btnPick').addEventListener('click', ()=> window.QtBridge?.pickData?.());

  el('mode').addEventListener('change', ()=>{
    if(mode() === 'queue' && __selRows.size > 1){
      const first = Array.from(__selRows.values())[0];
      __selRows = new Set([first]);
    }
    updateModeUI();
    renderAll();
  });

  el('searchRow').addEventListener('input', renderSelectionLists);
  el('searchCol').addEventListener('input', renderSelectionLists);

  el('btnAllRows').addEventListener('click', ()=>{
    if(mode() === 'queue'){
      if(__vars.length) __selRows = new Set([__vars[0]]);
    } else {
      __vars.forEach(v => __selRows.add(v));
    }
    renderAll();
  });

  el('btnNoneRows').addEventListener('click', ()=>{
    __selRows.clear();
    renderAll();
  });

  el('btnAllCols').addEventListener('click', ()=>{
    __vars.forEach(v => __selCols.add(v));
    renderAll();
  });

  el('btnNoneCols').addEventListener('click', ()=>{
    __selCols.clear();
    renderAll();
  });

  el('btnPreview').addEventListener('click', ()=>{
    const weightVar = __selectedWeight || '';
    if(!weightVar){
      el('status').textContent = 'Selecione o peso.';
      return;
    }

    const rowVars = Array.from(__selRows.values());
    const colVars = Array.from(__selCols.values());

    if(mode() === 'queue' && rowVars.length !== 1){
      el('status').textContent = 'No modo fila, selecione 1 variável LINHA.';
      return;
    }

    if(!rowVars.length || !colVars.length){
      el('status').textContent = 'Selecione pelo menos 1 LINHA e 1 COLUNA.';
      return;
    }

    const payload = {
      mode: mode(),
      rowVars,
      colVars,
      weightVar,
      pctMode: el('pctMode').value || 'col',
      decimals: Math.max(0, Math.min(10, Number(el('decimals').value) || 1))
    };

    el('status').textContent = 'Calculando...';
    window.QtBridge?.computePreview?.(JSON.stringify(payload));
  });

  el('btnAddQueue').addEventListener('click', ()=>{
    if(mode() !== 'queue') return;

    if(!__preview || !__preview.ok || !(__preview.tables || []).length){
      el('status').textContent = 'Gere a prévia antes de adicionar à fila.';
      return;
    }

    const t = __preview.tables[0];
    window.QtBridge?.addToQueueFromPreview?.(JSON.stringify(t));
  });

  el('btnClearQueue').addEventListener('click', ()=> window.QtBridge?.clearQueue?.());

  el('btnExportXlsx').addEventListener('click', ()=>{
    const payload = {
      mode: mode(),
      displayMap: (__dispMap || {}),
      previewTables: (__preview && __preview.ok) ? (__preview.tables || []) : []
    };
    window.QtBridge?.exportXlsx?.(JSON.stringify(payload));
  });

  el('btnModalCancel').addEventListener('click', closeModal);
  el('btnModalSave').addEventListener('click', applyRowModal);

  window.__setData = function(payloadJson){
    try{
      const p = JSON.parse(payloadJson);
      el('dataStatus').textContent = p.label || '';

      __vars = (p.columns || []).map(c => c.name);
      __labels = p.labels || {};

      __dispMap = {};
      __vars.forEach(v=>{
        const lab = (__labels[v] || '').trim();
        __dispMap[v] = lab ? lab : v;
      });

      __selRows = new Set();
      __selCols = new Set();

      __weights = p.weight_candidates || [];
      __selectedWeight = '';

      if(__weights.length){
        el('weightHint').textContent = 'Mostra primeiro colunas que contêm “peso”.';
      } else {
        el('weightHint').textContent = 'Não achei colunas com “peso”. Mostrando todas para você escolher.';
      }

      renderWeightPicker();

      __preview = null;
      el('status').textContent = '';
      el('expStatus').textContent = '';
      el('meta').textContent = '—';

      updateModeUI();
      renderAll();
      renderQueue();

      el('wrap').innerHTML = '<div class="muted">Pronto para gerar a prévia.</div>';
    }catch(e){
      console.error(e);
    }
  };

  window.__setPreview = function(payloadJson){
    try{
      const p = JSON.parse(payloadJson);
      __preview = p;

      if(!p.ok){
        el('status').textContent = p.error || 'Erro.';
        return;
      }

      el('status').textContent = 'Pronto.';
      el('meta').textContent = `${(p.tables || []).length} tabela(s)`;
      renderPreview(p);
    }catch(e){
      console.error(e);
      el('status').textContent = 'Erro ao receber prévia.';
    }
  };

  window.__setQueueState = function(payloadJson){
    try{
      const p = JSON.parse(payloadJson);
      __queue = p.items || [];
      el('expStatus').textContent = p.message || '';
      renderQueue();

      if(p.resetRowSelection){
        __selRows.clear();
        renderAll();
      }
    }catch(e){
      console.error(e);
    }
  };
</script>
</body>
</html>
'''


class ResidualsBridge(QObject):
    pushData = Signal(str)
    pushPreview = Signal(str)
    pushQueueState = Signal(str)

    def __init__(self, main_window: QMainWindow, web_view: QWebEngineView):
        super().__init__()
        self.main_window = main_window
        self.web_view = web_view
        self.df: Optional[pd.DataFrame] = None
        self.file_label: str = ""
        self.var_labels: Dict[str, str] = {}
        self.value_labels: Dict[str, Dict[Any, Any]] = {}
        self.missing_ranges: Dict[str, List[Any]] = {}
        self.queue: List[Dict[str, Any]] = []

    def _push_queue_state(self, message: str = "", reset_row: bool = False) -> None:
        items = []
        for it in self.queue:
            items.append({
                "title": str(it.get("title") or ""),
            })
        payload = {
            "items": items,
            "message": message,
            "resetRowSelection": bool(reset_row),
        }
        self.web_view.page().runJavaScript(
            f"window.__setQueueState({json.dumps(json.dumps(payload, ensure_ascii=False))});"
        )

    def _apply_value_labels_for_var(self, var: str, s: pd.Series) -> Tuple[pd.Series, Dict[Any, str]]:
        labels = self.value_labels.get(var) or {}
        out = s.copy()
        label_map: Dict[Any, str] = {}

        def _label_of(v: Any) -> str:
            if pd.isna(v):
                return ""
            if v in labels:
                return _safe_str(labels[v])
            try:
                fv = float(v)
                if fv in labels:
                    return _safe_str(labels[fv])
                iv = int(fv)
                if iv in labels:
                    return _safe_str(labels[iv])
            except Exception:
                pass
            return _safe_str(v)

        out = s.map(_label_of)
        for raw in s.dropna().unique().tolist():
            label_map[raw] = _label_of(raw)

        return out, label_map

    def _missing_specs_for_var(self, var: str) -> List[Any]:
        raw = self.missing_ranges.get(str(var), []) or []
        if isinstance(raw, list):
            return raw
        return [raw]

    def _values_match_for_missing(self, value: Any, target: Any) -> bool:
        try:
            if pd.isna(value) or pd.isna(target):
                return False
        except Exception:
            pass

        if isinstance(value, str) or isinstance(target, str):
            v = _safe_str(value)
            t = _safe_str(target)
            return (v == t) or (_canon(v) == _canon(t))

        try:
            return float(value) == float(target)
        except Exception:
            return _safe_str(value) == _safe_str(target)

    def _is_spss_user_missing(self, var: str, value: Any) -> bool:
        try:
            if pd.isna(value):
                return False
        except Exception:
            pass

        specs = self._missing_specs_for_var(var)
        if not specs:
            return False

        for spec in specs:
            if isinstance(spec, dict):
                values = spec.get("values")
                if isinstance(values, (list, tuple, set)):
                    for vv in values:
                        if self._values_match_for_missing(value, vv):
                            return True

                lo = spec.get("lo")
                hi = spec.get("hi")
                if lo is None and hi is None and "value" in spec:
                    lo = spec.get("value")
                    hi = spec.get("value")
            elif isinstance(spec, (list, tuple)) and len(spec) >= 2:
                lo, hi = spec[0], spec[1]
            else:
                lo = spec
                hi = spec

            if lo is None and hi is None:
                continue
            if hi is None:
                hi = lo
            if lo is None:
                lo = hi

            if isinstance(value, str) or isinstance(lo, str) or isinstance(hi, str):
                if self._values_match_for_missing(value, lo) or self._values_match_for_missing(value, hi):
                    return True
                continue

            try:
                fv = float(value)
                flo = float(lo)
                fhi = float(hi)
                if flo <= fv <= fhi:
                    return True
            except Exception:
                if self._values_match_for_missing(value, lo) or self._values_match_for_missing(value, hi):
                    return True

        return False

    def _user_missing_mask(self, s: pd.Series, var: str) -> pd.Series:
        specs = self._missing_specs_for_var(var)
        if not specs:
            return pd.Series(False, index=s.index, dtype=bool)
        return s.map(lambda x: bool(self._is_spss_user_missing(var, x))).astype(bool)

    def _weighted_crosstab(self, df: pd.DataFrame, row_var: str, col_var: str, w: pd.Series) -> pd.DataFrame:
        tmp = pd.DataFrame({
            "_row": df[row_var],
            "_col": df[col_var],
            "_w": w,
        })
        ct = pd.pivot_table(
            tmp,
            index="_row",
            columns="_col",
            values="_w",
            aggfunc="sum",
            fill_value=0.0,
            dropna=False,
        )
        ct = ct.sort_index(axis=0).sort_index(axis=1)
        return ct

    def _unweighted_bases(self, df: pd.DataFrame, row_var: str, col_var: str) -> Tuple[List[Any], List[Any]]:
        base_rows = [x for x in pd.unique(df[row_var].dropna())]
        base_cols = [x for x in pd.unique(df[col_var].dropna())]
        return base_rows, base_cols

    def _sort_rows_by_total_desc(self, ct: pd.DataFrame, row_display_labels: List[str]) -> List[Any]:
        row_totals = ct.sum(axis=1)
        aux = []
        for i, code in enumerate(ct.index.tolist()):
            lab = row_display_labels[i] if i < len(row_display_labels) else _safe_str(code)
            aux.append((code, float(row_totals.loc[code]), lab))

        aux.sort(key=lambda t: (-t[1], _canon(t[2])))

        ordered_labels = [t[2] for t in aux]
        ordered_labels = _apply_fixed_bottom_order(ordered_labels, FIXED_BOTTOM_LABELS)

        final_codes = []
        used = set()
        for lab in ordered_labels:
            for code, _, lb in aux:
                if lb == lab and code not in used:
                    final_codes.append(code)
                    used.add(code)
                    break

        for code, _, _ in aux:
            if code not in used:
                final_codes.append(code)

        return final_codes

    def _prepare_pair(
        self,
        row_var: str,
        col_var: str,
        weight_var: str,
    ) -> Tuple[pd.DataFrame, pd.Series, Dict[Any, str], Dict[Any, str]]:
        assert self.df is not None

        cols = [row_var, col_var, weight_var]
        df = self.df[cols].copy()

        for v in [row_var, col_var, weight_var]:
            if v in df.columns:
                mask = self._user_missing_mask(df[v], v)
                if bool(mask.any()):
                    df.loc[mask, v] = np.nan

        df = df.dropna(subset=[row_var, col_var, weight_var])

        w = pd.to_numeric(df[weight_var], errors="coerce").fillna(0.0).astype(float)
        mask = (w > 0)
        df = df.loc[mask, [row_var, col_var]].copy()
        w = w.loc[mask]

        df[row_var], row_map = self._apply_value_labels_for_var(row_var, df[row_var])
        df[col_var], col_map = self._apply_value_labels_for_var(col_var, df[col_var])

        return df, w, row_map, col_map

    def _compute_block(
        self,
        row_var: str,
        col_var: str,
        weight_var: str,
        pct_mode: str,
    ) -> Dict[str, Any]:
        dfp, w, row_map, col_map = self._prepare_pair(row_var, col_var, weight_var)

        if dfp.empty:
            return {
                "col_var": col_var,
                "row_cats": [],
                "col_cats": [],
                "percent": [],
                "z": [],
                "debug_row_codes": [],
                "debug_col_codes": [],
                "debug_row_display": [],
                "debug_col_display": [],
                "debug_obs": [],
                "debug_exp": [],
            }

        ct = self._weighted_crosstab(dfp, row_var, col_var, w)

        base_rows, base_cols = self._unweighted_bases(dfp, row_var, col_var)
        ct = ct.reindex(index=base_rows, columns=base_cols, fill_value=0.0)

        row_disp = [row_map.get(code, _safe_str(code)) for code in ct.index]
        ordered_codes = self._sort_rows_by_total_desc(ct, row_disp)
        ct = ct.reindex(index=ordered_codes, fill_value=0.0)

        ct = _round_obs_to_integer(ct)

        row_disp = [row_map.get(code, _safe_str(code)) for code in ct.index]
        col_disp = [col_map.get(code, _safe_str(code)) for code in ct.columns]

        pct = _pct_matrix(ct, pct_mode)
        z = _adjusted_residuals(ct.to_numpy(dtype=float))

        return {
            "col_var": col_var,
            "row_cats": row_disp,
            "col_cats": col_disp,
            "percent": pct.values.tolist(),
            "z": np.round(z, 6).tolist(),
            "debug_row_codes": [_safe_str(code) for code in ct.index],
            "debug_col_codes": [_safe_str(code) for code in ct.columns],
            "debug_row_display": row_disp,
            "debug_col_display": col_disp,
            "debug_obs": ct.to_numpy(dtype=float).tolist(),
            "debug_exp": _expected_from_ct(ct).tolist(),
        }

    def _reorder_block_rows_to_match(self, block: Dict[str, Any], ordered_row_labels: List[str]) -> Dict[str, Any]:
        old_rows = list(block.get("row_cats") or [])
        idx_map = {lab: i for i, lab in enumerate(old_rows)}

        col_count = len(block.get("col_cats") or [])

        def blank_num() -> List[float]:
            return [0.0] * col_count

        def blank_nan() -> List[Any]:
            return [None] * col_count

        new_percent = []
        new_z = []
        new_obs = []
        new_exp = []

        for lab in ordered_row_labels:
            idx = idx_map.get(lab)
            if idx is None:
                new_percent.append(blank_num())
                new_z.append(blank_nan())
                new_obs.append(blank_num())
                new_exp.append(blank_nan())
            else:
                new_percent.append((block.get("percent") or [])[idx] if idx < len(block.get("percent") or []) else blank_num())
                new_z.append((block.get("z") or [])[idx] if idx < len(block.get("z") or []) else blank_nan())
                new_obs.append((block.get("debug_obs") or [])[idx] if idx < len(block.get("debug_obs") or []) else blank_num())
                new_exp.append((block.get("debug_exp") or [])[idx] if idx < len(block.get("debug_exp") or []) else blank_nan())

        block["row_cats"] = ordered_row_labels[:]
        block["percent"] = new_percent
        block["z"] = new_z
        block["debug_obs"] = new_obs
        block["debug_exp"] = new_exp
        return block

    @Slot()
    def pickData(self) -> None:
        try:
            path, _ = QFileDialog.getOpenFileName(
                self.main_window,
                "Selecione o banco",
                "",
                "Dados (*.csv *.xlsx *.xls *.sav)"
            )
            if not path:
                return

            self.queue = []
            self.var_labels = {}
            self.value_labels = {}
            self.missing_ranges = {}

            if path.lower().endswith(".csv"):
                df = pd.read_csv(path)
            elif path.lower().endswith(".xlsx"):
                df = pd.read_excel(path)
            elif path.lower().endswith(".xls"):
                df = pd.read_excel(path)
            elif path.lower().endswith(".sav"):
                try:
                    import pyreadstat  # type: ignore
                except Exception:
                    raise RuntimeError("Para ler .sav, instale: pip install pyreadstat")

                df = None
                meta = None
                last_err = None

                file_enc = None
                try:
                    _, m0 = pyreadstat.read_sav(path, metadataonly=True)  # type: ignore
                    file_enc = (getattr(m0, "file_encoding", None) or "").strip() or None
                except Exception:
                    file_enc = None

                candidates = []
                for enc in [None, file_enc, "utf-8", "cp1252", "latin1", "iso-8859-1", "iso-8859-15", "utf-16le", "ucs-2le"]:
                    if enc and enc not in candidates:
                        candidates.append(enc)

                for enc in candidates:
                    try:
                        if enc is None:
                            df, meta = pyreadstat.read_sav(path, user_missing=True)  # type: ignore
                        else:
                            df, meta = pyreadstat.read_sav(path, encoding=enc, user_missing=True)  # type: ignore
                        break
                    except Exception as e:
                        last_err = e
                        msg = str(e)
                        if (
                            "Unable to convert string" in msg
                            or "invalid byte sequence" in msg
                            or "codec can't decode" in msg
                            or "invalid continuation byte" in msg
                        ):
                            continue
                        raise

                if df is None or meta is None:
                    raise last_err or Exception("Falha ao ler o arquivo .sav.")

                try:
                    m = getattr(meta, "column_names_to_labels", None)
                    if isinstance(m, dict):
                        self.var_labels = {str(k): str(v) for k, v in m.items() if v}
                except Exception:
                    self.var_labels = {}

                try:
                    for var, m in (meta.variable_value_labels or {}).items():
                        self.value_labels[str(var)] = dict(m)
                except Exception:
                    self.value_labels = {}

                try:
                    raw_missing = getattr(meta, "missing_ranges", None) or {}
                    self.missing_ranges = {str(var): list(specs) for var, specs in raw_missing.items()}
                except Exception:
                    self.missing_ranges = {}
            else:
                raise RuntimeError("Formato não suportado. Use CSV, XLSX ou SAV.")

            df = df.dropna(axis=1, how="all")
            self.df = df
            self.file_label = os.path.basename(path)

            cols = [{"name": str(c), "dtype": str(df[c].dtype)} for c in df.columns]
            weights = [str(c) for c in df.columns if _is_peso_name(str(c))]

            payload = {
                "label": self.file_label,
                "columns": cols,
                "labels": self.var_labels,
                "weight_candidates": weights,
            }

            self.web_view.page().runJavaScript(
                f"window.__setData({json.dumps(json.dumps(payload, ensure_ascii=False))});"
            )
            self._push_queue_state(message="Fila resetada.")

        except Exception as e:
            QMessageBox.critical(self.main_window, "Erro", f"Falha ao carregar dados.\n\n{e}")

    @Slot(str)
    def computePreview(self, payload_json: str) -> None:
        try:
            if self.df is None:
                raise RuntimeError("Nenhum arquivo carregado.")

            payload = json.loads(payload_json) if payload_json else {}
            mode = str(payload.get("mode") or "batch").strip().lower()
            row_vars = payload.get("rowVars") or []
            col_vars = payload.get("colVars") or []
            weight_var = str(payload.get("weightVar") or "").strip()
            pct_mode = str(payload.get("pctMode") or "col").strip().lower()
            decimals = int(payload.get("decimals") or 1)

            if not isinstance(row_vars, list) or not isinstance(col_vars, list):
                raise RuntimeError("Seleção inválida.")
            row_vars = [str(x) for x in row_vars if str(x) in self.df.columns]
            col_vars = [str(x) for x in col_vars if str(x) in self.df.columns]

            if not weight_var or weight_var not in self.df.columns:
                raise RuntimeError("Selecione um peso válido.")
            if not row_vars or not col_vars:
                raise RuntimeError("Selecione ao menos 1 LINHA e 1 COLUNA.")
            if mode == "queue" and len(row_vars) != 1:
                raise RuntimeError("No modo fila, selecione 1 variável LINHA.")

            tables: List[Dict[str, Any]] = []
            for rv in row_vars:
                valid_cols = [cv for cv in col_vars if cv != rv]
                blocks = [self._compute_block(rv, cv, weight_var, pct_mode) for cv in valid_cols]

                if blocks:
                    ordered_rows = list(blocks[0].get("row_cats") or [])
                    blocks = [self._reorder_block_rows_to_match(b, ordered_rows) for b in blocks]

                title = f"A.R {self.var_labels.get(rv, rv) or rv}"
                tables.append({
                    "row_var": rv,
                    "weight_var": weight_var,
                    "decimals": decimals,
                    "title": title,
                    "blocks": blocks,
                })

            out = {"ok": True, "tables": tables, "decimals": decimals}

        except Exception as e:
            out = {"ok": False, "error": str(e)}

        self.web_view.page().runJavaScript(
            f"window.__setPreview({json.dumps(json.dumps(out, ensure_ascii=False))});"
        )

    @Slot(str)
    def addToQueueFromPreview(self, table_json: str) -> None:
        try:
            t = json.loads(table_json) if table_json else {}
            if not isinstance(t, dict):
                raise RuntimeError("Tabela inválida.")

            title = str(t.get("title") or "").strip()
            if not title:
                title = "A.R"
            if not title.upper().startswith("A.R"):
                title = "A.R " + title

            item = {
                "title": title,
                "row_var": t.get("row_var"),
                "weight_var": t.get("weight_var"),
                "decimals": int(t.get("decimals") or 1),
                "blocks": t.get("blocks") or [],
            }

            self.queue.append(item)
            self._push_queue_state(message="Adicionado à fila.", reset_row=True)

        except Exception as e:
            self._push_queue_state(message=f"Erro: {e}")

    @Slot(int)
    def removeQueueItem(self, idx: int) -> None:
        try:
            if 0 <= idx < len(self.queue):
                self.queue.pop(idx)
            self._push_queue_state(message="Removido.")
        except Exception as e:
            self._push_queue_state(message=f"Erro: {e}")

    @Slot()
    def clearQueue(self) -> None:
        self.queue = []
        self._push_queue_state(message="Fila limpa.")

    def _export_tables(self, tables: List[Dict[str, Any]], display_map: Dict[str, str], out_path: str) -> None:
        wb = Workbook()
        wb.remove(wb.active)

        used_names = set()

        font_title = Font(name="DIN", size=15, bold=True, color="000000")
        font_head = Font(name="DIN", size=11, bold=True, color="000000")
        font_base = Font(name="DIN", size=10, bold=False, color="000000")
        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

        thin_black = Side(style="thin", color="000000")
        medium_black = Side(style="medium", color="000000")

        fill_head = PatternFill(fill_type="solid", fgColor="F2F2F2")
        fill_sub = PatternFill(fill_type="solid", fgColor="EAF3FF")

        def disp(v: str) -> str:
            v = str(v)
            return (display_map.get(v) or self.var_labels.get(v) or v)

        def apply_border(cell, left=None, right=None, top=None, bottom=None):
            cell.border = Border(
                left=left or thin_black,
                right=right or thin_black,
                top=top or thin_black,
                bottom=bottom or thin_black,
            )

        for t in tables:
            title = str(t.get("title") or "A.R").strip()
            if not title.upper().startswith("A.R"):
                title = "A.R " + title

            sheet_name = _unique_sheet_name(title, used_names)
            ws = wb.create_sheet(sheet_name)

            blocks = t.get("blocks") or []
            decimals = int(t.get("decimals") or 1)
            row_var = str(t.get("row_var") or "")
            row_labels = list((blocks[0].get("row_cats") if blocks else []) or [])

            ws.cell(1, 1, title)
            ws.cell(1, 1).font = font_title
            ws.cell(1, 1).alignment = align_left

            if not blocks:
                continue

            start_row = 3

            ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row + 1, end_column=1)
            c = ws.cell(start_row, 1, disp(row_var))
            c.font = font_head
            c.alignment = align_center
            c.fill = fill_head
            apply_border(c, left=medium_black, right=medium_black, top=medium_black, bottom=medium_black)

            col_cursor = 2
            total_data_cols = 0
            for b in blocks:
                col_cats = b.get("col_cats") or []
                span = max(1, len(col_cats))
                total_data_cols += span

                ws.merge_cells(start_row=start_row, start_column=col_cursor, end_row=start_row, end_column=col_cursor + span - 1)
                top_cell = ws.cell(start_row, col_cursor, disp(str(b.get("col_var") or "")))
                top_cell.font = font_head
                top_cell.alignment = align_center
                top_cell.fill = fill_sub

                for cc in range(col_cursor, col_cursor + span):
                    apply_border(
                        ws.cell(start_row, cc),
                        left=medium_black if cc == col_cursor else thin_black,
                        right=medium_black if cc == (col_cursor + span - 1) else thin_black,
                        top=medium_black,
                        bottom=medium_black
                    )

                for j, cc_label in enumerate(col_cats, start=col_cursor):
                    cell = ws.cell(start_row + 1, j, str(cc_label))
                    cell.font = font_head
                    cell.alignment = align_center
                    cell.fill = fill_head
                    apply_border(
                        cell,
                        left=medium_black if j == col_cursor else thin_black,
                        right=medium_black if j == (col_cursor + span - 1) else thin_black,
                        top=thin_black,
                        bottom=medium_black
                    )

                col_cursor += span

            row_cursor = start_row + 2

            for i, row_lab in enumerate(row_labels):
                tr_pct = row_cursor
                tr_z = row_cursor + 1

                ws.merge_cells(start_row=tr_pct, start_column=1, end_row=tr_z, end_column=1)
                c0 = ws.cell(tr_pct, 1, str(row_lab))
                c0.font = font_head
                c0.alignment = align_left
                c0.fill = fill_head
                apply_border(
                    c0,
                    left=medium_black,
                    right=medium_black,
                    top=medium_black if i == 0 else thin_black,
                    bottom=medium_black if i == len(row_labels) - 1 else thin_black
                )

                col_cursor = 2
                for b in blocks:
                    pct = b.get("percent") or []
                    z = b.get("z") or []
                    col_cats = b.get("col_cats") or []

                    start_block_col = col_cursor
                    end_block_col = col_cursor + len(col_cats) - 1

                    for j in range(len(col_cats)):
                        cellp = ws.cell(tr_pct, col_cursor)
                        cellz = ws.cell(tr_z, col_cursor)

                        cellp.alignment = align_center
                        cellp.font = font_base
                        cellz.alignment = align_center
                        cellz.font = font_base

                        apply_border(
                            cellp,
                            left=medium_black if col_cursor == start_block_col else thin_black,
                            right=medium_black if col_cursor == end_block_col else thin_black,
                            top=medium_black if i == 0 else thin_black,
                            bottom=thin_black
                        )
                        apply_border(
                            cellz,
                            left=medium_black if col_cursor == start_block_col else thin_black,
                            right=medium_black if col_cursor == end_block_col else thin_black,
                            top=thin_black,
                            bottom=medium_black if i == len(row_labels) - 1 else thin_black
                        )

                        try:
                            pv = float((pct[i] or [])[j])
                            cellp.value = pv
                            cellp.number_format = '0.0"%"'
                        except Exception:
                            cellp.value = None

                        try:
                            zv = float((z[i] or [])[j])
                            cellz.value = round(float(zv), decimals)
                            cellz.number_format = "0" if decimals <= 0 else ("0." + ("0" * decimals))

                            st = _cell_style_for_z(zv)
                            if st is not None:
                                fill_hex, font_hex = st
                                cellz.fill = PatternFill("solid", fgColor=fill_hex)
                                cellz.font = Font(name="DIN", size=10, bold=True, color=font_hex)

                                cellp.fill = PatternFill("solid", fgColor=fill_hex)
                                cellp.font = Font(name="DIN", size=10, bold=True, color=font_hex)
                        except Exception:
                            cellz.value = None

                        col_cursor += 1

                row_cursor += 2

            max_col = 1 + total_data_cols
            for c in range(1, max_col + 1):
                ws.column_dimensions[get_column_letter(c)].width = 16 if c == 1 else 12

            for r in range(1, ws.max_row + 1):
                ws.row_dimensions[r].height = 22

            ws.freeze_panes = "B5"

        wb.save(out_path)

    # ============================================================
    # DEBUG-ONLY START
    # ============================================================
    def _build_case_debug_for_pair(
        self,
        row_var: str,
        col_var: str,
        weight_var: str,
    ) -> pd.DataFrame:
        assert self.df is not None

        cols = [row_var, col_var, weight_var]
        df = self.df[cols].copy()

        df["_row_isna_orig"] = df[row_var].isna()
        df["_col_isna_orig"] = df[col_var].isna()
        df["_w_isna_orig"] = df[weight_var].isna()

        row_user_missing = self._user_missing_mask(df[row_var], row_var)
        col_user_missing = self._user_missing_mask(df[col_var], col_var)
        weight_user_missing = self._user_missing_mask(df[weight_var], weight_var)

        row_series, row_map = self._apply_value_labels_for_var(row_var, df[row_var])
        col_series, col_map = self._apply_value_labels_for_var(col_var, df[col_var])
        _ = row_series, col_series

        weight_clean = df[weight_var].mask(weight_user_missing)
        w_num = pd.to_numeric(weight_clean, errors="coerce")

        out = pd.DataFrame({
            "case_index": df.index,
            "row_code": df[row_var],
            "row_label": [_safe_label_from_map(row_map, v) for v in df[row_var]],
            "col_code": df[col_var],
            "col_label": [_safe_label_from_map(col_map, v) for v in df[col_var]],
            "weight_raw": df[weight_var],
            "weight_num": w_num,
            "row_isna_orig": df["_row_isna_orig"],
            "col_isna_orig": df["_col_isna_orig"],
            "weight_isna_orig": df["_w_isna_orig"],
            "row_is_user_missing": row_user_missing,
            "col_is_user_missing": col_user_missing,
            "weight_is_user_missing": weight_user_missing,
        })

        included = (
            (~out["row_isna_orig"])
            & (~out["col_isna_orig"])
            & (~out["weight_isna_orig"])
            & (~out["row_is_user_missing"])
            & (~out["col_is_user_missing"])
            & (~out["weight_is_user_missing"])
            & (pd.to_numeric(out["weight_num"], errors="coerce").fillna(0.0) > 0)
        )
        out["included"] = included

        reasons = []
        for _, r in out.iterrows():
            if bool(r["included"]):
                reasons.append("IN")
            elif bool(r["row_is_user_missing"]):
                reasons.append("OUT: row user-missing")
            elif bool(r["col_is_user_missing"]):
                reasons.append("OUT: col user-missing")
            elif bool(r["weight_is_user_missing"]):
                reasons.append("OUT: weight user-missing")
            elif bool(r["row_isna_orig"]):
                reasons.append("OUT: row NA")
            elif bool(r["col_isna_orig"]):
                reasons.append("OUT: col NA")
            elif bool(r["weight_isna_orig"]):
                reasons.append("OUT: weight NA")
            else:
                try:
                    wn = float(r["weight_num"])
                    if not np.isfinite(wn):
                        reasons.append("OUT: weight invalid")
                    elif wn <= 0:
                        reasons.append("OUT: weight <= 0")
                    else:
                        reasons.append("OUT: unknown")
                except Exception:
                    reasons.append("OUT: weight invalid")
        out["include_reason"] = reasons

        return out

    def _write_debug_matrix_sheet(self, ws, title: str, row_codes: List[str], col_codes: List[str], matrix: List[List[Any]], decimals: int) -> None:
        font_base = Font(name="DIN", size=10, bold=False, color="000000")
        font_head = Font(name="DIN", size=10, bold=True, color="000000")
        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

        thin = Side(style="thin", color="D9D9D9")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        num_fmt = "0." + ("0" * decimals) if decimals > 0 else "0"

        ws["A1"] = title
        ws["A1"].font = font_head
        ws["A1"].alignment = align_left

        ws.cell(3, 1, "ROW\\COL").font = font_head
        ws.cell(3, 1).alignment = align_center
        ws.cell(3, 1).border = border

        for j, cc in enumerate(col_codes, start=2):
            cell = ws.cell(3, j, cc)
            cell.font = font_head
            cell.alignment = align_center
            cell.border = border

        for i, rc in enumerate(row_codes, start=4):
            cell = ws.cell(i, 1, rc)
            cell.font = font_head
            cell.alignment = align_left
            cell.border = border

            row = matrix[i - 4] if (i - 4) < len(matrix) else []
            for j, v in enumerate(row, start=2):
                cell = ws.cell(i, j)
                cell.border = border
                cell.alignment = align_center
                cell.font = font_base

                try:
                    fv = float(v)
                    if np.isfinite(fv):
                        cell.value = fv
                        cell.number_format = num_fmt
                    else:
                        cell.value = None
                except Exception:
                    cell.value = None

        max_col = max(1, ws.max_column)
        max_row = max(1, ws.max_row)

        for c in range(1, max_col + 1):
            ws.column_dimensions[get_column_letter(c)].width = 12
        for r in range(1, max_row + 1):
            ws.row_dimensions[r].height = 22

        ws.freeze_panes = "B4"

    def _export_debug_tables(self, tables: List[Dict[str, Any]], debug_path: str) -> None:
        wb = Workbook()
        wb.remove(wb.active)

        sheet_idx = 1

        for t in tables:
            row_var = str(t.get("row_var") or "")
            blocks = t.get("blocks") or []

            for b in blocks:
                col_var = str(b.get("col_var") or "")

                row_codes = b.get("debug_row_codes") or []
                col_codes = b.get("debug_col_codes") or []
                row_disp = b.get("debug_row_display") or []
                col_disp = b.get("debug_col_display") or []
                obs = b.get("debug_obs") or []
                exp = b.get("debug_exp") or []
                z = b.get("z") or []

                tag = f"{sheet_idx:02d}"

                ws_map = wb.create_sheet(f"MAP {tag}")
                ws_map["A1"] = f"Bloco {tag}"
                ws_map["A2"] = f"Linha: {row_var}"
                ws_map["A3"] = f"Coluna: {col_var}"

                ws_map["A6"] = "ROW_CODE"
                ws_map["B6"] = "ROW_DISPLAY"
                for i, (rc, rd) in enumerate(zip(row_codes, row_disp), start=7):
                    ws_map.cell(i, 1, rc)
                    ws_map.cell(i, 2, rd)

                ws_map["D6"] = "COL_CODE"
                ws_map["E6"] = "COL_DISPLAY"
                for i, (cc, cd) in enumerate(zip(col_codes, col_disp), start=7):
                    ws_map.cell(i, 4, cc)
                    ws_map.cell(i, 5, cd)

                for c in range(1, 6):
                    ws_map.column_dimensions[get_column_letter(c)].width = 22

                ws_obs = wb.create_sheet(f"OBS {tag}")
                self._write_debug_matrix_sheet(
                    ws_obs,
                    f"Observado ponderado | {row_var} x {col_var}",
                    row_codes,
                    col_codes,
                    obs,
                    decimals=6
                )

                ws_exp = wb.create_sheet(f"EXP {tag}")
                self._write_debug_matrix_sheet(
                    ws_exp,
                    f"Esperado | {row_var} x {col_var}",
                    row_codes,
                    col_codes,
                    exp,
                    decimals=6
                )

                ws_z = wb.create_sheet(f"Z {tag}")
                self._write_debug_matrix_sheet(
                    ws_z,
                    f"Residual Ajustado | {row_var} x {col_var}",
                    row_codes,
                    col_codes,
                    z,
                    decimals=6
                )

                sheet_idx += 1

        wb.save(debug_path)

    def _write_case_debug_sheet(self, ws, title: str, df_cases: pd.DataFrame) -> None:
        font_base = Font(name="DIN", size=10, bold=False, color="000000")
        font_head = Font(name="DIN", size=10, bold=True, color="000000")
        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

        thin = Side(style="thin", color="D9D9D9")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        ws["A1"] = title
        ws["A1"].font = font_head
        ws["A1"].alignment = align_left

        headers = list(df_cases.columns)
        for j, h in enumerate(headers, start=1):
            cell = ws.cell(3, j, h)
            cell.font = font_head
            cell.alignment = align_center
            cell.border = border

        for i, row in enumerate(df_cases.itertuples(index=False), start=4):
            for j, v in enumerate(row, start=1):
                cell = ws.cell(i, j, v)
                cell.font = font_base
                cell.alignment = align_left if j <= 8 else align_center
                cell.border = border

        for c in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(c)].width = 16
        ws.freeze_panes = "A4"

    def _export_case_level_debug(self, tables: List[Dict[str, Any]], debug_cases_path: str) -> None:
        wb = Workbook()
        wb.remove(wb.active)

        sheet_idx = 1

        for t in tables:
            row_var = str(t.get("row_var") or "")
            weight_var = str(t.get("weight_var") or "")
            blocks = t.get("blocks") or []

            for b in blocks:
                col_var = str(b.get("col_var") or "")
                tag = f"{sheet_idx:02d}"

                df_cases = self._build_case_debug_for_pair(
                    row_var,
                    col_var,
                    weight_var,
                )

                ws_cases = wb.create_sheet(f"CASES {tag}")
                self._write_case_debug_sheet(
                    ws_cases,
                    f"Casos usados | {row_var} x {col_var} | peso={weight_var}",
                    df_cases
                )

                ws_group = wb.create_sheet(f"GROUP {tag}")
                grouped = (
                    df_cases[df_cases["included"] == True]
                    .groupby(["row_code", "row_label", "col_code", "col_label"], dropna=False)["weight_num"]
                    .sum()
                    .reset_index()
                    .rename(columns={"weight_num": "weighted_sum"})
                )
                self._write_case_debug_sheet(
                    ws_group,
                    f"Resumo agrupado | {row_var} x {col_var} | peso={weight_var}",
                    grouped
                )

                sheet_idx += 1

        wb.save(debug_cases_path)
    # ============================================================
    # DEBUG-ONLY END
    # ============================================================

    @Slot(str)
    def exportXlsx(self, payload_json: str) -> None:
        try:
            if self.df is None:
                raise RuntimeError("Nenhum arquivo carregado.")

            payload = json.loads(payload_json) if payload_json else {}
            mode = str(payload.get("mode") or "batch").strip().lower()
            display_map = payload.get("displayMap") or {}
            preview_tables = payload.get("previewTables") or []

            if not isinstance(display_map, dict):
                display_map = {}

            save_path, _ = QFileDialog.getSaveFileName(
                self.main_window,
                "Salvar XLSX",
                "analise_residual.xlsx",
                "Excel (*.xlsx)"
            )
            if not save_path:
                return
            if not save_path.lower().endswith(".xlsx"):
                save_path += ".xlsx"

            if mode == "queue":
                if not self.queue:
                    raise RuntimeError("Fila vazia.")
                tables = self.queue
            else:
                if not isinstance(preview_tables, list) or not preview_tables:
                    raise RuntimeError("No modo 1, gere a prévia antes de exportar.")
                tables = preview_tables

            self._export_tables(tables, display_map, save_path)

            debug_path = save_path[:-5] + "_DEBUG.xlsx"
            self._export_debug_tables(tables, debug_path)

            debug_cases_path = save_path[:-5] + "_CASES_DEBUG.xlsx"
            self._export_case_level_debug(tables, debug_cases_path)

            self._push_queue_state(
                message=(
                    "XLSX exportado com sucesso. "
                    f"Debugs salvos em: {os.path.basename(debug_path)} e {os.path.basename(debug_cases_path)}"
                )
            )

        except Exception as e:
            self._push_queue_state(message=f"Erro ao exportar: {e}")


class ResidualsPage(QWidget):
    def __init__(self, main_window: QMainWindow, profile: Optional[QWebEngineProfile] = None):
        super().__init__()
        self.main_window = main_window

        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(0)

        self.topbar = TopBar(
            title="Funcionalidade: Análise Residual (A.R)",
            icon_path=resource_path("assets/icons/table.ico"),
            back_icon_path=resource_path("assets/icons/arrow-left.ico"),
            show_help=True
        )
        self.btn_back = self.topbar.btn_back
        layout.addWidget(self.topbar)

        self.view = QWebEngineView(self)
        layout.addWidget(self.view, 1)

        if profile is not None:
            page = QWebEnginePage(profile, self.view)
            self.view.setPage(page)

        self.view.setHtml(HTML, QUrl("https://residuals.app.local/"))

        self.bridge = ResidualsBridge(main_window, self.view)
        channel = QWebChannel(self.view.page())
        channel.registerObject("QtBridge", self.bridge)
        self.view.page().setWebChannel(channel)

        self.view.loadFinished.connect(lambda ok: self.bridge._push_queue_state())

    def cleanup(self):
        pass