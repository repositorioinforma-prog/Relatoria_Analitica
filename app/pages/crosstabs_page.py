from app.core.resources import resource_base, resource_path
import json
import math
import re
import unicodedata
from decimal import Decimal, ROUND_HALF_UP
from typing import Any, Dict, List, Optional, Tuple
from pathlib import Path
import sys

import pandas as pd

from PySide6.QtCore import QUrl, QObject, Signal, Slot
from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QFileDialog, QMessageBox
)
from PySide6.QtWebEngineWidgets import QWebEngineView
from PySide6.QtWebEngineCore import QWebEnginePage, QWebEngineProfile
from PySide6.QtWebChannel import QWebChannel

from app.ui.topbar import TopBar

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ============================================================
# CONFIG DE ORDENAÇÃO AUTOMÁTICA
# ============================================================
FIXED_BOTTOM_LABELS = [
    "Não sei",
    "Branco/Nulo",
    "Branco / Nulo",
    "Nulo / Branco",
    "Não votarei",
    "Não irá votar",
    "Não sabe",
    "NS",
]

# Ordem fixa para escalas específicas.
# Se a tabela contiver pelo menos 2 itens dessa sequência,
# a ordem abaixo prevalece sobre a ordenação por Total.
PRIORITY_ROW_SEQUENCES = [
    [
        "Certamente voto",
        "Provavelmente voto",
        "Provavelmente não voto",
        "Certamente não voto",
        "Não conheço suficiente para votar",
        "Não sabe",
        "NS",
    ],
]

# Itens herdados do VBA que não entram na ordenação por Total.
# Eles ficam preservados e vão para o fim, mantendo a ordem original entre si.
VBA_EXCLUDE_LABELS = [
    "Base",
    "Base reduzida",
    "Outros",
    "Direita",
    "Centro",
    "Esquerda",
    "Branco/Nulo",
    "Votaria em todos",
    "Nenhum",
    "Aumenta",
    "Diminui",
    "Não sabe",
    "Concordo totalmente",
    "Nem concordo nem discordo",
    "Discordo totalmente",
    "Um candidato opositor a Família Reis ",
    "Um candidato da Família Reis ",
    "Um candidato aliado aos Reis, mas que não seja da família; ",
    "Concordo plenamente",
    "Concordo",
    "Discordo",
    "Discordo plenamente",
    "Nenhuma  palavra",
    "É conhecedor de Brasília e das prioridades das pessoas ",
    "Não conhece a cidade e o que as pessoas precisam ",
    "Tem vontade de trazer melhorias para Brasília",
    "Não tem vontade de trazer melhorias para Brasília",
    "É conhecedor de Brasília e das prioridades das pessoas",
    "Não conhece a cidade e o que as pessoas precisam",
    "Está preparado para exercer o cargo",
    "Não está preparado para exercer o cargo",
    "Alta associação",
    "Associação moderada",
    "Baixa associação",
    "Diariamente",
    "Algumas vezes por semana",
    "Algumas vezes no mês",
    "Raramente",
    "Nunca",
    "Grande interesse",
    "Interesse moderado",
    "Baixo interesse",
    "Certamente voto",
    "Certamente não voto",
    "Não conheço suficiente para votar",
    "Provavelmente voto",
    "Provavelmente não voto",
    "Todas as alternativas",
    "Nenhum desses",
    "Branco / Nulo",
    "Ajuda muito na decisão",
    "Ajuda na decisão",
    "Atrapalha na decisão",
    "Atrapalha muito na decisão",
    "Preparado",
    "Despreparado",
    "Experiente",
    "Inexperiente",
    "Tem influência política",
    "Não tem influência política",
    "Conhece a cidade",
    "Não sei",
    "Não votarei",
    "Regular",
    "Aprovação",
    "Reprovação",
    "Nenhum deles",
    "Melhorou",
    "Piorou",
    "Continua da mesma forma",
    "Nada a melhorar",
    "Vai melhorar",
    "Ficar como está",
    "Vai piorar",
    "Confia muito",
    "Confia",
    "Confia mais ou menos",
    "Confia pouco",
    "Não confia",
    "Trocaria",
    "Não trocaria",
    "Outras cidades",
    "Não faz nada",
    "Nada contra a gestão atual",
    "Votou branco/nulo",
    "Não irá votar",
    "Está preparado",
    "Não está preparado",
    "Tem capacidade",
    "Não tem capacidade",
    "Tem força política",
    "Não tem força política",
    "Tem autonomia",
    "Não tem autonomia",
    "Ajudaria",
    "Não ajudaria",
    "Demais, deveria diminuir",
    "Na medida certa",
    "Pouca, deveria aumentar",
    "Lembra",
    "Não lembra",
    "Desde que nasceu",
    "Mora há mais de 20 anos",
    "Mora entre 20 e 10 anos",
    "Mora entre 10 e 5 anos",
    "Mora há menos de 5 anos",
    "Irá cumprir todas as promessas que faz",
    "Irá cumprir somente algumas",
    "Não irá cumprir nenhuma promessa que faz",
    "Promete muito mais do que conseguirá fazer em Maricá",
    "Promete na medida certa",
    "Promete pouco para Maricá, pois conseguiria fazer mais",
    "Sim",
    "Não",
    "Indiferente",
    "Outros aspectos",
    "Nada de bom",
    "Não vê nada de ruim",
    "Não sente falta de nada",
    "Muito interesse",
    "Interesse regular",
    "Nenhum interesse",
    "Excelente",
    "Boa",
    "Ruim",
    "Péssima",
    "Votou Branco ou Nulo",
    "Repro- vação",
    "Média",
    "Otimista",
    "Pessimista",
    "Nada importante",
    "Pouco importante",
    "Importante",
    "Muito importante",
    "NS",
    "Evoluindo",
    "Mesma coisa / parada",
    "Regredindo",
    "Da esquerda",
    "Da direita",
    "Do centro",
    "Conheço bem",
    "Conheço de ouvir falar",
    "Conheço das redes sociais",
    "Não conheço",
    "Indeciso",
    "Aumenta minha chance de votar nesse candidato",
    "É indiferente o apoio",
    "Diminui a minha chance de votar nesse candidato",
    "Não lembra/ Não tem candidato",
    "Não sabe em quem votar",
    "Outros candidatos",
    "Nulo / Branco",
    "Outras áreas",
    "Nenhuma área",
    "Todas as áreas",
    "Outras necessidades",
    "Nenhuma necessidade",
    "Conhece um pouco",
    "Nunca ouvi falar",
    "Conhece muito",
    "Influencia meu voto no candidato que apoiar",
    "Posso considerar o voto no candidato apoiado",
    "Posso considerar não votar no candidato apoiado",
    "Influencia a não votar no candidato apoiado",
    "Conheço muito",
    "Conheço um pouco",
    "Conheço somente de foto",
    "Conheço de foto, mas, não sabia que era Rodrigo Bacellar",
    "Nunca ouvi falar e nunca vi",
    "Tenho informações positivas",
    "Tenho informações negativas",
    "Não tenho informações",
    "Já viu",
    "Nunca viu",
    "Nenhum dos dois governos está bem",
    "Concorda totalmente",
    "Concorda em parte",
    "Não concorda e nem discorda",
    "Discorda em parte",
    "Discorda totalmente",
    "Melhor",
    "Igual",
    "Pior",
    "Todos devem permanecer presos e ser condenados",
    "A prisão deve ser apenas para quem danificou prédios",
    "Todos devem receber anistia",
    "Maricá oferece vagas de emprego até demais para os moradores",
    "Maricá oferece emprego na medida certa, todos poderiam estar empregados",
    "Maricá oferece poucas vagas de emprego, e vai oferecer cada vez menos",
    "Maricá oferece poucas vagas de emprego, mas, começará a oferecer mais",
    "O turismo está crescendo, mas falta gerar emprego e renda aos moradores",
    "O turismo está crescendo e tem ajudado a gerar emprego e renda aos moradores",
    "O turismo está estagnado",
    "O turismo está diminuindo",
    "Maricá está crescendo, mas sem planejamento e ficará pior",
    "Maricá está crescendo, sem planejamento, mas haverá uma organização",
    "Maricá já está crescendo de forma organizada",
    "Maricá está estagnada",
    "Maricá está regredindo",
]


HTML = r"""<!doctype html>
<html lang="pt-br">
<head>
  <meta charset="utf-8"/>
  <meta name="viewport" content="width=device-width,initial-scale=1"/>
  <title>Cruzamentos</title>
  <style>
    html,body{height:100%;margin:0;font-family:system-ui,Arial,sans-serif;background:#fff;}
    .app{display:grid;grid-template-columns:460px 1fr;height:100%}
    .panel{border-right:1px solid #eee;padding:14px 16px 16px;overflow:auto;background:#fff}
    h1{font-size:18px;margin:0 0 12px}
    h2{font-size:14px;margin:16px 0 8px}
    label{font-size:13px;display:block;margin:8px 0 6px}
    small{color:#666;font-size:12px}
    button,input[type="number"],textarea,input[type="text"]{
      font:inherit;padding:8px 10px;border-radius:10px;border:1px solid #ddd;background:#fff
    }
    button{cursor:pointer}
    .btn-primary{background:#1E90FF;border-color:#1E90FF;color:#fff}
    .btn-outline{background:#fff;color:#111;border-color:#ccc}
    .btn-sm{padding:6px 8px;border-radius:10px;font-size:12px}
    .inline{display:flex;gap:8px;align-items:center;flex-wrap:wrap}
    .muted{color:#666;font-size:12px}
    .box{border:1px solid #eee;border-radius:14px;padding:10px 10px 8px;background:#fafafa;margin-bottom:10px;}
    .grid2{display:grid;grid-template-columns:1fr 1fr;gap:10px}
    .grid3{display:grid;grid-template-columns:1fr 1fr 1fr;gap:10px}
    .stack{display:grid;gap:10px}
    .actionsRow{display:flex;gap:8px;align-items:center;justify-content:flex-end;flex-wrap:wrap}
    @media (max-width: 980px){
      .grid2,.grid3{grid-template-columns:1fr;}
      .actionsRow{justify-content:flex-start}
    }

    .list{
      border:1px solid #ddd;border-radius:12px;background:#fff;max-height:220px;overflow:auto;padding:6px;
    }
    .item{display:flex;gap:8px;align-items:center;padding:6px;border-radius:10px}
    .item:hover{background:#f4f6fb}
    .item input{width:auto}
    .pill{padding:4px 8px;border-radius:999px;background:#fff;border:1px solid #e6e6e6;font-weight:700;cursor:help;white-space:nowrap}
    .sep{height:1px;background:#eee;margin:10px 0}

    .content{padding:14px 16px;background:#fff;overflow:auto}
    .card{border:1px solid #eee;border-radius:14px;padding:12px;background:#fff;margin-bottom:12px}
    .cardHead{display:flex;gap:10px;justify-content:space-between;align-items:center;flex-wrap:wrap}
    .cardTitleWrap{display:flex;gap:8px;align-items:center;flex-wrap:wrap;min-width:340px;flex:1}
    .cardTitleLabel{font-weight:900;font-size:13px;white-space:nowrap}
    .titleInput{
      flex:1; min-width:240px;
      padding:8px 10px;border-radius:10px;border:1px solid #ddd;
      font:inherit;
    }
    .tableWrap{overflow:auto;border:1px solid #eee;border-radius:12px;margin-top:10px}
    table{border-collapse:collapse;min-width:100%;background:#fff}
    th,td{border:1px solid #eee;padding:10px 10px;text-align:center;vertical-align:middle;font-size:12px;white-space:nowrap}
    th{background:#f6f6f6;font-weight:900}
    td.rowlab,th.rowlab{text-align:left;font-weight:900;background:#fafafa;position:sticky;left:0;z-index:2}
    th.rowlab{z-index:3}
    tr.baseRow td, tr.baseRow th{background:#fafafa;font-weight:900}
    .num{font-variant-numeric: tabular-nums;}

    input:focus, textarea:focus, button:focus{outline:none !important;box-shadow:none !important;}

    .dd{position:relative;}
    .dd-btn{
      width:100%;text-align:left;padding:10px 10px;border-radius:12px;border:1px solid #ddd;background:#fff;
      cursor:pointer;font:inherit;display:flex;align-items:center;justify-content:space-between;gap:10px;
    }
    .dd-btn:disabled{opacity:.6;cursor:not-allowed}
    .dd-arrow{color:#666;font-size:12px}
    .dd-menu{
      position:absolute; left:0; right:0; top:calc(100% + 6px);
      background:#fff;border:1px solid #ddd;border-radius:12px;max-height:280px;overflow:auto;
      box-shadow:0 12px 28px rgba(0,0,0,.10);z-index:2000;
    }
    .dd-item{padding:10px 10px;cursor:pointer;font-size:13px;color:#111;user-select:none;}
    .dd-item:hover{background:#f2f2f2}
    .dd-item.is-selected{background:#1E90FF; color:#fff; font-weight:900;}

    .combo{position:relative;}
    .combo input{
      width:100%;padding:10px 10px;border-radius:12px;border:1px solid #ddd;background:#fff;font:inherit;
    }
    .combo-menu{
      position:absolute; left:0; right:0; top:calc(100% + 6px);
      background:#fff;border:1px solid #ddd;border-radius:12px;max-height:280px;overflow:auto;
      box-shadow:0 12px 28px rgba(0,0,0,.10);z-index:2100;
    }
    .combo-item{padding:10px 10px;cursor:pointer;font-size:13px;}
    .combo-item:hover{background:#f2f2f2}
    .combo-item.is-selected{background:#1E90FF;color:#fff;font-weight:900}

    .modal-backdrop{
      position:fixed; inset:0; background:rgba(0,0,0,.35);
      display:none; align-items:center; justify-content:center; z-index:3000;
    }
    .modal{
      width:min(780px, 92vw);
      max-height:86vh;
      background:#fff; border-radius:16px; border:1px solid #eee;
      box-shadow:0 20px 60px rgba(0,0,0,.25);
      padding:14px;

      display:flex;
      flex-direction:column;
      gap:10px;
    }
    .modal h3{margin:0;font-size:14px}
    .modal .rows{
      border:1px solid #eee;border-radius:12px;
      overflow:auto;
      flex:1;
      min-height:120px;
    }
    .modal .r{
      display:flex; align-items:center; justify-content:space-between;
      gap:10px; padding:10px 10px; border-bottom:1px solid #f0f0f0;
      background:#fff;
    }
    .modal .r:last-child{border-bottom:none}
    .modal .name{font-weight:800;font-size:13px;flex:1;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}
    .modal .meta{color:#666;font-size:12px;white-space:nowrap}
    .modal .actions{display:flex;gap:6px;flex-wrap:nowrap}
    .modal .actions button{padding:6px 8px;border-radius:10px}
    .modal .footer{display:flex;justify-content:flex-end;gap:8px}
    .is-removed{opacity:.45;text-decoration:line-through}
    .hintRow{
      padding:10px 10px;
      font-size:12px;
      color:#666;
      border-top:1px solid #f0f0f0;
      background:#fff;
    }
  </style>
</head>
<body>
<div class="app">
  <aside class="panel">
    <h1>Cruzamentos</h1>

    <div class="box">
      <h2 style="margin:0 0 8px">0) Dados</h2>
      <label>Importar Banco com peso do SPSS</label>
      <div class="inline">
        <button id="btnPickData" class="btn-outline">Carregar dados (CSV/XLSX/SAV)</button>
        <small id="dataStatus" class="muted">Nenhum arquivo carregado.</small>
      </div>

      <div class="sep"></div>

      <h2 style="margin:0 0 8px">1) Configurações das Variáveis</h2>
      <div class="stack">
        <div>
          <label>Variável (Linhas)</label>
          <div class="combo" id="rowCombo"></div>
          <small class="muted">Você pode rolar a lista ou digitar para buscar</small>
        </div>
        <div class="grid2">
          <div>
            <label>Peso</label>
            <div class="dd" id="weightVarDD" data-placeholder="(Sem peso)"></div>
            <small class="muted">Mostra somente colunas que contêm “peso”</small>
          </div>
          <div>
            <label>Percentual</label>
            <div class="dd" id="pctModeDD" data-placeholder="Coluna (%)"></div>
          </div>
        </div>
      </div>

      <div style="margin-top:10px">
        <label>Buscar variável (Colunas)</label>
        <input id="colSearch" type="text" placeholder="Digite para filtrar..." />
        <div class="actionsRow" style="margin-top:8px">
          <button id="btnAll" class="btn-outline btn-sm" type="button">Marcar tudo</button>
          <button id="btnNone" class="btn-outline btn-sm" type="button">Desmarcar</button>
        </div>
      </div>

      <label>Variáveis (Colunas)</label>
      <div id="colList" class="list"></div>

      <div class="inline" style="margin-top:10px">
        <span class="pill"><input id="includeNA" type="checkbox" style="width:auto;margin:0"> Incluir NA</span>
      </div>

      <div class="inline" style="margin-top:12px">
        <button id="btnRun" class="btn-primary" type="button">Gerar cruzamentos</button>
        <button id="btnClear" class="btn-outline" type="button">Limpar resultados</button>
        <small id="runStatus" class="muted"></small>
      </div>
    </div>

    <div class="box">
      <h2 style="margin:0 0 8px">2) Exportação para Excel</h2>
      <div class="inline" style="margin-bottom:8px">
        <button id="btnAddAllQueue" class="btn-outline" type="button">Adicionar todos</button>
        <button id="btnClearQueue" class="btn-outline" type="button">Limpar pilha</button>
        <button id="btnExportQueue" class="btn-primary" type="button">Exportar Excel</button>
      </div>
      <small id="queueStatus" class="muted">Fila: 0 tabela(s)</small>
      <div id="queueList" class="muted" style="margin-top:10px"></div>
    </div>

    <small class="muted">
      Dica: gere os cruzamentos, edite título/ordem se quiser, e adicione à pilha para exportar tudo de uma vez.
    </small>
  </aside>

  <main class="content">
    <div id="results" class="muted">Nenhum resultado ainda.</div>
  </main>
</div>

<div class="modal-backdrop" id="modalBackdrop">
  <div class="modal">
    <h3 id="modalTitle">Ordenar</h3>
    <div class="rows" id="modalRows"></div>
    <div class="footer">
      <button class="btn-outline" id="btnModalCancel" type="button">Cancelar</button>
      <button class="btn-primary" id="btnModalSave" type="button">Salvar</button>
    </div>
  </div>
</div>

<script src="qrc:///qtwebchannel/qwebchannel.js"></script>
<script>
  let QtBridge = null;
  let __columns = [];
  let __results = [];
  let __queue = [];
  let __varLabels = {};
  let __selectedCols = new Set();
  let __colSearchTimer = null;
  let __lastColFilter = "";

  function esc(s){
    return String(s ?? "").replace(/[&<>"']/g, m => ({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[m]));
  }
  function el(id){ return document.getElementById(id); }
  function norm(s){ return String(s||"").normalize('NFD').replace(/[\u0300-\u036f]/g,'').toLowerCase(); }
  function canon(s){ return norm(s).replace(/[^a-z0-9]+/g,' ').trim(); }
  function isPeso(name){ return canon(name).includes("peso"); }
  function dispVar(name){
    return (__varLabels && __varLabels[name]) ? __varLabels[name] : name;
  }

  function Dropdown(rootEl){
    const placeholder = rootEl.dataset.placeholder || '—';
    let items = [];
    let value = '';
    let enabled = true;

    const btn = document.createElement('button');
    btn.type = 'button';
    btn.className = 'dd-btn';
    btn.innerHTML = `<span class="dd-text">${placeholder}</span><span class="dd-arrow">▼</span>`;

    const menu = document.createElement('div');
    menu.className = 'dd-menu';
    menu.hidden = true;

    rootEl.appendChild(btn);
    rootEl.appendChild(menu);

    function setText(t){ btn.querySelector('.dd-text').textContent = t || placeholder; }
    function close(){ menu.hidden = true; }
    function open(){ if(enabled && items.length) menu.hidden = false; }

    function render(){
      menu.innerHTML = '';
      items.forEach(it=>{
        const div = document.createElement('div');
        div.className = 'dd-item' + (it.value===value ? ' is-selected' : '');
        div.textContent = it.label;
        div.addEventListener('click', ()=>{
          value = it.value;
          setText(it.label);
          render();
          close();
          api.onChange && api.onChange(value);
        });
        menu.appendChild(div);
      });
    }

    btn.addEventListener('click', ()=> menu.hidden ? open() : close());
    document.addEventListener('click', (e)=>{ if(!rootEl.contains(e.target)) close(); });

    const api = {
      onChange: null,
      setItems(arr){
        items = Array.isArray(arr) ? arr.slice() : [];
        value = '';
        setText('');
        render();
        close();
      },
      setValue(v){
        value = v || '';
        const found = items.find(x => x.value === value);
        setText(found ? found.label : '');
        render();
        close();
      },
      getValue(){ return value; },
      setEnabled(v){
        enabled = !!v;
        btn.disabled = !enabled;
        if(!enabled){ value=''; setText(''); close(); }
      }
    };
    return api;
  }

  function ComboSearch(rootEl){
    let items = [];
    let value = '';

    const input = document.createElement('input');
    input.type = 'text';
    input.placeholder = 'Digite para buscar...';
    rootEl.appendChild(input);

    const menu = document.createElement('div');
    menu.className = 'combo-menu';
    menu.hidden = true;
    rootEl.appendChild(menu);

    function close(){ menu.hidden = true; }
    function open(){ if(items.length) menu.hidden = false; }

    function render(){
      const q = canon(input.value || '');
      menu.innerHTML = '';
      const filtered = items.filter(it => !q || canon(it.label).includes(q));
      filtered.forEach(it=>{
        const div = document.createElement('div');
        div.className = 'combo-item' + (it.value===value ? ' is-selected' : '');
        div.textContent = it.label;
        div.addEventListener('click', ()=>{
          value = it.value;
          input.value = it.label;
          render();
          close();
          api.onChange && api.onChange(value);
        });
        menu.appendChild(div);
      });
    }

    input.addEventListener('focus', ()=>{ open(); render(); });
    input.addEventListener('click', ()=>{ open(); render(); });
    input.addEventListener('input', ()=>{ open(); render(); });
    document.addEventListener('click', (e)=>{ if(!rootEl.contains(e.target)) close(); });

    const api = {
      onChange: null,
      setItems(arr){
        items = Array.isArray(arr) ? arr.slice() : [];
        value = '';
        input.value = '';
        render();
        close();
      },
      setValue(v){
        value = v || '';
        const found = items.find(x => x.value === value);
        input.value = found ? found.label : '';
        render();
      },
      getValue(){ return value; },
    };
    return api;
  }

  let ddWeight = null;
  let ddPct = null;
  let comboRow = null;

  function filteredColumns(filterText=""){
    const ft = canon(filterText || "");
    return __columns.filter(c => {
      const label = dispVar(c.name);
      return !ft || canon(label).includes(ft) || canon(c.name).includes(ft);
    });
  }

  function renderColList(filterText=""){
    const host = el("colList");
    const frag = document.createDocumentFragment();
    __lastColFilter = filterText || "";
    host.innerHTML = "";

    for(const c of filteredColumns(__lastColFilter)){
      const label = dispVar(c.name);
      const row = document.createElement("div");
      row.className = "item";
      row.innerHTML = `<input type="checkbox" data-var="${esc(c.name)}" ${__selectedCols.has(c.name) ? "checked" : ""}>
                       <div style="flex:1">
                         <div style="font-weight:900;font-size:12px">${esc(label)}</div>
                         <div class="muted">${esc(c.name)} · ${esc(c.dtype)}</div>
                       </div>`;
      frag.appendChild(row);
    }

    host.appendChild(frag);
  }

  function scheduleRenderColList(filterText=""){
    if(__colSearchTimer) clearTimeout(__colSearchTimer);
    __colSearchTimer = setTimeout(()=> renderColList(filterText || ""), 140);
  }

  function markAll(val){
    for(const c of filteredColumns(__lastColFilter)){
      if(val) __selectedCols.add(c.name);
      else __selectedCols.delete(c.name);
    }
    el("colList").querySelectorAll('input[type="checkbox"]').forEach(cb => {
      cb.checked = val;
    });
  }

  function selectedCols(){
    return Array.from(__selectedCols);
  }

  function fmtPct(x){
    if(x == null || Number.isNaN(x)) return "";
    return Number(x).toFixed(1).replace('.', ',');
  }

  function buildTableHTML(t){
    let html = `<div class="card" data-key="${esc(t.key)}">
      <div class="cardHead">
        <div class="cardTitleWrap">
          <div class="cardTitleLabel">${esc(t.row_var_display || t.row_var)} × ${esc(t.col_var_display || t.col_var)}</div>
          <input class="titleInput" type="text" value="${esc(t.title)}" data-action="editTitle" />
        </div>
        <div class="inline">
          <span class="pill">${esc(t.base_display)}</span>
          <button class="btn-outline btn-sm" data-action="orderRows" type="button">Ordenar linhas</button>
          <button class="btn-outline btn-sm" data-action="orderCols" type="button">Ordenar colunas</button>
          <button class="btn-outline btn-sm" data-action="addQueue" type="button"
            style="background:#1E90FF; border-color:#1E90FF; color:#fff;">
            Adicionar à pilha
          </button>
          <button class="btn-outline btn-sm" data-action="exportOne" type="button">Exportar (1)</button>
        </div>
      </div>

      <div class="tableWrap">
      <table>
        <thead>
          <tr>
            <th class="rowlab">Categoria</th>
            <th>Total</th>`;

    for(const c of t.col_labels) html += `<th>${esc(c)}</th>`;
    html += `</tr></thead><tbody>`;

    for(let i=0;i<t.row_labels.length;i++){
      html += `<tr>
        <td class="rowlab">${esc(t.row_labels[i])}</td>
        <td class="num">${fmtPct(t.pct_total[i])}</td>`;

      for(let j=0;j<t.col_labels.length;j++){
        html += `<td class="num">${fmtPct(t.pct_matrix[i][j])}</td>`;
      }
      html += `</tr>`;
    }

    html += `<tr class="baseRow">
      <td class="rowlab">Base</td>
      <td class="num">${esc(t.base_total_n)}</td>`;
    for(let j=0;j<t.col_labels.length;j++){
      html += `<td class="num">${esc(t.base_col_n[j])}</td>`;
    }
    html += `</tr>`;

    html += `</tbody></table></div></div>`;
    return html;
  }

  function renderResults(){
    const host = el("results");
    if(!__results.length){
      host.textContent = "Nenhum resultado ainda.";
      return;
    }
    let html = "";
    for(const t of __results) html += buildTableHTML(t);
    host.innerHTML = html;
  }

  function renderSingleResult(tableKey){
    const t = __results.find(x => x.key === tableKey);
    if(!t) return;

    const host = el("results");
    const current = host.querySelector(`.card[data-key="${CSS.escape(tableKey)}"]`);
    if(!current){
      renderResults();
      return;
    }

    const wrap = document.createElement("div");
    wrap.innerHTML = buildTableHTML(t);
    const nextCard = wrap.firstElementChild;
    if(!nextCard){
      renderResults();
      return;
    }
    current.replaceWith(nextCard);
  }

  function renderQueueState(payloadJson){
    const payload = JSON.parse(payloadJson);
    __queue = payload.items || [];
    el("queueStatus").textContent = `Fila: ${__queue.length} tabela(s)`;

    const host = el("queueList");
    host.innerHTML = "";

    __queue.forEach((it, idx)=>{
      const row = document.createElement("div");
      row.className = "inline";
      row.style.justifyContent = "space-between";
      row.style.gap = "10px";
      row.style.padding = "6px 0";
      row.style.borderBottom = "1px dashed #eee";

      const txt = document.createElement("div");
      txt.textContent = `${idx+1}. ${it.title}`;
      txt.style.flex = "1";
      txt.style.whiteSpace = "nowrap";
      txt.style.overflow = "hidden";
      txt.style.textOverflow = "ellipsis";

      const btn = document.createElement("button");
      btn.className = "btn-outline btn-sm";
      btn.type = "button";
      btn.textContent = "Remover";
      btn.dataset.action = "removeQueue";
      btn.dataset.index = String(idx);

      row.appendChild(txt);
      row.appendChild(btn);
      host.appendChild(row);
    });
  }

  // ===== MODAL =====
  let __modalMode = null;
  let __modalTableKey = null;
  let __modalRowsArr = [];
  let __modalCols = [];

  function openModal(){ el("modalBackdrop").style.display = "flex"; }
  function closeModal(){
    el("modalBackdrop").style.display = "none";
    __modalMode = null; __modalTableKey = null;
    __modalRowsArr = []; __modalCols = [];
  }

  function applyOrderToTableRows(tableKey, newRowLabels){
    const t = __results.find(x => x.key === tableKey);
    if(!t) return;

    const old = t.row_labels.slice();
    const idxMap = new Map(old.map((lab, i)=>[lab, i]));

    const newPctTotal = [];
    const newPctMatrix = [];
    newRowLabels.forEach(lab=>{
      const i = idxMap.get(lab);
      if(i == null) return;
      newPctTotal.push(t.pct_total[i]);
      newPctMatrix.push(t.pct_matrix[i]);
    });

    t.row_labels = newRowLabels.slice();
    t.pct_total = newPctTotal;
    t.pct_matrix = newPctMatrix;

    renderSingleResult(tableKey);
  }

  function applyOrderToTableCols(tableKey, activeIdxs){
    const t = __results.find(x => x.key === tableKey);
    if(!t) return;

    const newColLabels = activeIdxs.map(i => t.col_labels[i]);
    const newBaseColN  = activeIdxs.map(i => t.base_col_n[i]);
    const newPctMatrix = t.pct_matrix.map(row => activeIdxs.map(i => row[i]));

    t.col_labels = newColLabels;
    t.base_col_n = newBaseColN;
    t.pct_matrix = newPctMatrix;

    renderSingleResult(tableKey);
  }

  function openOrderRowsModal(table){
    __modalMode = "rows";
    __modalTableKey = table.key;
    el("modalTitle").textContent = `Ordenar linhas: ${table.row_var_display || table.row_var} × ${table.col_var_display || table.col_var}`;

    const totals = table.pct_total || [];
    __modalRowsArr = (table.row_labels || []).map((lab, i)=>({
      label: lab,
      total: (totals[i] != null) ? Number(totals[i]) : null
    }));

    const hint = (table.fixed_bottom_labels || []).join(" → ");

    function redraw(){
      const host = el("modalRows");
      host.innerHTML = "";

      __modalRowsArr.forEach((obj, i)=>{
        const r = document.createElement("div");
        r.className = "r";
        const meta = (obj.total == null) ? "" : `Total: ${fmtPct(obj.total)}`;
        r.innerHTML = `
          <div class="name">${esc(obj.label)}</div>
          <div class="meta">${esc(meta)}</div>
          <div class="actions">
            <button class="btn-outline btn-sm" data-act="up" data-i="${i}">↑</button>
            <button class="btn-outline btn-sm" data-act="down" data-i="${i}">↓</button>
          </div>
        `;
        host.appendChild(r);
      });

      if(hint){
        const h = document.createElement("div");
        h.className = "hintRow";
        h.textContent = `Sugestão: itens comuns no final: ${hint} (você pode mover se quiser)`;
        host.appendChild(h);
      }
    }

    el("modalRows").onclick = (ev)=>{
      const btn = ev.target.closest("button[data-act]");
      if(!btn) return;
      const act = btn.dataset.act;
      const i = parseInt(btn.dataset.i, 10);
      if(!Number.isFinite(i)) return;

      if(act === "up" && i > 0){
        const tmp = __modalRowsArr[i-1];
        __modalRowsArr[i-1] = __modalRowsArr[i];
        __modalRowsArr[i] = tmp;
        redraw();
      }
      if(act === "down" && i < __modalRowsArr.length-1){
        const tmp = __modalRowsArr[i+1];
        __modalRowsArr[i+1] = __modalRowsArr[i];
        __modalRowsArr[i] = tmp;
        redraw();
      }
    };

    redraw();
    openModal();
  }

  function openOrderColsModal(table){
    __modalMode = "cols";
    __modalTableKey = table.key;
    el("modalTitle").textContent = `Ordenar/Remover colunas: ${table.row_var_display || table.row_var} × ${table.col_var_display || table.col_var}`;

    __modalCols = (table.col_labels || []).map((lab, idx)=>({
      label: lab,
      base_n: (table.base_col_n || [])[idx],
      removed: false,
      origIndex: idx,
    }));

    function redraw(){
      const host = el("modalRows");
      host.innerHTML = "";

      const lock = document.createElement("div");
      lock.className = "r";
      lock.innerHTML = `<div class="name">Total</div><div class="meta">Fixo na coluna B</div>`;
      host.appendChild(lock);

      __modalCols.forEach((col, i)=>{
        const r = document.createElement("div");
        r.className = "r" + (col.removed ? " is-removed" : "");
        const baseTxt = (col.base_n != null) ? `Base: ${col.base_n}` : "";
        r.innerHTML = `
          <div class="name">${esc(col.label)}</div>
          <div class="meta">${esc(baseTxt)}</div>
          <div class="actions">
            <button class="btn-outline btn-sm" data-act="up" data-i="${i}" ${col.removed ? "disabled" : ""}>↑</button>
            <button class="btn-outline btn-sm" data-act="down" data-i="${i}" ${col.removed ? "disabled" : ""}>↓</button>
            <button class="btn-outline btn-sm" data-act="toggle" data-i="${i}">
              ${col.removed ? "Restaurar" : "Remover"}
            </button>
          </div>
        `;
        host.appendChild(r);
      });
    }

    el("modalRows").onclick = (ev)=>{
      const btn = ev.target.closest("button[data-act]");
      if(!btn) return;
      const act = btn.dataset.act;
      const i = parseInt(btn.dataset.i, 10);
      if(!Number.isFinite(i)) return;

      if(act === "toggle"){
        __modalCols[i].removed = !__modalCols[i].removed;
        redraw();
        return;
      }

      if(__modalCols[i].removed) return;

      if(act === "up" && i > 0){
        let j = i - 1;
        while(j >= 0 && __modalCols[j].removed) j--;
        if(j >= 0){
          const tmp = __modalCols[j];
          __modalCols[j] = __modalCols[i];
          __modalCols[i] = tmp;
          redraw();
        }
      }
      if(act === "down" && i < __modalCols.length-1){
        let j = i + 1;
        while(j < __modalCols.length && __modalCols[j].removed) j++;
        if(j < __modalCols.length){
          const tmp = __modalCols[j];
          __modalCols[j] = __modalCols[i];
          __modalCols[i] = tmp;
          redraw();
        }
      }
    };

    redraw();
    openModal();
  }

  function bindGlobalEvents(){

    const required = ["btnPickData","btnRun","btnClear","colSearch","btnAll","btnNone",
                      "btnAddAllQueue","btnClearQueue","btnExportQueue",
                      "btnModalCancel","btnModalSave","results","queueList"];
    for(const id of required){
      if(!el(id)){
        console.warn("Elemento não encontrado:", id);
        return;
      }
    }

    el("btnPickData").addEventListener("click", () => QtBridge.pickData());

    el("btnRun").addEventListener("click", () => {
      const rowVar = comboRow.getValue();
      const weightVar = ddWeight.getValue() || "";
      const pctMode = ddPct.getValue() || "col";
      const includeNA = el("includeNA").checked;
      const colVars = selectedCols();

      if(!rowVar || !colVars.length){
        alert("Selecione a variável de Linhas e ao menos uma variável de Colunas.");
        return;
      }

      if(!weightVar){
        const ok = confirm("Nenhum peso foi selecionado. O cruzamento será gerado sem ponderação. Deseja continuar?");
        if(!ok) return;
      }

      el("runStatus").textContent = "Calculando...";
      QtBridge.computeCrosstabs(rowVar, JSON.stringify(colVars), weightVar, includeNA, pctMode);
      setTimeout(()=> el("runStatus").textContent = "", 900);
    });

    el("btnClear").addEventListener("click", () => {
      __results = [];
      el("results").textContent = "Nenhum resultado ainda.";
      __selectedCols.clear();
      renderColList(__lastColFilter);
      el("runStatus").textContent = "";
    });

    el("colSearch").addEventListener("input", (e)=> scheduleRenderColList(e.target.value || ""));
    el("btnAll").addEventListener("click", ()=> markAll(true));
    el("btnNone").addEventListener("click", ()=> markAll(false));

    el("colList").addEventListener("change", (ev)=>{
      const cb = ev.target.closest('input[type="checkbox"][data-var]');
      if(!cb) return;
      const v = cb.dataset.var || "";
      if(!v) return;
      if(cb.checked) __selectedCols.add(v);
      else __selectedCols.delete(v);
    });

    el("btnAddAllQueue").addEventListener("click", ()=>{
      if(!__results.length) return;
      __results.forEach(t => QtBridge.addToQueue(JSON.stringify(t)));
    });
    el("btnClearQueue").addEventListener("click", ()=> QtBridge.clearQueue());
    el("btnExportQueue").addEventListener("click", ()=> QtBridge.exportQueueExcel());

    el("btnModalCancel").addEventListener("click", closeModal);
    el("btnModalSave").addEventListener("click", ()=>{
      if(!__modalTableKey || !__modalMode) return closeModal();

      if(__modalMode === "rows"){
        const newLabels = __modalRowsArr.map(x => x.label);
        applyOrderToTableRows(__modalTableKey, newLabels);
        closeModal();
        return;
      }

      if(__modalMode === "cols"){
        const active = __modalCols.filter(x => !x.removed);
        if(!active.length){
          alert("Você removeu todas as colunas. Deixe ao menos 1 coluna para salvar.");
          return;
        }
        const idxs = active.map(x => x.origIndex);
        applyOrderToTableCols(__modalTableKey, idxs);
        closeModal();
        return;
      }

      closeModal();
    });

    el("results").addEventListener("click", (ev)=>{
      const card = ev.target.closest(".card");
      if(!card) return;
      const key = card.dataset.key;
      const t = __results.find(x => x.key === key);
      if(!t) return;

      const btnOrderR = ev.target.closest('button[data-action="orderRows"]');
      if(btnOrderR){ openOrderRowsModal(t); return; }

      const btnOrderC = ev.target.closest('button[data-action="orderCols"]');
      if(btnOrderC){ openOrderColsModal(t); return; }

      const btnAdd = ev.target.closest('button[data-action="addQueue"]');
      if(btnAdd){ QtBridge.addToQueue(JSON.stringify(t)); return; }

      const btnExp = ev.target.closest('button[data-action="exportOne"]');
      if(btnExp){ QtBridge.exportSingleExcel(JSON.stringify(t)); return; }
    });

    el("results").addEventListener("input", (ev)=>{
      const inp = ev.target.closest('input[data-action="editTitle"]');
      if(!inp) return;
      const card = ev.target.closest(".card");
      if(!card) return;
      const key = card.dataset.key;
      const t = __results.find(x => x.key === key);
      if(t) t.title = inp.value || t.title;
    });

    el("queueList").addEventListener("click", (ev)=>{
      const btn = ev.target.closest('button[data-action="removeQueue"]');
      if(!btn) return;
      const idx = parseInt(btn.dataset.index, 10);
      if(Number.isFinite(idx)) QtBridge.removeQueueItem(idx);
    });
  }

  function onDataLoaded(payloadJson){
    const payload = JSON.parse(payloadJson);
    __columns = payload.columns || [];
    __varLabels = payload.labels || {};
    el("dataStatus").textContent = payload.file_info || "Arquivo carregado.";

    const names = __columns.map(c => c.name);

    comboRow.setItems(names.map(n => ({value:n, label:dispVar(n)})));

    const weightNames = names.filter(isPeso);
    ddWeight.setItems(weightNames.map(n => ({value:n, label:dispVar(n)})));
    ddWeight.setEnabled(weightNames.length > 0);
    ddWeight.setValue("");

    ddPct.setItems([
      {value:"col", label:"Coluna (%)"},
      {value:"row", label:"Linha (%)"},
      {value:"total", label:"Total (%)"},
    ]);
    ddPct.setValue("col");

    __selectedCols.clear();
    __lastColFilter = "";
    el("colSearch").value = "";
    renderColList("");

    __results = [];
    el("results").textContent = "Selecione variáveis e clique em Gerar cruzamentos.";
  }

  function onCrosstabsReady(payloadJson){
    const payload = JSON.parse(payloadJson);
    __results = payload.tables || [];
    renderResults();
  }

  new QWebChannel(qt.webChannelTransport, function(channel){
    QtBridge = channel.objects.QtBridge;

    comboRow = ComboSearch(el("rowCombo"));
    ddWeight = Dropdown(el("weightVarDD"));
    ddPct = Dropdown(el("pctModeDD"));

    QtBridge.dataLoaded.connect(onDataLoaded);
    QtBridge.crosstabsReady.connect(onCrosstabsReady);
    QtBridge.queueState.connect(renderQueueState);

    bindGlobalEvents();
  });
</script>
</body>
</html>
"""




# ==========================
# Helpers (Python)
# ==========================
def _round_half_up(x: float, ndigits: int = 1) -> float:
    q = Decimal("1").scaleb(-ndigits)
    return float(Decimal(str(x)).quantize(q, rounding=ROUND_HALF_UP))


def _canon(s: str) -> str:
    s = (s or "").strip()
    s = unicodedata.normalize("NFD", s)
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    s = s.lower()
    s = re.sub(r"[^a-z0-9]+", " ", s).strip()
    return s


def _is_peso_name(name: str) -> bool:
    return "peso" in _canon(name)


def _safe_str(x: Any) -> str:
    if x is None:
        return ""
    try:
        if isinstance(x, float) and math.isnan(x):
            return ""
    except Exception:
        pass
    return str(x)


def _dedupe_keep_order(values: List[str]) -> List[str]:
    seen = set()
    out: List[str] = []
    for v in values:
        key = _canon(v)
        if key in seen:
            continue
        seen.add(key)
        out.append(v)
    return out


def _apply_priority_sequences(labels: List[str], sequences: List[List[str]], min_present: int = 2) -> Tuple[List[str], List[str]]:
    remaining = labels[:]
    prefix: List[str] = []

    for seq in sequences:
        present: List[str] = []
        for wanted in seq:
            wanted_c = _canon(wanted)
            match = next((lab for lab in remaining if _canon(lab) == wanted_c), None)
            if match is not None:
                present.append(match)

        present = _dedupe_keep_order(present)
        if len(present) < min_present:
            continue

        for lab in present:
            remaining = [x for x in remaining if _canon(x) != _canon(lab)]
        prefix.extend(present)

    return prefix, remaining


def _apply_fixed_bottom_order(labels: List[str], fixed_bottom: List[str]) -> List[str]:
    fixed_canon = [_canon(x) for x in fixed_bottom]
    fixed_canon_set = set(fixed_canon)
    present = labels[:]

    fixed_present: List[str] = []
    for f in fixed_canon:
        for lab in present:
            if _canon(lab) == f and lab not in fixed_present:
                fixed_present.append(lab)
                break

    rest = [lab for lab in present if _canon(lab) not in fixed_canon_set]
    return rest + fixed_present


class QtBridge(QObject):
    dataLoaded = Signal(str)
    crosstabsReady = Signal(str)
    queueState = Signal(str)

    def __init__(self, main_window, web_view: QWebEngineView):
        super().__init__()
        self.main_window = main_window
        self.web_view = web_view

        self.df: Optional[pd.DataFrame] = None
        self.file_info: str = ""
        self.value_labels: Dict[str, Dict[Any, str]] = {}
        self.var_labels: Dict[str, str] = {}

        self.queue: List[Dict[str, Any]] = []

    def _push_queue_state(self) -> None:
        items = [{"title": (it.get("title") or it.get("key") or "").strip()} for it in self.queue]
        self.queueState.emit(json.dumps({"items": items}, ensure_ascii=False))

    @Slot()
    def pickData(self) -> None:
        path, _ = QFileDialog.getOpenFileName(
            self.main_window,
            "Carregar dados",
            "",
            "Dados (*.csv *.xlsx *.xls *.sav);;CSV (*.csv);;Excel (*.xlsx *.xls);;SPSS (*.sav)"
        )
        if not path:
            return

        try:
            self.value_labels = {}
            self.var_labels = {}

            if path.lower().endswith(".csv"):
                df = pd.read_csv(path)

            elif path.lower().endswith((".xlsx", ".xls")):
                df = pd.read_excel(path)

            elif path.lower().endswith(".sav"):
                try:
                    import pyreadstat  # type: ignore
                except Exception:
                    QMessageBox.critical(self.main_window, "Erro", "pyreadstat não está instalado para ler .sav.")
                    return

                df = None
                meta = None
                last_err = None

                try:
                    df, meta = pyreadstat.read_sav(path)  # type: ignore
                except Exception as e:
                    last_err = e
                    msg = str(e)
                    if not (
                        "Unable to convert string" in msg
                        or "invalid byte sequence" in msg
                        or "codec can't decode" in msg
                        or "invalid continuation byte" in msg
                    ):
                        raise

                if df is None or meta is None:
                    file_enc = None
                    try:
                        _, m0 = pyreadstat.read_sav(path, metadataonly=True)  # type: ignore
                        file_enc = (getattr(m0, "file_encoding", None) or "").strip() or None
                    except Exception:
                        file_enc = None

                    candidates = []
                    for enc in [file_enc, "utf-8", "cp1252", "latin1", "iso-8859-1", "iso-8859-15", "utf-16le", "ucs-2le"]:
                        if enc and enc not in candidates:
                            candidates.append(enc)

                    for enc in candidates:
                        try:
                            df, meta = pyreadstat.read_sav(path, encoding=enc)  # type: ignore
                            break
                        except Exception as e:
                            last_err = e
                            msg = str(e)
                            if (
                                "Unable to convert string" in msg
                                or "invalid byte sequence" in msg
                                or "codec can't decode" in msg
                                or "invalid continuation byte" in msg
                            ):
                                continue
                            raise

                if df is None or meta is None:
                    raise last_err or Exception("Falha ao ler o arquivo .sav.")

                try:
                    for var, m in (meta.variable_value_labels or {}).items():
                        self.value_labels[var] = dict(m)
                except Exception:
                    self.value_labels = {}

                labels_map: Dict[str, str] = {}
                try:
                    m = getattr(meta, "column_names_to_labels", None)
                    if isinstance(m, dict):
                        labels_map = {str(k): str(v) for k, v in m.items() if v}
                    else:
                        cn = getattr(meta, "column_names", None)
                        cl = getattr(meta, "column_labels", None)
                        if cn and cl and len(cn) == len(cl):
                            labels_map = {str(k): str(v) for k, v in zip(cn, cl) if v}
                except Exception:
                    labels_map = {}

                self.var_labels = labels_map

            else:
                QMessageBox.warning(self.main_window, "Formato não suportado", "Formato de arquivo não suportado.")
                return

            self.df = df
            self.file_info = path

            cols = [{"name": str(c), "dtype": str(dtype)} for c, dtype in df.dtypes.items()]
            labels = {str(c): (self.var_labels.get(str(c)) or "") for c in df.columns}
            payload = {"file_info": self.file_info, "columns": cols, "labels": labels}
            self.dataLoaded.emit(json.dumps(payload, ensure_ascii=False))

        except Exception as e:
            QMessageBox.critical(self.main_window, "Erro ao carregar", f"Falha ao carregar dados:\n{e}")

    def _prepare_df(self, row_var: str, col_var: str, weight_var: str, include_na: bool) -> Tuple[pd.DataFrame, Optional[pd.Series]]:
        assert self.df is not None

        needed = [row_var, col_var]
        if weight_var:
            needed.append(weight_var)

        src = self.df.loc[:, needed]
        if include_na:
            df = src.copy()
        else:
            mask = src[row_var].notna() & src[col_var].notna()
            df = src.loc[mask].copy()

        for c in (row_var, col_var):
            m = self.value_labels.get(c)
            if m:
                df[c] = df[c].map(lambda v, labels=m: labels.get(v, v))

        if include_na:
            df[row_var] = df[row_var].where(df[row_var].notna(), "(Sem resposta)")
            df[col_var] = df[col_var].where(df[col_var].notna(), "(Sem resposta)")

        df[row_var] = df[row_var].map(_safe_str)
        df[col_var] = df[col_var].map(_safe_str)

        w = None
        if weight_var:
            w = pd.to_numeric(df[weight_var], errors="coerce").fillna(0.0)
        return df, w

    def _weighted_crosstab(self, df: pd.DataFrame, row_var: str, col_var: str, w: Optional[pd.Series]) -> pd.DataFrame:
        if w is None:
            return pd.crosstab(df[row_var], df[col_var], dropna=False).astype(float)

        ct = pd.crosstab(
            df[row_var],
            df[col_var],
            values=w,
            aggfunc="sum",
            dropna=False,
        ).fillna(0.0)
        return ct.astype(float)

    def _unweighted_bases(self, df: pd.DataFrame, row_var: str, col_var: str) -> pd.DataFrame:
        return pd.crosstab(df[row_var], df[col_var], dropna=False)

    def _pct_matrix(self, ct: pd.DataFrame, mode: str) -> pd.DataFrame:
        if mode == "col":
            denom = ct.sum(axis=0).replace(0, float("nan"))
            p = ct.divide(denom, axis=1) * 100.0
        elif mode == "row":
            denom = ct.sum(axis=1).replace(0, float("nan"))
            p = ct.divide(denom, axis=0) * 100.0
        elif mode == "total":
            denom = float(ct.values.sum())
            p = (ct / denom) * 100.0 if denom else ct * 0.0
        else:
            p = ct * 0.0

        p = p.fillna(0.0)
        return p.apply(lambda s: s.map(lambda v: _round_half_up(float(v), 1)))

    def _pct_total_col(self, ct: pd.DataFrame) -> List[float]:
        row_sums = ct.sum(axis=1)
        grand = float(ct.values.sum()) if ct.values.size else 0.0
        if grand <= 0:
            return [0.0 for _ in row_sums]
        out = (row_sums / grand) * 100.0
        return [_round_half_up(float(v), 1) for v in out.tolist()]

    def _sort_rows_by_total_desc(
        self,
        ct: pd.DataFrame,
        fixed_bottom: List[str],
        exclude_labels: Optional[List[str]] = None,
        priority_sequences: Optional[List[List[str]]] = None,
    ) -> List[str]:
        original_labels = [str(x) for x in list(ct.index.astype(str))]
        exclude_canon = {_canon(x) for x in (exclude_labels or [])}
        priority_sequences = priority_sequences or []

        priority_prefix, remaining_labels = _apply_priority_sequences(original_labels, priority_sequences)

        grand = float(ct.values.sum()) if ct.values.size else 0.0
        if grand > 0:
            row_sums = ct.sum(axis=1)
            pct = (row_sums / grand) * 100.0
        else:
            pct = pd.Series(0.0, index=ct.index)

        pinned = [lab for lab in remaining_labels if _canon(lab) in exclude_canon]
        sortable = [lab for lab in remaining_labels if _canon(lab) not in exclude_canon]

        if grand > 0:
            sortable.sort(key=lambda lab: float(pct.get(lab, 0.0)), reverse=True)

        ordered_tail = _apply_fixed_bottom_order(sortable + pinned, fixed_bottom)
        return priority_prefix + ordered_tail

    @Slot(str, str, str, bool, str)
    def computeCrosstabs(self, rowVar: str, colVarsJson: str, weightVar: str, includeNA: bool, pctMode: str) -> None:
        if self.df is None:
            QMessageBox.warning(self.main_window, "Aviso", "Carregue um arquivo primeiro.")
            return

        try:
            col_vars = json.loads(colVarsJson)
            if not isinstance(col_vars, list) or not col_vars:
                raise ValueError("Lista de colunas inválida.")

            rowVar = (rowVar or "").strip()
            weightVar = (weightVar or "").strip()

            if weightVar and not _is_peso_name(weightVar):
                weightVar = ""
            if weightVar and weightVar not in self.df.columns:
                weightVar = ""

            tables: List[Dict[str, Any]] = []

            rowVarDisplay = (self.var_labels.get(rowVar) or "").strip() or rowVar

            for colVar in col_vars:
                colVar = (colVar or "").strip()
                if not colVar:
                    continue

                df_prep, w = self._prepare_df(rowVar, colVar, weightVar, includeNA)
                base_ct = self._unweighted_bases(df_prep, rowVar, colVar)
                if base_ct.empty:
                    continue

                base_row_labels = [str(x) for x in base_ct.index.tolist()]
                base_col_labels = [str(x) for x in base_ct.columns.tolist()]
                base_total_n = int(base_ct.to_numpy().sum())

                ct = self._weighted_crosstab(df_prep, rowVar, colVar, w)
                ct = ct.reindex(index=base_row_labels, columns=base_col_labels, fill_value=0.0)

                ordered_rows = self._sort_rows_by_total_desc(
                    ct,
                    FIXED_BOTTOM_LABELS,
                    exclude_labels=VBA_EXCLUDE_LABELS,
                    priority_sequences=PRIORITY_ROW_SEQUENCES,
                )
                ct = ct.reindex(index=ordered_rows, fill_value=0.0)
                base_ct = base_ct.reindex(index=ordered_rows, columns=base_col_labels, fill_value=0)

                pct_m = self._pct_matrix(ct, pctMode)
                pct_total = self._pct_total_col(ct)

                col_labels = base_col_labels
                row_labels = [str(x) for x in ct.index.tolist()]
                base_col_n_aligned = base_ct.sum(axis=0).astype(int).tolist()

                weighted_total = float(ct.values.sum())
                base_display = f"Base (N): {base_total_n}"
                if weightVar:
                    base_display += f" | Ponderada: {weighted_total:.1f}"
                else:
                    base_display += " | Sem ponderação"

                colVarDisplay = (self.var_labels.get(colVar) or "").strip() or colVar
                title = f"{rowVarDisplay} × {colVarDisplay}"

                tables.append({
                    "key": f"{rowVar}__{colVar}",
                    "title": title,
                    "row_var": rowVar,
                    "col_var": colVar,
                    "row_var_display": rowVarDisplay,
                    "col_var_display": colVarDisplay,
                    "pct_mode": pctMode,
                    "weight_var": weightVar,
                    "include_na": bool(includeNA),

                    "fixed_bottom_labels": FIXED_BOTTOM_LABELS,
                    "priority_row_sequences": PRIORITY_ROW_SEQUENCES,

                    "row_labels": row_labels,
                    "col_labels": col_labels,

                    "pct_total": pct_total,
                    "pct_matrix": pct_m.values.tolist(),

                    "base_total_n": int(base_total_n),
                    "base_col_n": base_col_n_aligned,

                    "base_display": base_display,
                })

            self.crosstabsReady.emit(json.dumps({"tables": tables}, ensure_ascii=False))

        except Exception as e:
            QMessageBox.critical(self.main_window, "Erro", f"Falha ao calcular cruzamentos:\n{e}")

    @Slot(str)
    def addToQueue(self, tableJson: str) -> None:
        try:
            t = json.loads(tableJson)
            if not isinstance(t, dict):
                return
            self.queue.append(t)
            self._push_queue_state()
        except Exception as e:
            QMessageBox.critical(self.main_window, "Erro", f"Falha ao adicionar à pilha:\n{e}")

    @Slot(int)
    def removeQueueItem(self, index: int) -> None:
        try:
            if 0 <= index < len(self.queue):
                del self.queue[index]
            self._push_queue_state()
        except Exception as e:
            QMessageBox.critical(self.main_window, "Erro", f"Falha ao remover item da pilha:\n{e}")

    @Slot()
    def clearQueue(self) -> None:
        self.queue.clear()
        self._push_queue_state()

    # ==========================
    # Excel Export (borders minimal)
    # ==========================
    def _excel_styles(self) -> Dict[str, Any]:
        font_title = Font(name="DIN", size=12, bold=False)
        font_header = Font(name="DIN", size=10, bold=False)
        font_rowlab = Font(name="DIN", size=10, bold=False)
        font_num = Font(name="DIN", size=10, bold=False)

        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

        thin_black = Side(style="thin", color="000000")
        thick_black = Side(style="thick", color="000000")
        none_side = Side(style=None)

        return {
            "font_title": font_title,
            "font_header": font_header,
            "font_rowlab": font_rowlab,
            "font_num": font_num,
            "align_center": align_center,
            "align_left": align_left,
            "thin": thin_black,
            "thick": thick_black,
            "none": none_side,
        }

    def _compute_colA_width(self, tables: List[Dict[str, Any]]) -> float:
        max_len = len("Categoria")
        for t in tables:
            for lab in (t.get("row_labels") or []):
                max_len = max(max_len, len(str(lab)))
            max_len = max(max_len, len("Base"))
        return min(max(12.0, max_len * 1.05), 70.0)

    def _apply_table_borders_minimal(
        self,
        ws,
        top: int,
        bottom: int,
        left: int,
        right: int,
        none: Side,
        thin: Side,
        thick: Side
    ) -> None:
        for r in range(top, bottom + 1):
            for c in range(left, right + 1):
                l = none
                t = none
                b = none
                rr = thin

                if c == left:
                    l = thick
                if c == right:
                    rr = thick
                if r == top:
                    t = thick
                if r == bottom:
                    b = thick

                if r == top:
                    l = thick
                    rr = thick
                    t = thick
                    b = thick

                ws.cell(r, c).border = Border(left=l, right=rr, top=t, bottom=b)

    def _write_table_block(self, ws, start_row: int, t: Dict[str, Any], styles: Dict[str, Any]) -> int:
        title = (t.get("title") or "").strip()
        row_labels: List[str] = t["row_labels"]
        col_labels: List[str] = t["col_labels"]
        pct_total: List[float] = t["pct_total"]
        pct_matrix: List[List[float]] = t["pct_matrix"]
        base_total_n: int = int(t["base_total_n"])
        base_col_n: List[int] = [int(x) for x in t["base_col_n"]]

        last_col = 2 + len(col_labels)

        r_title = start_row
        ctitle = ws.cell(r_title, 1, title)
        ctitle.font = styles["font_title"]
        ctitle.alignment = styles["align_center"]
        ws.merge_cells(start_row=r_title, start_column=1, end_row=r_title, end_column=last_col)

        r_blank = r_title + 1

        r_head = r_blank + 1
        headers = ["Categoria", "Total"] + col_labels
        for j, h in enumerate(headers, start=1):
            c = ws.cell(r_head, j, h)
            c.font = styles["font_header"]
            c.alignment = styles["align_left"] if j == 1 else styles["align_center"]

        r0 = r_head + 1
        for i, lab in enumerate(row_labels):
            rr = r0 + i

            c0 = ws.cell(rr, 1, lab)
            c0.font = styles["font_rowlab"]
            c0.alignment = styles["align_left"]

            cT = ws.cell(rr, 2, float(pct_total[i]))
            cT.font = styles["font_num"]
            cT.alignment = styles["align_center"]
            cT.number_format = "0.0"

            for j in range(len(col_labels)):
                v = float(pct_matrix[i][j])
                cc = ws.cell(rr, 3 + j, v)
                cc.font = styles["font_num"]
                cc.alignment = styles["align_center"]
                cc.number_format = "0.0"

        r_base = r0 + len(row_labels)
        cB0 = ws.cell(r_base, 1, "Base")
        cB0.font = styles["font_rowlab"]
        cB0.alignment = styles["align_left"]

        cBT = ws.cell(r_base, 2, int(base_total_n))
        cBT.font = styles["font_rowlab"]
        cBT.alignment = styles["align_center"]
        cBT.number_format = "0"

        for j in range(len(col_labels)):
            cc = ws.cell(r_base, 3 + j, int(base_col_n[j]))
            cc.font = styles["font_rowlab"]
            cc.alignment = styles["align_center"]
            cc.number_format = "0"

        self._apply_table_borders_minimal(
            ws,
            top=r_head,
            bottom=r_base,
            left=1,
            right=last_col,
            none=styles["none"],
            thin=styles["thin"],
            thick=styles["thick"],
        )

        return r_base

    @Slot()
    def exportQueueExcel(self) -> None:
        if not self.queue:
            QMessageBox.information(self.main_window, "Excel", "A pilha está vazia.")
            return

        save_path, _ = QFileDialog.getSaveFileName(
            self.main_window,
            "Salvar Excel (Pilha)",
            "cruzamentos.xlsx",
            "Excel (*.xlsx)"
        )
        if not save_path:
            return
        if not save_path.lower().endswith(".xlsx"):
            save_path += ".xlsx"

        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Cruzamentos"

            styles = self._excel_styles()

            ws.column_dimensions["A"].width = self._compute_colA_width(self.queue)

            max_cols = 0
            for t in self.queue:
                max_cols = max(max_cols, 2 + len(t.get("col_labels") or []))
            for c in range(2, max_cols + 1):
                ws.column_dimensions[get_column_letter(c)].width = 8.89

            r = 1
            for t in self.queue:
                last_used = self._write_table_block(ws, r, t, styles)
                r = last_used + 3

            wb.save(save_path)
            QMessageBox.information(self.main_window, "Excel", f"Arquivo salvo com {len(self.queue)} tabela(s).")

        except Exception as e:
            QMessageBox.critical(self.main_window, "Erro", f"Falha ao exportar Excel:\n{e}")

    @Slot(str)
    def exportSingleExcel(self, tableJson: str) -> None:
        try:
            t = json.loads(tableJson)
            if not isinstance(t, dict):
                QMessageBox.warning(self.main_window, "Excel", "Tabela inválida para exportação.")
                return

            save_path, _ = QFileDialog.getSaveFileName(
                self.main_window,
                "Salvar Excel (1 tabela)",
                "cruzamento.xlsx",
                "Excel (*.xlsx)"
            )
            if not save_path:
                return
            if not save_path.lower().endswith(".xlsx"):
                save_path += ".xlsx"

            wb = Workbook()
            ws = wb.active
            ws.title = "Cruzamento"

            styles = self._excel_styles()

            ws.column_dimensions["A"].width = self._compute_colA_width([t])
            max_cols = 2 + len(t.get("col_labels") or [])
            for c in range(2, max_cols + 1):
                ws.column_dimensions[get_column_letter(c)].width = 8.89

            self._write_table_block(ws, 1, t, styles)

            wb.save(save_path)
            QMessageBox.information(self.main_window, "Excel", "Arquivo salvo.")

        except Exception as e:
            QMessageBox.critical(self.main_window, "Erro", f"Falha ao exportar:\n{e}")


class CrosstabsPage(QWidget):
    def __init__(self, main_window, profile: Optional[QWebEngineProfile] = None):
        super().__init__()
        self.main_window = main_window

        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(0)

        topbar = TopBar(
            title="Funcionalidade: Cruzamentos",
            icon_path=resource_path("assets/icons/table.ico"),
            back_icon_path=resource_path("assets/icons/arrow-left.ico")
        )

        self.btn_back = topbar.btn_back
        layout.addWidget(topbar)

        self.view = QWebEngineView(self)
        layout.addWidget(self.view, 1)

        if profile is not None:
            page = QWebEnginePage(profile, self.view)
            self.view.setPage(page)

        self.view.setHtml(HTML, QUrl("https://crosstabs.app.local/"))

        self.bridge = QtBridge(main_window, self.view)
        channel = QWebChannel(self.view.page())
        channel.registerObject("QtBridge", self.bridge)
        self.view.page().setWebChannel(channel)

        self.view.loadFinished.connect(lambda ok: self.bridge._push_queue_state())

    def cleanup(self):
        pass